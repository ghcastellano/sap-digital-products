from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

SRC = Path('/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento/Digital Products - Prior Materials Review .pptx')
OUT = Path('/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento/Digital Products - Prior Materials Review - complete-beautiful.pptx')

NAVY = RGBColor(0x00, 0x20, 0x60)
BLUE = RGBColor(0x00, 0x5B, 0x96)
CYAN = RGBColor(0x19, 0xA3, 0xFC)
GREEN = RGBColor(0x00, 0xA4, 0x43)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHAR = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0xF5, 0xF7, 0xFA)
MID = RGBColor(0x6B, 0x72, 0x80)
DPBG1 = RGBColor(0xE8, 0xF1, 0xFB)
DPBG2 = RGBColor(0xDC, 0xEB, 0xFA)

src = Presentation(str(SRC))
prs = Presentation()
prs.slide_width = src.slide_width
prs.slide_height = src.slide_height


def first_text(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                return t
    return ''


def all_texts(slide):
    vals = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                vals.append(t)
    return vals


def draw_header(slide, title, subtitle=''):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.85))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY
    band.line.fill.background()

    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.78), prs.slide_width, Inches(0.07))
    acc.fill.solid(); acc.fill.fore_color.rgb = CYAN
    acc.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(11.8), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    r.font.name = 'Aptos'

    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(12.0), Inches(0.3))
        sp = sb.text_frame.paragraphs[0]
        sr = sp.add_run(); sr.text = subtitle; sr.font.size = Pt(9); sr.font.color.rgb = MID; sr.font.name = 'Aptos'


# Slide 1 cover
s = prs.slides.add_slide(prs.slide_layouts[6])
# background
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0x0B, 0x1E, 0x3A); bg.line.fill.background()
# decorative bars
for i, c in enumerate([BLUE, CYAN, GREEN]):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5 + i*0.28), Inches(0.8), Inches(0.2), Inches(5.7))
    sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.line.fill.background()

title_text = all_texts(src.slides[0])
main = title_text[0] if title_text else 'Digital Products'
sub = title_text[1] if len(title_text) > 1 else 'Review of existing assets'

tb = s.shapes.add_textbox(Inches(1.5), Inches(1.7), Inches(10.5), Inches(2.8))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = main; r.font.size = Pt(46); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = 'Aptos Display'
p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(20); r2.font.color.rgb = RGBColor(0xC7, 0xD9, 0xF5); r2.font.name = 'Aptos'

# Slides 2-3 narrative cards
for i in [1, 2]:
    ss = prs.slides.add_slide(prs.slide_layouts[6])
    texts = all_texts(src.slides[i])
    draw_header(ss, texts[0][:120] if texts else 'Context', 'Strategic framing and recommendation')

    y = 1.35
    for idx, t in enumerate(texts[1:9], 1):
        h = 0.55 if len(t) < 90 else 0.85
        card = ss.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(y), Inches(12.0), Inches(h))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

        tx = ss.shapes.add_textbox(Inches(0.85), Inches(y + 0.08), Inches(11.5), Inches(h - 0.1))
        p = tx.text_frame.paragraphs[0]
        rr = p.add_run(); rr.text = t[:360]; rr.font.size = Pt(12 if idx <= 2 else 10); rr.font.color.rgb = CHAR; rr.font.name = 'Aptos'
        y += h + 0.12
        if y > 6.6:
            break

# Table slides 4-6: copy all text but with premium styling
for src_idx in [3, 4, 5]:
    src_slide = src.slides[src_idx]
    ss = prs.slides.add_slide(prs.slide_layouts[6])

    title = first_text(src_slide) or 'Reference Matrix'
    draw_header(ss, title, 'All references consolidated for reuse planning')

    st = None
    for sh in src_slide.shapes:
        if getattr(sh, 'has_table', False) and sh.has_table:
            st = sh.table
            break
    if st is None:
        continue

    rows, cols = len(st.rows), len(st.columns)
    tbl_shape = ss.shapes.add_table(rows, cols, Inches(0.2), Inches(1.2), Inches(12.9), Inches(6.0))
    t = tbl_shape.table

    widths = [Inches(1.1), Inches(3.0), Inches(2.8), Inches(2.2), Inches(3.8)] if cols == 5 else [Inches(1.4), Inches(4.2), Inches(3.9), Inches(3.0)]
    for ci in range(min(cols, len(widths))):
        t.columns[ci].width = int(widths[ci])

    for r in range(rows):
        for c in range(cols):
            src_text = st.cell(r, c).text
            cell = t.cell(r, c)
            cell.text = src_text
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)

            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = BLUE if not (cols == 5 and c == 4) else NAVY
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for run in p.runs:
                        run.font.size = Pt(8)
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                        run.font.name = 'Aptos'
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            else:
                base = WHITE if r % 2 else GRAY
                if cols == 5 and c == 4:
                    base = DPBG1 if r % 2 else DPBG2
                cell.fill.solid(); cell.fill.fore_color.rgb = base
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(7)
                        run.font.color.rgb = CHAR
                        run.font.name = 'Aptos'
                cell.vertical_anchor = MSO_ANCHOR.TOP

    # footer tag
    tag = ss.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.4), Inches(7.0), Inches(2.4), Inches(0.3))
    tag.fill.solid(); tag.fill.fore_color.rgb = RGBColor(0xEE, 0xF6, 0xFF); tag.line.fill.background()
    tp = tag.text_frame.paragraphs[0]
    tr = tp.add_run(); tr.text = 'Complete + Stable Build'; tr.font.size = Pt(8); tr.font.color.rgb = BLUE; tr.font.bold = True; tr.font.name = 'Aptos'
    tp.alignment = PP_ALIGN.CENTER

prs.save(str(OUT))
print(str(OUT))
