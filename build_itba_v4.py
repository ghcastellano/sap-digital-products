"""
build_itba_v4.py — ITBA deck mirroring CS+VD structure exactly (2 proposals).

Structure (14 slides, mirrors CS+VD minus Value Lens):
  1.  Cover
  2.  PROPOSAL 1 cards (1 OVS umbrella + 3 VS + Enabling band) — 3-VS proposal
  3-5. 3 VS detail (Network & Asset Operations · Corporate & Regulatory · Iberdrola)
  6.  PROPOSAL 2 cards (1 OVS umbrella + 4 DVS + Enabling band) — 4-DVS proposal
  7.  PROPOSAL 2 RATIONALE (Big Picture — 4 DVS swim lanes)
  8.  Value Stream Framing
  9.  OVS detail (ITBA umbrella OVS)
  10-13. 4 DVS detail (Engineering & Asset · Network Ops · Corporate & Reg · Iberdrola)
  14. Enabling Team

ITBA Proposal 1: 3 consolidated value streams (simpler, fewer PMs)
ITBA Proposal 2: 4 development value streams (current sub-team alignment, recommended)
"""
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC = "Material em andamento/Customer Services - Proposed Value Streams + Value Definition.pptx"
DST = "Material em andamento/IT Business Applications - Proposed Value Streams + Value Definition.pptx"

NAVY  = RGBColor(0x0F, 0x1B, 0x4F)
BLUE2 = RGBColor(0x2C, 0x57, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x2C, 0x3E, 0x50)
GRAY  = RGBColor(0x95, 0xA5, 0xA6)
RED   = RGBColor(0x98, 0x00, 0x03)
ENAB  = RGBColor(0x6C, 0x7A, 0x89)
LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)
LIGHTBG_AREA = RGBColor(0xEE, 0xEE, 0xEE)
GREEN = RGBColor(0x9C, 0xC4, 0x4A)
TEAL = RGBColor(0x0B, 0x76, 0xA0)


# ---------------- helpers ----------------

def add_text(slide, x, y, w, h, text, *, size=10, bold=False, italic=False,
             color=DARK, name="Arial", align=PP_ALIGN.LEFT, anchor=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    if anchor: tf.vertical_anchor = anchor
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.name = name; r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color
    return tb


def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.5)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def copy_bg(prs, src_idx, new_slide):
    src = prs.slides[src_idx]
    src_bg = src._element.find(qn('p:cSld') + '/' + qn('p:bg'))
    new_cSld = new_slide._element.find(qn('p:cSld'))
    if src_bg is not None and new_cSld is not None:
        existing = new_cSld.find(qn('p:bg'))
        if existing is not None: new_cSld.remove(existing)
        new_cSld.insert(0, deepcopy(src_bg))


def replace_textbox(shape, text):
    tf = shape.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        f0 = tf.paragraphs[0].runs[0].font
        sz = f0.size; bold = f0.bold; italic = f0.italic
        try: color = f0.color.rgb if f0.color and f0.color.rgb else None
        except Exception: color = None
        name = f0.name
    else:
        sz, bold, italic, color, name = None, None, None, None, "Arial"
    tf.clear()
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        if name: r.font.name = name
        if sz: r.font.size = sz
        if bold is not None: r.font.bold = bold
        if italic is not None: r.font.italic = italic
        if color: r.font.color.rgb = color


def edit_by_idx(slide, mapping):
    by_idx = {i: s for i, s in enumerate(slide.shapes)}
    for i, text in mapping.items():
        if i in by_idx and by_idx[i].has_text_frame:
            replace_textbox(by_idx[i], text)


def delete_slide_safe(prs, idx):
    slide_id = list(prs.slides._sldIdLst)[idx]
    rId = slide_id.rId
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(slide_id)


def clear_slide(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


# ---------------- Cards Slide builder (comprehensive view) ----------------

def build_cards_slide(slide, *, proposal_label, ovs_title, ovs_tag,
                       ovs_domain, ovs_desc, banner_text, dvs_list, n_dvs):
    """Comprehensive cards slide (mirror CS slide 2 layout): umbrella OVS card on top,
    N DVS cards below, Enabling Team band at bottom."""
    # Header
    add_text(slide, 0.10, 0.10, 12.53, 0.38,
             proposal_label, size=20, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    add_text(slide, 0.40, 0.53, 12.53, 0.23,
             "IT BUSINESS APPLICATIONS",
             size=11, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

    # OVS umbrella card
    add_rect(slide, 0.51, 1.37, 12.33, 1.21, fill=WHITE,
             line=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    add_rect(slide, 0.49, 1.32, 12.37, 0.50, fill=NAVY,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    add_text(slide, 0.71, 1.40, 9.23, 0.40, ovs_title,
             size=13, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    add_text(slide, 9.74, 1.42, 2.90, 0.36, ovs_tag,
             size=9, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_text(slide, 0.71, 1.94, 11.93, 0.20, ovs_domain,
             size=8.5, bold=True, italic=True, color=DARK)
    add_text(slide, 0.71, 2.16, 11.93, 0.34, ovs_desc,
             size=11, italic=True, color=DARK)

    # Banner
    add_rect(slide, 0.51, 2.76, 12.33, 0.55, fill=NAVY,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    add_text(slide, 0.51, 2.76, 12.33, 0.55, banner_text,
             size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # DVS cards in a row
    margin = 0.51
    gap = 0.10
    total_w = 13.33 - 2 * margin
    card_w = (total_w - (n_dvs - 1) * gap) / n_dvs

    card_y = 3.55
    card_h = 2.45

    for i, (title, domain, desc, apps) in enumerate(dvs_list):
        x = margin + i * (card_w + gap)

        # Body
        add_rect(slide, x, card_y + 0.08, card_w, card_h - 0.08, fill=WHITE,
                 line=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)

        # Header bar
        add_rect(slide, x - 0.02, card_y, card_w + 0.04, 0.78, fill=NAVY,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)

        # DVS label
        add_text(slide, x, card_y + 0.05, card_w, 0.22,
                 "DEVELOPMENT VALUE STREAM", size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # Title
        add_text(slide, x, card_y + 0.27, card_w, 0.50, title,
                 size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        inner_y = card_y + 0.92
        # Domain
        add_text(slide, x + 0.10, inner_y, card_w - 0.20, 0.22, domain,
                 size=8, bold=True, italic=True, color=DARK, align=PP_ALIGN.CENTER)
        # Description
        add_text(slide, x + 0.10, inner_y + 0.25, card_w - 0.20, 0.50, desc,
                 size=9.5, italic=True, color=DARK, align=PP_ALIGN.CENTER)
        # Divider
        add_rect(slide, x + 0.20, inner_y + 0.85, card_w - 0.40, 0.04, fill=NAVY)
        # APPS label
        add_text(slide, x + 0.10, inner_y + 0.93, card_w - 0.20, 0.22,
                 "EXAMPLE APPLICATIONS", size=8, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        # Apps text
        add_text(slide, x + 0.10, inner_y + 1.18, card_w - 0.20, 0.30,
                 apps, size=9.5, color=DARK, align=PP_ALIGN.CENTER)

    # Enabling Team band
    et_x = 0.51
    et_y = 6.20
    et_w = total_w
    et_h = 0.90
    add_rect(slide, et_x, et_y, et_w, et_h, fill=RED,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    add_text(slide, et_x + 0.15, et_y + 0.06, 3.50, 0.22,
             "ENABLING TEAM", size=11, bold=True, color=WHITE)
    add_text(slide, et_x + 0.15, et_y + 0.30, 3.50, 0.55,
             f"Shared services · supports the {n_dvs} DVS\nNo Product Manager · governed by Director ITBA",
             size=9, color=WHITE)
    add_text(slide, et_x + 3.85, et_y + 0.10, et_w - 3.95, 0.20,
             "FOUNDATION SERVICES", size=9, bold=True, color=WHITE)
    add_text(slide, et_x + 3.85, et_y + 0.32, et_w - 3.95, 0.55,
             "SWF (Software Factory)  ·  Demand Intake (PDD · Jira · IT Now)  ·  "
             "Test & Release Coordination  ·  Third-Party Risk Gating  ·  ALM Tooling",
             size=9, color=WHITE)


# ---------------- Big Picture (Proposal 2 Rationale) ----------------

def build_big_picture(slide):
    """Big Picture for ITBA Proposal 2 — 4 DVS swim lanes (no chevrons since multi-OVS)."""
    add_text(slide, 0.40, 0.10, 12.53, 0.38,
             "PROPOSAL 2 – RATIONALE",
             size=18, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    add_text(slide, 0.40, 0.58, 12.53, 0.23,
             "Four ITBA Development Value Streams supporting multiple business operational value streams · "
             "one shared Enabling Team",
             size=10.5, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

    # 4 DVS swim lanes
    LANES = [
        ("Engineering & Asset Mgmt", "Asset / CAD / GIS",
         "Engineer / planner designs with reliable data",
         "ArcFM · ProjectWise · Bentley · CYME",
         "George Porto + crossover", "~6 FTE"),
        ("Network Operations & Control", "SCADA / Real-time",
         "Grid operator prevents and minimizes outages",
         "Echelon · Cimplicity · SLAM · Switching Orders",
         "Bruno Moreto + Andrew McGibbon", "~7 FTE"),
        ("Corporate & Regulatory", "Workflow / Compliance",
         "HR / employee / regulatory user self-service",
         "OpenText · Rate Case Advisor · Utilisphere",
         "Sarath Garikapati + Amit", "~8 FTE"),
        ("Iberdrola Corporate Services", "SAP modules",
         "Spain finance / procurement / budget",
         "SAP FI · PO · Budget · S/4HANA",
         "Vanessa Colombo + Michael Dubois", "~5 FTE"),
    ]

    lane_top = 1.10
    lane_h = 1.20
    lane_gap = 0.10

    # DVS color palette
    DVS_COLORS = [
        RGBColor(0x16, 0xA0, 0x85),  # green
        RGBColor(0xD3, 0x54, 0x00),  # orange
        RGBColor(0x8E, 0x44, 0xAD),  # purple
        RGBColor(0x2C, 0x57, 0x8B),  # blue
    ]

    for i, (name, tag, desc, apps, lead, fte) in enumerate(LANES):
        y = lane_top + i * (lane_h + lane_gap)
        color = DVS_COLORS[i]

        # Light bg lane
        add_rect(slide, 0.27, y - 0.05, 12.79, lane_h, fill=LIGHTBG_AREA,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)

        # Left: DVS name box (color-coded)
        add_rect(slide, 0.40, y, 2.50, 1.10, fill=color,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        add_text(slide, 0.42, y + 0.05, 2.46, 0.20, "DVS",
                 size=7.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, 0.42, y + 0.22, 2.46, 0.45, name,
                 size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 0.42, y + 0.70, 2.46, 0.18, tag,
                 size=7.5, italic=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, 0.42, y + 0.88, 2.46, 0.18, f"{lead} · {fte}",
                 size=7, italic=True, color=WHITE, align=PP_ALIGN.CENTER)

        # Middle: description
        add_text(slide, 3.00, y + 0.08, 4.50, 0.20, "OUTCOME",
                 size=7.5, bold=True, color=GRAY)
        add_text(slide, 3.00, y + 0.28, 4.50, 0.55, desc,
                 size=10, bold=True, italic=True, color=DARK)

        # Right: apps
        add_text(slide, 3.00, y + 0.70, 4.50, 0.20, "APPLICATIONS",
                 size=7.5, bold=True, color=GRAY)
        add_text(slide, 3.00, y + 0.88, 4.50, 0.30, apps,
                 size=9.5, color=DARK)

        # Far right: business OVS supported
        add_text(slide, 7.70, y + 0.08, 5.30, 0.20,
                 "BUSINESS OVS SUPPORTED",
                 size=7.5, bold=True, color=GRAY)
        ovs_text = {
            0: "Network Build · Asset Replacement · Capital Planning · Asset Decommissioning",
            1: "Outage Response · Grid Optimization · Routine Switching · Storm Restoration",
            2: "Source-to-Pay · Hire-to-Retire · Rate Case Filing · Regulatory Reporting",
            3: "Procure-to-Pay (Spain) · Period Close (Spain) · Plan-to-Budget (Spain)",
        }[i]
        add_text(slide, 7.70, y + 0.28, 5.30, 0.80, ovs_text,
                 size=9, italic=True, color=DARK)

    # Enabling Team band
    et_y = lane_top + 4 * (lane_h + lane_gap) + 0.10
    add_rect(slide, 0.40, et_y, 12.53, 0.50, fill=RED,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.20)
    add_text(slide, 0.40, et_y, 12.53, 0.50,
             "Enabling Team  ·  shared services consumed by the 4 DVS  "
             "(SWF · demand intake · test/release · third-party risk · ALM tooling)",
             size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Footer
    add_text(slide, 0.40, et_y + 0.55, 12.53, 0.20,
             "Pattern: each ITBA Development Value Stream supports multiple business OVS within its domain · "
             "shared Enabling Team underneath all DVS",
             size=8, italic=True, color=GRAY)


# ---------------- DVS detail content (Prop 1 — 3 VS) ----------------

PROP1_VS_DETAIL = [
    {  # VS 1 — Network & Asset Operations (consolidated)
        0: "DEVELOPMENT VALUE STREAM",
        1: "NETWORK & ASSET OPERATIONS — UNIFIED ENGINEERING + GRID OPS PLATFORM",
        3: "ITBA · George Porto + Bruno Moreto + Andrew McGibbon (combined leadership)",
        5: "Consolidated proposal · combines Engineering & Asset + Network Operations",
        8: ("Single value stream covering all field-and-asset technology — engineering content, "
            "spatial data, grid control. Combines what Proposal 2 splits into Engineering & Asset Mgmt "
            "and Network Operations & Control DVSs."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  ArcFM (GIS / asset management)\n"
             "•  ProjectWise (engineering content)\n"
             "•  Bentley · Microstation (CAD)\n"
             "•  CYME (power systems analysis)\n"
             "•  Cimplicity (SCADA / HMI)\n"
             "•  Echelon · SLAM · Switching Orders"),
        15: "DEMAND SOURCES",
        16: ("•  Engineering — design new builds / modifications\n"
             "•  Planning — forecasts and capacity planning\n"
             "•  Grid operators — real-time monitoring & switching\n"
             "•  Field crews — work execution\n"
             "•  Vendor-driven upgrades (ProjectWise · Bentley)"),
        17: "Routed via Jira / IT Now / PDD",
        20: "TEAM",
        21: ("•  Combined Tech Lead — to be confirmed (George + Bruno)\n"
             "•  Engineers spanning GIS + G&E disciplines\n"
             "•  Squad size · ~13 FTE (combined GIS + G&E)\n"
             "•  Single PM for combined scope"),
        22: "Trade-off: simpler PM structure but mixes very different solution lifecycles (CAD vs SCADA)",
        25: "VENDOR LANDSCAPE",
        26: ("•  Bentley (ProjectWise / Microstation / CYME)\n"
             "•  Esri / ArcFM partner\n"
             "•  GE (Cimplicity)\n"
             "•  Echelon vendor\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = owner · vendor mix per sub-domain",
    },
    {  # VS 2 — Corporate & Regulatory (kept as-is)
        0: "DEVELOPMENT VALUE STREAM",
        1: "CORPORATE & REGULATORY — REGULATORY COMPLIANCE · EMPLOYEE ENABLEMENT",
        3: "ITBA · Sarath Garikapati + Amit (lead)",
        5: "Same scope in Proposal 2",
        8: ("Single value stream covering regulatory and employee-facing technology — operating under "
            "continuous compliance pressure with always-priority cadence."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  OpenText (content / records management)\n"
             "•  Rate Case Advisor (regulatory rate cases)\n"
             "•  Utilisphere (utility-specific platform)\n"
             "•  Corporate applications portfolio\n"
             "•  Employee-facing portals · regulatory reporting tools"),
        15: "DEMAND SOURCES",
        16: ("•  HR / Finance / Legal / Compliance — corporate process changes\n"
             "•  Regulators — mandated reporting cadence\n"
             "•  Employee population — enablement & training\n"
             "•  Continuous new requirements"),
        17: "Routed via Jira / IT Now / PDD",
        20: "TEAM",
        21: ("•  Tech Lead — Sarath Garikapati / Amit\n"
             "•  5–6 PMs + 2 tech leads (current composition)\n"
             "•  Squad size · ~8 FTE today\n"
             "•  Single PM for combined Corporate + Regulatory scope"),
        22: "Largest sub-team — internal sub-streams (HR vs Regulatory) may emerge",
        25: "VENDOR LANDSCAPE",
        26: ("•  OpenText vendor\n"
             "•  Accenture (corporate apps + projects)\n"
             "•  Regulatory tooling vendors\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = owner · vendor = delivery muscle",
    },
    {  # VS 3 — Iberdrola Corporate Services (kept as-is)
        0: "DEVELOPMENT VALUE STREAM",
        1: "IBERDROLA CORPORATE SERVICES — SPAIN-SIDE STABILITY · IBERDROLA ALIGNMENT",
        3: "ITBA · Vanessa Colombo + Michael Dubois (lead)",
        5: "Same scope in Proposal 2",
        8: ("Single value stream covering Spain-side SAP modules and Iberdrola integration — "
            "balancing incidents, small demands, and major projects."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  SAP FI (Finance) — Spain side\n"
             "•  SAP PO (Purchase Orders) — Spain side\n"
             "•  SAP Budget control\n"
             "•  S/4HANA (Iberdrola roadmap)\n"
             "•  Cross-org transport coordination"),
        15: "DEMAND SOURCES",
        16: ("•  Iberdrola corporate (Spain) — process changes\n"
             "•  Avangrid Finance — US-side integration needs\n"
             "•  Incidents (IT Now)\n"
             "•  Small Demands · Major projects (PDD)"),
        17: "Routed via Jira / IT Now / PDD",
        20: "TEAM",
        21: ("•  Tech Lead — Vanessa Colombo / Michael Dubois\n"
             "•  Business Analyst — SAP FI/PO SME\n"
             "•  Squad size · ~5 FTE today\n"
             "•  Single PM for Iberdrola scope"),
        22: "Hiring lag (~1 year) + knowledge-transfer loss flagged as top risks",
        25: "VENDOR LANDSCAPE",
        26: ("•  Accenture (Global SAP support)\n"
             "•  SAP partner for Spain modules\n"
             "•  Iberdrola Spain integration teams (peer)\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = owner · vendor = delivery muscle",
    },
]


# ---------------- DVS detail content (Prop 2 — 4 DVS) ----------------

PROP2_DVS_DETAIL = [
    {  # DVS 1 — Engineering & Asset Management
        0: "DEVELOPMENT VALUE STREAM (DVS) 1 / 4",
        1: "ENGINEERING & ASSET MANAGEMENT — RELIABLE ASSET DATA FOR PLANNERS AND ENGINEERS",
        3: "ITBA · George Porto (lead) + crossover with G&E team for ProjectWise/Bentley",
        5: "Run · Grow · Transform mix · to validate with team",
        8: ("Operate and evolve the engineering & asset platform — capture, model and publish "
            "reliable data so engineers and planners can design and operate the network."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  ArcFM (GIS / asset management)\n"
             "•  ProjectWise (engineering content / drawings)\n"
             "•  Bentley · Microstation (CAD)\n"
             "•  CYME (power systems analysis)\n"
             "•  SAP HANA spatial · GIS web apps"),
        15: "DEMAND SOURCES — who sends the work",
        16: ("•  Engineers — design new builds / modifications\n"
             "•  Planners — forecasts and capacity planning\n"
             "•  Asset Management — lifecycle decisions\n"
             "•  Field operations — asset updates\n"
             "•  Vendor-driven upgrades (ProjectWise · Bentley)"),
        17: "Routed via the Jira request type for this DVS",
        20: "AVANGRID CORE TEAM",
        21: ("•  Product Manager (NEW)\n"
             "•  Tech Lead — George Porto\n"
             "•  Crossover engineers from current G&E team for Bentley/ProjectWise scope\n"
             "•  Squad size · to validate (currently ~6 FTE in GIS team)"),
        22: "Cross-team scope — clean PM accountability requires crossing today's GIS/G&E line",
        25: "VENDOR EXTENSION (delivery)",
        26: ("•  Bentley (ProjectWise / Microstation / CYME)\n"
             "•  Esri / ArcFM partner\n"
             "•  SAP HANA partner\n"
             "•  SWF (Software Factory) — execution muscle"),
        27: "Avangrid = product owner · vendor = delivery muscle",
    },
    {  # DVS 2 — Network Operations & Control
        0: "DEVELOPMENT VALUE STREAM (DVS) 2 / 4",
        1: "NETWORK OPERATIONS & CONTROL — KEEP THE GRID UP · PREVENT AND MINIMIZE OUTAGES",
        3: "ITBA · Bruno Moreto + Andrew McGibbon (lead)",
        5: "Run · Grow · Transform mix · to validate with team",
        8: ("Operate and evolve the real-time grid control platform — SCADA, switching workflow, "
            "load analysis. Safety-critical."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  Cimplicity (SCADA / HMI)\n"
             "•  Echelon (control / metering platform)\n"
             "•  SLAM (System Loading and Analysis Module — to confirm)\n"
             "•  Switching Orders (utility-specific switching workflow)"),
        15: "DEMAND SOURCES — who sends the work",
        16: ("•  Grid operators — real-time monitoring needs\n"
             "•  Network engineers — analysis & load planning\n"
             "•  Field crews — switching execution\n"
             "•  Regulators — outage reporting cadence"),
        17: "Real-time / safety-critical · high-availability constraints apply",
        20: "AVANGRID CORE TEAM",
        21: ("•  Product Manager (NEW)\n"
             "•  Tech Lead — Bruno Moreto / Andrew McGibbon\n"
             "•  Business Analyst — grid operations SME\n"
             "•  Squad size · to validate (currently ~7 FTE in G&E team)"),
        22: "OT/IT boundary to confirm — some apps may sit in OT scope, not ITBA",
        25: "VENDOR EXTENSION (delivery)",
        26: ("•  Echelon vendor\n"
             "•  GE (Cimplicity)\n"
             "•  SLAM vendor (to confirm)\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = product owner · vendor = delivery muscle",
    },
    {  # DVS 3 — Corporate & Regulatory
        0: "DEVELOPMENT VALUE STREAM (DVS) 3 / 4",
        1: "CORPORATE & REGULATORY — REGULATORY COMPLIANCE · EMPLOYEE ENABLEMENT",
        3: "ITBA · Sarath Garikapati + Amit (lead)",
        5: "Run · Grow · Transform mix · to validate with team",
        8: ("Deliver and maintain regulatory and employee-facing technology across Avangrid US — "
            "operating under continuous compliance pressure."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  OpenText (content / records management)\n"
             "•  Rate Case Advisor (regulatory rate cases)\n"
             "•  Utilisphere (utility-specific platform)\n"
             "•  Corporate applications portfolio\n"
             "•  Employee-facing portals · regulatory reporting tools"),
        15: "DEMAND SOURCES — who sends the work",
        16: ("•  HR / Finance / Legal / Compliance — corporate process changes\n"
             "•  Regulators — mandated reporting cadence\n"
             "•  Employee population — enablement & training\n"
             "•  Continuous new requirements"),
        17: "Routed via the Jira request type for this DVS",
        20: "AVANGRID CORE TEAM",
        21: ("•  Product Manager (NEW)\n"
             "•  Tech Lead — Sarath Garikapati / Amit\n"
             "•  5–6 PMs + 2 tech leads (current composition)\n"
             "•  Squad size · to validate (~8 FTE today)"),
        22: "Largest sub-team — possible internal sub-streams (HR vs Regulatory) to validate",
        25: "VENDOR EXTENSION (delivery)",
        26: ("•  OpenText vendor\n"
             "•  Accenture (corporate apps + projects)\n"
             "•  Regulatory tooling vendors\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = product owner · vendor = delivery muscle",
    },
    {  # DVS 4 — Iberdrola Corporate Services
        0: "DEVELOPMENT VALUE STREAM (DVS) 4 / 4",
        1: "IBERDROLA CORPORATE SERVICES — SPAIN-SIDE STABILITY · IBERDROLA ALIGNMENT",
        3: "ITBA · Vanessa Colombo + Michael Dubois (lead)",
        5: "Run · Grow · Transform mix · to validate with team",
        8: ("Run Spain-side SAP modules and keep US/Iberdrola corporate processes aligned — "
            "balancing incidents, small demands, and major projects under hiring constraints."),
        11: "APPLICATIONS SUPPORTED",
        12: ("•  SAP FI (Finance) — Spain side\n"
             "•  SAP PO (Purchase Orders) — Spain side\n"
             "•  SAP Budget control\n"
             "•  S/4HANA (Iberdrola roadmap)\n"
             "•  Cross-org transport coordination"),
        15: "DEMAND SOURCES — who sends the work",
        16: ("•  Iberdrola corporate (Spain) — process changes\n"
             "•  Avangrid Finance — US-side integration needs\n"
             "•  Incidents (IT Now)\n"
             "•  Small Demands & Major projects (PDD)"),
        17: "Routed via the Jira request type for this DVS",
        20: "AVANGRID CORE TEAM",
        21: ("•  Product Manager (NEW)\n"
             "•  Tech Lead — Vanessa Colombo / Michael Dubois\n"
             "•  Business Analyst — SAP FI/PO SME\n"
             "•  Squad size · to validate (~5 FTE today)"),
        22: "Hiring lag (~1 year) + knowledge-transfer loss flagged as top risks",
        25: "VENDOR EXTENSION (delivery)",
        26: ("•  Accenture (Global SAP support)\n"
             "•  SAP partner for Spain modules\n"
             "•  Iberdrola Spain integration teams (peer)\n"
             "•  SWF (Software Factory)"),
        27: "Avangrid = product owner · vendor = delivery muscle",
    },
]


# ---------------- main ----------------

def main():
    prs = Presentation(SRC)

    # Drop only the Value Lens slide (idx 14) — keep all other slides (modify content)
    delete_slide_safe(prs, 14)
    # Now 14 slides — same structure as CS+VD without Value Lens

    # Original CS+VD remaining structure (0-indexed):
    # 0=Cover, 1=Prop1Cards, 2-4=3 VS detail Prop1, 5=Prop2Cards, 6=Prop2Rationale,
    # 7=Framing, 8=OVSdetail, 9-12=4 DVS detail Prop2, 13=Enabling

    # Step 1: Cover
    cover = prs.slides[0]
    for sh in cover.shapes:
        if sh.has_text_frame and "PROPOSED VALUE STREAMS" in sh.text_frame.text and "CUSTOMER SERVICES" in sh.text_frame.text:
            replace_textbox(sh,
                "PROPOSED VALUE STREAMS – IT BUSINESS APPLICATIONS\n"
                "Proposal 1: 3 Value Streams (consolidated)  ·  Proposal 2: 4 Development Value Streams\n"
                "April 2026")

    # Step 2: Rebuild Slide 1 — PROPOSAL 1 cards (3 VS)
    p1_cards = prs.slides[1]
    clear_slide(p1_cards)
    copy_bg(prs, 1, p1_cards)
    build_cards_slide(p1_cards,
        proposal_label="PROPOSAL 1 – 3 CONSOLIDATED VALUE STREAMS",
        ovs_title="IT Business Applications",
        ovs_tag="ITBA UMBRELLA · 3 VS",
        ovs_domain="DOMAIN · Network & Asset · Corporate & Regulatory · Iberdrola Corporate Services",
        ovs_desc=("Operate and evolve the technology platforms that support Avangrid's enterprise "
                  "operations — consolidating Engineering & Network Ops into a single value stream"),
        banner_text="ITBA  ·  3 CONSOLIDATED VALUE STREAMS",
        dvs_list=[
            ("Network & Asset Operations",
             "DOMAIN · Asset · Network",
             "Engineering, GIS, grid operations — combined",
             "ArcFM · ProjectWise · Bentley · CYME · Echelon · Cimplicity · SLAM · Switching Orders"),
            ("Corporate & Regulatory",
             "DOMAIN · Finance · HR · Procurement · Regulatory",
             "HR / employee / regulatory user self-service",
             "OpenText · Rate Case Advisor · Utilisphere"),
            ("Iberdrola Corporate Services",
             "DOMAIN · Finance · Procurement (Iberdrola)",
             "Spain finance / procurement / budget",
             "SAP FI · PO · Budget · S/4HANA"),
        ],
        n_dvs=3)

    # Step 3: Edit slides 2-4 (3 VS detail Prop 1)
    edit_by_idx(prs.slides[2], PROP1_VS_DETAIL[0])
    edit_by_idx(prs.slides[3], PROP1_VS_DETAIL[1])
    edit_by_idx(prs.slides[4], PROP1_VS_DETAIL[2])

    # Step 4: Rebuild Slide 5 — PROPOSAL 2 cards (4 DVS)
    p2_cards = prs.slides[5]
    clear_slide(p2_cards)
    copy_bg(prs, 5, p2_cards)
    build_cards_slide(p2_cards,
        proposal_label="PROPOSAL 2 – 1 OPERATIONAL VS UMBRELLA + 4 DEVELOPMENT VALUE STREAMS",
        ovs_title="IT Business Applications",
        ovs_tag="ITBA UMBRELLA · 4 DVS",
        ovs_domain="DOMAIN · Asset · Network · Corporate · Iberdrola — multiple business operations",
        ovs_desc=("Operate and evolve the technology platforms that support Avangrid's enterprise "
                  "operations — separating Engineering & Asset from Network Ops for clearer accountability"),
        banner_text="ITBA  ·  DEVELOPMENT VALUE STREAMS (4)",
        dvs_list=[
            ("Engineering & Asset Management",
             "DOMAIN · Asset Management",
             "Engineer / planner designs with reliable data",
             "ArcFM · ProjectWise · Bentley · CYME"),
            ("Network Operations & Control",
             "DOMAIN · Network Operations",
             "Grid operator prevents and minimizes outages",
             "Echelon · Cimplicity · SLAM · Switching Orders"),
            ("Corporate & Regulatory",
             "DOMAIN · Finance · HR · Procurement · Regulatory",
             "HR / employee / regulatory user self-service",
             "OpenText · Rate Case Advisor · Utilisphere"),
            ("Iberdrola Corporate Services",
             "DOMAIN · Finance · Procurement (Iberdrola)",
             "Spain finance / procurement / budget",
             "SAP FI · PO · Budget · S/4HANA"),
        ],
        n_dvs=4)

    # Step 5: Rebuild Slide 6 — PROPOSAL 2 RATIONALE (Big Picture)
    p2_rat = prs.slides[6]
    clear_slide(p2_rat)
    copy_bg(prs, 6, p2_rat)
    build_big_picture(p2_rat)

    # Step 6: Edit slide 7 (VALUE STREAM FRAMING)
    edit_by_idx(prs.slides[7], {
        0: "VALUE STREAM FRAMING",
        1: "ONE OPERATIONAL VS UMBRELLA + 4 DEVELOPMENT VALUE STREAMS  ·  SHARED ENABLING TEAM",
        3: "Director ITBA · Carlos Garbin  ·  4 DVS Product Managers (NEW)",
        5: "ITBA umbrella OVS · 4 DVS aligned to current sub-teams",
        8: ("ITBA serves multiple business operational value streams (one per domain). For governance "
            "purposes we wrap them under a single ITBA Operations OVS umbrella. Each ITBA Development "
            "Value Stream owns a coherent solution family aligned to one business domain."),
        11: "OPERATIONAL VALUE STREAM (UMBRELLA)",
        12: ("•  IT Business Applications Operations\n"
             "    – Owner: Director ITBA (Carlos Garbin)\n"
             "    – Customer-facing: enterprise business operations across 4 domains\n"
             "    – Forum: monthly portfolio review across the 4 DVS"),
        15: "DEVELOPMENT VALUE STREAMS (4)",
        16: ("•  DVS 1 — Engineering & Asset Management  ·  George Porto + crossover\n"
             "•  DVS 2 — Network Operations & Control  ·  Bruno Moreto + Andrew McGibbon\n"
             "•  DVS 3 — Corporate & Regulatory  ·  Sarath Garikapati + Amit\n"
             "•  DVS 4 — Iberdrola Corporate Services  ·  Vanessa Colombo + Michael Dubois\n"
             "+ Enabling Team — Foundation services (SWF · intake · test · risk)"),
        17: "Sizes to validate post-design",
        20: "MAPPING — DVS → APPLICATIONS",
        21: ("•  Engineering & Asset Mgmt → ArcFM · ProjectWise · Bentley · CYME\n"
             "•  Network Operations & Control → Echelon · Cimplicity · SLAM · Switching Orders\n"
             "•  Corporate & Regulatory → OpenText · Rate Case Advisor · Utilisphere\n"
             "•  Iberdrola Corporate Services → SAP FI · PO · Budget · S/4HANA\n"
             "•  Enabling Team → SWF · intake · test/release · third-party risk · ALM"),
        22: "1 OVS umbrella : 4 DVS — clean accountability · enabling team is shared support",
        25: "WHY THIS WORKS",
        26: ("•  4 dedicated Product Managers (one per DVS)\n"
             "•  Foundation services run as small enabling team — no extra PM\n"
             "•  Each DVS keeps current FTE base (5–8) — feasible squad size\n"
             "•  Director ITBA (Carlos) = existing role · no new layer\n"
             "•  Vendor extension model preserved per DVS"),
        27: "Same FTE base · clearer outcome ownership · single shared Enabling Team",
    })

    # Step 7: Edit slide 8 (OVS detail) — ITBA umbrella OVS
    edit_by_idx(prs.slides[8], {
        0: "OPERATIONAL VALUE STREAM (OVS)",
        1: "IT BUSINESS APPLICATIONS OPERATIONS — UMBRELLA OVS · 4 DVS · MULTIPLE BUSINESS DOMAINS",
        3: "ITBA · Carlos Garbin · OVS owner cross-DVS",
        5: "ITBA umbrella · supports multiple business OVSs across domains",
        8: ("Single ITBA umbrella operational value stream that wraps the 4 development value streams. "
            "Provides governance and portfolio view across the 4 domains while allowing each DVS to focus "
            "on its solution family."),
        11: "STAGES OF THE OVS",
        12: ("•  Capture demand from business stakeholders (PDD · Jira · IT Now)\n"
             "•  Route to the appropriate DVS by domain\n"
             "•  Plan and prioritize per DVS backlog\n"
             "•  Build & evolve solutions in DVS squads\n"
             "•  Coordinate cross-DVS dependencies (governed by Director ITBA)\n"
             "•  Deliver to business · operate · feedback loop"),
        15: "WHY THIS IS THE OPERATIONAL VS",
        16: ("•  Single umbrella for governance — one Director, one portfolio view\n"
             "•  Allows differentiated DVS strategies per business domain\n"
             "•  Enables Enabling Team to serve all 4 DVS uniformly\n"
             "•  Aligns with current org structure (Carlos Garbin = existing role)\n"
             "•  Provides escalation path for cross-DVS issues"),
        17: "OVS = governance & portfolio layer · DVS = solution-family delivery",
        20: "OVS GOVERNANCE (cross-DVS)",
        21: ("•  OVS Owner — Carlos Garbin (Director ITBA)\n"
             "•  Portfolio Forum — monthly cross-DVS prioritisation\n"
             "•  Architecture coherence — cross-DVS standards (SWF · ALM · risk)\n"
             "•  Flow metrics — lead time, throughput, business outcome KPIs (per DVS)\n"
             "•  Cross-DVS dependency resolution"),
        22: "Forum participants: 4 DVS Product Managers + Enabling Team coordinator",
        25: "OVS-LEVEL DEPENDENCIES",
        26: ("•  Cross-DVS resource sharing — SWF capacity contention\n"
             "•  Vendor concentration (Accenture, SWF, Bentley, SAP, Esri)\n"
             "•  Third-party risk gating — applies to all DVS\n"
             "•  Iberdrola integration cadence — driven by Spain side\n"
             "•  Compliance calendar — drives Corporate & Regulatory DVS"),
        27: "OVS = governance umbrella · DVS slides 10–13 detail the 4 development value streams",
    })

    # Step 8: Edit slides 9-12 (4 DVS detail Prop 2)
    edit_by_idx(prs.slides[9], PROP2_DVS_DETAIL[0])
    edit_by_idx(prs.slides[10], PROP2_DVS_DETAIL[1])
    edit_by_idx(prs.slides[11], PROP2_DVS_DETAIL[2])
    edit_by_idx(prs.slides[12], PROP2_DVS_DETAIL[3])

    # Step 9: Edit slide 13 (Enabling Team)
    edit_by_idx(prs.slides[13], {
        0:  "ENABLING TEAM",
        1:  "FOUNDATION SERVICES — SHARED CAPABILITIES THAT KEEP THE 4 ITBA DVS RUNNING",
        3:  "ITBA · governed by Director ITBA (Carlos Garbin) + cross-DVS coordination",
        5:  "No PM · no DVS-level backlog · operates as a small shared-service team",
        8:  ("Provide the technical and process foundation — Software Factory governance, demand "
             "intake, test & release coordination, third-party risk gating — that the 4 ITBA DVS rely on."),
        11: "FOUNDATION SERVICES",
        12: ("•  SWF (Software Factory) governance — used across all 4 DVS\n"
             "•  Demand intake orchestration (PDD · Jira · IT Now)\n"
             "•  Test & release coordination (cross-DVS)\n"
             "•  Third-party risk assessment gating\n"
             "•  Vendor-management cadence (Accenture · SWF · business vendors)\n"
             "•  ALM tooling administration"),
        15: "WHO USES THESE SERVICES",
        16: ("•  All 4 ITBA DVS — for SWF capacity · environments · test cycles\n"
             "•  Engineering & Asset Mgmt DVS — third-party risk gating for vendor upgrades\n"
             "•  Iberdrola Corp Services DVS — transport coordination · Iberdrola integration\n"
             "•  Director ITBA — portfolio view across DVS\n"
             "•  Optionally shared with CS Apps enabling team (Digital Products consolidation)"),
        17: "Foundation services support all 4 ITBA DVS — they do not constitute a separate product line",
        20: "TEAM SHAPE",
        21: ("•  Coordination led by Director ITBA (Carlos Garbin)\n"
             "•  SWF governance lead\n"
             "•  Demand intake / portfolio admin\n"
             "•  Test & risk coordinators — embedded in DVS where feasible\n"
             "•  No dedicated Product Manager\n"
             "Squad size · to validate"),
        22: "Why not a DVS: at this headcount, productizing foundation services would create overhead without internal customer roadmap",
        25: "VENDOR EXTENSION",
        26: ("•  SWF (Software Factory) vendor — capacity & tooling\n"
             "•  Accenture — for cross-DVS major projects\n"
             "•  Atlassian Jira admins\n"
             "•  Third-party risk assessment vendors"),
        27: "Avangrid governs the team operationally · vendors execute specific capabilities",
    })

    prs.save(DST)
    print(f"Saved → {DST}")
    print(f"Total slides: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        title = ""
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title = sh.text_frame.text.strip().split('\n')[0][:80]
                break
        print(f"  {i:2d}. {title}")


if __name__ == "__main__":
    main()
