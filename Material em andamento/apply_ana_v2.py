"""
Integrate Ana's feedback by REPLACING existing paragraph text (not appending).
This avoids text-box overflow that broke the previous version.
Targets: slides 12, 13, 15, 33, 39, 40, 41, 42, 44 + CONCLUSIONS shapes on 12/39.
Skips slide 34 (partner's domain).
"""
from pptx import Presentation
from copy import deepcopy
from pathlib import Path
from lxml import etree

OUT = Path("/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento")
SRC = OUT / "2604_AS-IS deliverable vrevisada - Final 2.pptx"
DST = OUT / "2604_AS-IS deliverable vrevisada - Final 3 (Ana integrated).pptx"

a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": a_ns, "p": p_ns}
prs = Presentation(str(SRC))


def shape_full_text(sp_xml):
    return "".join((t.text or "") for t in sp_xml.findall(".//a:t", NS))


def find_shape(slide, contains, occurrence=0):
    matches = [sp for sp in slide.element.findall(".//p:sp", NS) if contains in shape_full_text(sp)]
    return matches[occurrence] if occurrence < len(matches) else None


def replace_paragraph_text(para_elem, new_text):
    """Replace the visible text of a paragraph, preserving formatting of its first run."""
    runs = para_elem.findall("a:r", NS)
    if not runs:
        # No runs — likely just <a:t> in <a:fld> or similar; try direct
        ts = para_elem.findall(".//a:t", NS)
        if ts:
            ts[0].text = new_text
            for t in ts[1:]:
                t.text = ""
        return
    # Set first run's text to new_text, blank out the rest
    first_t = runs[0].find("a:t", NS)
    if first_t is None:
        # Create one
        t_elem = etree.SubElement(runs[0], f"{{{a_ns}}}t")
        t_elem.text = new_text
    else:
        first_t.text = new_text
    # Blank subsequent runs (so layout stays stable)
    for r in runs[1:]:
        for t in r.findall("a:t", NS):
            t.text = ""
    # Remove explicit line breaks <a:br>
    for br in para_elem.findall("a:br", NS):
        para_elem.remove(br)


def replace_para_in_shape(slide, shape_substring, target_substring, new_text, occurrence=0):
    """In the matching shape, find the paragraph containing target_substring and replace its text."""
    sh = find_shape(slide, shape_substring, occurrence)
    if sh is None:
        print(f"  ✗ shape with '{shape_substring[:50]}' not found")
        return False
    paras = sh.findall(".//a:p", NS)
    for pa in paras:
        ptext = "".join((t.text or "") for t in pa.findall(".//a:t", NS))
        if target_substring in ptext:
            replace_paragraph_text(pa, new_text)
            return True
    print(f"  ✗ paragraph containing '{target_substring[:50]}' not found in shape")
    return False


# ════════════════════════════════════════════════════════════════════════
# SLIDE 13 + 40 (Current Operating Constraints) — replace 2 PEOPLE bullets
# ════════════════════════════════════════════════════════════════════════
for idx in [12, 39]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — Operating Constraints")
    # PEOPLE bullets are in the 5-paragraph shape with "Limited and slow-to-scale capacity"
    # Replace bullet "Unclear ownership and overlapping..." with Run/Build content
    ok1 = replace_para_in_shape(
        s, "Limited and slow-to-scale capacity",
        "Unclear ownership and overlapping",
        "Run and Build are blended — same people fight incidents and build features, with no Run-team / Build-team separation"
    )
    print(f"   PEOPLE bullet 1 (Run/Build): {'OK' if ok1 else 'FAIL'}")
    # Replace another bullet — the "High vendor dependency" one — with missing biz roles
    ok2 = replace_para_in_shape(
        s, "Limited and slow-to-scale capacity",
        "High vendor dependency",
        "Business-side roles missing — no Process Owners, IT Business Partners (Finance · Supply Chain · Asset Mgmt) or Demand Manager"
    )
    print(f"   PEOPLE bullet 2 (biz roles): {'OK' if ok2 else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 12 + 39 (Future State Guiding Principles divider) — fold into existing bullet
# ════════════════════════════════════════════════════════════════════════
for idx in [11, 38]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — Future State Guiding Principles divider")
    # OPERATING CONSTRAINTS panel: existing People bullet
    ok = replace_para_in_shape(
        s, "These are the structural and operational limitations",
        "People: Capacity & accountability are structurally constrained",
        "People: Capacity & accountability constrained · Run/Build blended · business-side roles missing (Process Owners, IT Business Partners, Demand Manager)"
    )
    print(f"   OPERATING CONSTRAINTS bullet: {'OK' if ok else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 33 — Forums "What's Missing" — replace existing text with merged version
# ════════════════════════════════════════════════════════════════════════
s = prs.slides[32]
print(f"\nSlide 33 — Forums What's Missing")
ok = replace_para_in_shape(
    s, "No portfolio steering forum",
    "No portfolio steering forum",
    "No portfolio steering forum (decisions on trade-offs across branches, budget, capacity) • No ERP Steering Committee for SAP strategic direction and Architecture vision • Manager Check-in participants and scope pending definition"
)
print(f"   Missing forums bullet: {'OK' if ok else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 41 — App-Owner Mindset description — replace
# ════════════════════════════════════════════════════════════════════════
s = prs.slides[40]
print(f"\nSlide 41 — App-Owner Mindset")
ok = replace_para_in_shape(
    s, "demand is owned at the application level",
    "demand is owned at the application level",
    "Across both branches, demand is owned at the application level — not by Product Managers/POs accountable for end-to-end outcomes. The supporting roles that anchor a product mindset are also absent: no IT Business Partners by domain (Finance · Supply Chain · Asset Mgmt), no Process Owners, no Demand Manager."
)
print(f"   App-Owner Mindset description: {'OK' if ok else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 42 — CS Leadership Gap description — replace
# ════════════════════════════════════════════════════════════════════════
s = prs.slides[41]
print(f"\nSlide 42 — CS Leadership Gap")
ok = replace_para_in_shape(
    s, "Key leadership roles remain unfilled",
    "Key leadership roles remain unfilled",
    "Key leadership roles remain unfilled, with prioritization cascading inconsistently across sub-teams. The Run side is also undefined — no formal L1/L2/L3 SAP Application Support, no Service Delivery Manager, no Incident or Problem Manager — so 'everyone does everything'."
)
print(f"   CS Leadership Gap description: {'OK' if ok else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 15 + 44 (4 Guiding Principles) — replace last bullet of Principle 04
# ════════════════════════════════════════════════════════════════════════
for idx in [14, 43]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — Principle 04 (Governance & Ownership)")
    ok = replace_para_in_shape(
        s, "Defined roles (Product Owner / Product Manager / Tech Lead)",
        "KPIs tied to product outcomes and business value",
        "KPIs tied to outcomes · domain-based organization (Finance · Procurement · Asset Mgmt · Supply Chain) · Run vs Build separated · internal Data & Architecture capability"
    )
    print(f"   Principle 04 last bullet: {'OK' if ok else 'FAIL'}")


# ════════════════════════════════════════════════════════════════════════
# SLIDE 12 + 39 (CONCLUSIONS) — replace first sentence
# ════════════════════════════════════════════════════════════════════════
for idx in [11, 38]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — Conclusions")
    ok = replace_para_in_shape(
        s, "Our assessment highlights opportunities",
        "Our assessment highlights opportunities",
        "Our assessment highlights opportunities to strengthen alignment, governance, and prioritization across six operating-model blocks — Governance & Strategy · Business Relationship · Build · Run · Data & Analytics · Architecture & Integration."
    )
    print(f"   Conclusions sentence: {'OK' if ok else 'FAIL'}")


prs.save(str(DST))
print(f"\n✓ Saved: {DST}")
