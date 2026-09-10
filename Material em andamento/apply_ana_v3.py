"""
Extend Process / Tools / Governance categories with Ana's gaps so the
'six operating-model blocks' claim is internally consistent. Paint red.
"""
from pptx import Presentation
from pathlib import Path
from lxml import etree

SRC = Path("/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento/"
           "2604_AS-IS deliverable vrevisada - Final 3 (Ana integrated).pptx")
a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": a_ns, "p": p_ns}
RED = "C00000"
prs = Presentation(str(SRC))


def shape_text(sp): return "".join((t.text or "") for t in sp.findall(".//a:t", NS))


def find_shape(slide, contains, occ=0):
    m = [sp for sp in slide.element.findall(".//p:sp", NS) if contains in shape_text(sp)]
    return m[occ] if occ < len(m) else None


def replace_para_text(para, new_text):
    runs = para.findall("a:r", NS)
    if not runs:
        ts = para.findall(".//a:t", NS)
        if ts:
            ts[0].text = new_text
        return
    first_t = runs[0].find("a:t", NS)
    if first_t is None:
        first_t = etree.SubElement(runs[0], f"{{{a_ns}}}t")
    first_t.text = new_text
    for r in runs[1:]:
        for t in r.findall("a:t", NS):
            t.text = ""
    for br in para.findall("a:br", NS):
        para.remove(br)


def color_run_red(run, hex_color=RED):
    rPr = run.find("a:rPr", NS)
    if rPr is None:
        rPr = etree.SubElement(run, f"{{{a_ns}}}rPr")
        run.insert(0, rPr)
    for sf in rPr.findall("a:solidFill", NS):
        rPr.remove(sf)
    sf = etree.Element(f"{{{a_ns}}}solidFill")
    sc = etree.SubElement(sf, f"{{{a_ns}}}srgbClr"); sc.set("val", hex_color)
    ln = rPr.find("a:ln", NS)
    if ln is not None: ln.addnext(sf)
    else: rPr.insert(0, sf)
    rPr.set("b", "1")


def replace_and_paint(slide, sh_sub, target_sub, new_text, label=""):
    sh = find_shape(slide, sh_sub)
    if sh is None:
        print(f"  ✗ [{label}] shape not found"); return False
    for pa in sh.findall(".//a:p", NS):
        ptxt = "".join((t.text or "") for t in pa.findall(".//a:t", NS))
        if target_sub in ptxt:
            replace_para_text(pa, new_text)
            for r in pa.findall("a:r", NS):
                color_run_red(r)
            print(f"  ✓ [{label}]")
            return True
    print(f"  ✗ [{label}] paragraph not found"); return False


# ────────── Slide 13 + 40 (Operating Constraints detail) ──────────
for idx in [12, 39]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — detail")
    # PROCESS shape (shape[1]) — replace one bullet
    replace_and_paint(s, "Fragmented intake and ad-hoc prioritization",
                      "Inconsistent ways of working across teams",
                      "Run and Build are blended · work organized by module silos, not by functional domains (Finance · Procurement · Asset Mgmt · Supply Chain)",
                      label="PROCESS · domain")
    # TOOLS shape (shape[6]) — replace one bullet
    replace_and_paint(s, "Multiple disconnected intake and tracking tools",
                      "Inconsistent and informal tool",
                      "No internal Data & Analytics capability — BW/SAC and Data Engineers fully external · no internal Architecture or Integration capability",
                      label="TOOLS · data/arch")
    # GOVERNANCE shape (shape[5]) — replace one bullet
    replace_and_paint(s, "Multiple forums without portfolio-level ownership",
                      "Limited cross-team alignment mechanisms",
                      "No ERP Steering Committee for SAP-strategic decisions · no SAP Enterprise Architect role to define what goes in SAP vs other systems",
                      label="GOVERNANCE · ERP")


# ────────── Slide 12 + 39 (Future State Guiding Principles divider) ──────────
for idx in [11, 38]:
    s = prs.slides[idx]
    print(f"\nSlide {idx+1} — divider")
    # OPERATING CONSTRAINTS panel description shape — extend Process / Tools / Governance bullets
    replace_and_paint(s, "These are the structural and operational limitations",
                      "Process: Lack of standardization drives inconsistency and inefficiency",
                      "Process: Standardization gap · Run/Build blended · module silos, not functional domains",
                      label="DIV PROCESS")
    replace_and_paint(s, "These are the structural and operational limitations",
                      "Tools: Fragmented ecosystem limits visibility and control",
                      "Tools: Fragmented ecosystem · no internal Data & Analytics or Architecture capability",
                      label="DIV TOOLS")
    replace_and_paint(s, "These are the structural and operational limitations",
                      "Governance: Decision-making is fragmented and not value-driven",
                      "Governance: Fragmented decisions · no ERP Steering Committee · no SAP Enterprise Architect",
                      label="DIV GOVERNANCE")


prs.save(str(SRC))
print(f"\n✓ Saved")
