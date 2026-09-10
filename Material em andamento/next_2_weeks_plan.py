"""
Two-slide deck:
  Slide 1 — Next 2 Weeks · Lifecycle Map Definition (timeline)
  Slide 2 — The 8 Guiding Principles (reference)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- Brand palette ----------
NAVY       = RGBColor(0x1F, 0x3A, 0x5F)
NAVY_DEEP  = RGBColor(0x15, 0x2A, 0x47)
AMBER      = RGBColor(0xE8, 0x8B, 0x2F)
TEAL       = RGBColor(0x2B, 0x6B, 0x77)
GREEN      = RGBColor(0x4C, 0x8C, 0x3F)
GREEN_DARK = RGBColor(0x3B, 0x7D, 0x23)
LIGHT_BG   = RGBColor(0xF4, 0xF5, 0xF7)
GREY_TEXT  = RGBColor(0x4B, 0x55, 0x63)
DARK_TEXT  = RGBColor(0x1A, 0x1F, 0x2C)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BORDER     = RGBColor(0xD8, 0xDD, 0xE3)

# ---------- Canvas ----------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_rect(slide, left, top, width, height, fill, line=None, line_w=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w) if line_w else Pt(0.5)
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, text, *,
             size=10, bold=False, color=DARK_TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Emu(36000)
    tf.margin_top = tf.margin_bottom = Emu(18000)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        if i == 0:
            run = p.add_run()
        else:
            para = tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


# ================================================================
# SLIDE 1 — Next Two Weeks · Lifecycle Map Definition (timeline)
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])

# HEADER
add_rect(slide, 0, 0, prs.slide_width, Inches(0.85), NAVY)
add_text(slide, Inches(0.4), Inches(0.12), Inches(12), Inches(0.4),
         "NEXT TWO WEEKS — LIFECYCLE MAP DEFINITION",
         size=20, bold=True, color=WHITE)
add_text(slide, Inches(0.4), Inches(0.48), Inches(12), Inches(0.3),
         "Working model · how we will co-create the TO-BE Digital Products Lifecycle Map with Avangrid",
         size=11, color=RGBColor(0xC7, 0xD3, 0xE4))

# OBJECTIVE BAND
add_rect(slide, Inches(0.4), Inches(1.0), Inches(12.55), Inches(0.55), LIGHT_BG, BORDER, 0.75)
add_rect(slide, Inches(0.4), Inches(1.0), Inches(0.08), Inches(0.55), AMBER)
add_text(slide, Inches(0.6), Inches(1.05), Inches(1.8), Inches(0.45),
         "OBJECTIVE", size=9, bold=True, color=AMBER, anchor=MSO_ANCHOR.MIDDLE)
add_text(slide, Inches(2.3), Inches(1.05), Inches(10.5), Inches(0.45),
         "Co-create the TO-BE Lifecycle Map — 5 stages (Intake → Classification → Prioritization → Portfolio → Delivery) × 4 dimensions "
         "(Channels · Roles · Tools · Governance) — following the 8 Guiding Principles (see slide 2) and the hybrid Agile / Waterfall delivery model.",
         size=10, color=DARK_TEXT, anchor=MSO_ANCHOR.MIDDLE)

# TIMELINE
TL_TOP   = Inches(1.75)
TL_LEFT  = Inches(0.4)
wk_h     = Inches(0.34)
wk_w     = Inches(6.225)

add_rect(slide, TL_LEFT, TL_TOP, wk_w, wk_h, NAVY_DEEP)
add_text(slide, TL_LEFT, TL_TOP, wk_w, wk_h,
         "WEEK 1  ·  INTAKE · CLASSIFICATION · HYBRID MODEL",
         size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(slide, TL_LEFT + wk_w + Inches(0.1), TL_TOP, wk_w, wk_h, NAVY_DEEP)
add_text(slide, TL_LEFT + wk_w + Inches(0.1), TL_TOP, wk_w, wk_h,
         "WEEK 2  ·  PRIORITIZATION · PORTFOLIO · DELIVERY",
         size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

day_top  = TL_TOP + wk_h + Inches(0.12)
card_w   = Inches(1.205)
card_gap = Inches(0.04)
card_h   = Inches(3.0)

WEEK1 = [
    ("MON", "Draft TO-BE Model",
     "Build TO-BE model following the 8 Guiding Principles (see slide 2)",
     "NTT internal",
     "NTT team",
     NAVY),
    ("TUE", "Heather Validation",
     "Continue TO-BE · weekly 1-hour validation meeting with Heather (fixed slot every Tuesday)",
     "NTT  +  Heather",
     "Heather (1h fixed)",
     TEAL),
    ("WED", "Champions Joint Session",
     "Validate & adjust TO-BE · surface risks, issues and blockers",
     "NTT-led workshop",
     "All 6 Champions",
     TEAL),
    ("THU", "Refine TO-BE",
     "Apply Champion feedback · prepare materials for Friday review",
     "NTT internal",
     "NTT team",
     AMBER),
    ("FRI", "Weekly Status Meeting",
     "Review Intake · Classification · Hybrid Model",
     "NTT-led",
     "Vitor · Kumar · Heather · Carlos",
     GREEN),
]

WEEK2 = [
    ("MON", "Draft TO-BE — Week 2 topics",
     "Build TO-BE model for Prioritization · Portfolio · Delivery",
     "NTT internal",
     "NTT team",
     NAVY),
    ("TUE", "Heather Validation",
     "Continue TO-BE · weekly 1-hour validation meeting with Heather (fixed slot every Tuesday)",
     "NTT  +  Heather",
     "Heather (1h fixed)",
     TEAL),
    ("WED", "Champions Joint Session",
     "Validate & adjust TO-BE · surface risks, issues and blockers",
     "NTT-led workshop",
     "All 6 Champions",
     TEAL),
    ("THU", "Refine TO-BE",
     "Apply Champion feedback · prepare materials for Friday review",
     "NTT internal",
     "NTT team",
     AMBER),
    ("FRI", "Weekly Status Meeting",
     "Review Prioritization · Portfolio · Delivery",
     "NTT-led",
     "Vitor · Kumar · Heather · Carlos",
     GREEN),
]


def draw_day_card(slide, x, day, title, what, how, who, accent):
    add_rect(slide, x, day_top, card_w, card_h, WHITE, BORDER, 0.75)
    add_rect(slide, x, day_top, card_w, Inches(0.05), accent)
    add_text(slide, x, day_top + Inches(0.08), card_w, Inches(0.26),
             day, size=10, bold=True, color=accent,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x + Inches(0.08), day_top + Inches(0.38), card_w - Inches(0.16), Inches(0.55),
             title, size=9, bold=True, color=DARK_TEXT)
    add_text(slide, x + Inches(0.08), day_top + Inches(0.95), card_w - Inches(0.16), Inches(0.75),
             what, size=7.5, color=GREY_TEXT)
    sep = slide.shapes.add_connector(1, x + Inches(0.15), day_top + Inches(1.8),
                                     x + card_w - Inches(0.15), day_top + Inches(1.8))
    sep.line.color.rgb = BORDER
    sep.line.width = Pt(0.5)
    add_text(slide, x + Inches(0.08), day_top + Inches(1.85), card_w - Inches(0.16), Inches(0.18),
             "LEAD", size=6.5, bold=True, color=AMBER)
    add_text(slide, x + Inches(0.08), day_top + Inches(2.03), card_w - Inches(0.16), Inches(0.3),
             how, size=7.5, color=DARK_TEXT)
    add_text(slide, x + Inches(0.08), day_top + Inches(2.38), card_w - Inches(0.16), Inches(0.18),
             "PARTICIPANTS", size=6.5, bold=True, color=AMBER)
    add_text(slide, x + Inches(0.08), day_top + Inches(2.56), card_w - Inches(0.16), Inches(0.42),
             who, size=7.5, color=DARK_TEXT)


x = TL_LEFT
for d in WEEK1:
    draw_day_card(slide, x, *d)
    x += card_w + card_gap

x = TL_LEFT + wk_w + Inches(0.1)
for d in WEEK2:
    draw_day_card(slide, x, *d)
    x += card_w + card_gap

# BOTTOM: 3 panels
BOT_TOP = Inches(5.45)
BOT_H   = Inches(1.85)
panel_w = Inches(4.12)
panel_gap = Inches(0.1)

p1_left = Inches(0.4)
add_rect(slide, p1_left, BOT_TOP, panel_w, BOT_H, LIGHT_BG, BORDER, 0.75)
add_rect(slide, p1_left, BOT_TOP, panel_w, Inches(0.32), TEAL)
add_text(slide, p1_left, BOT_TOP, panel_w, Inches(0.32),
         "  HOW  ·  WORKING FORMATS",
         size=9.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

how_rows = [
    ("NTT internal design", "Mon & Thu · build and refine TO-BE model using the 8 Guiding Principles"),
    ("Fixed weekly with Heather", "Every Tuesday · 1 hour · validation & steering against business vision"),
    ("Champion joint session", "Every Wednesday · NTT-led · co-validation with all 6 Champions · risk surfacing"),
    ("Weekly Status Meeting", "Every Friday · review with Vitor · Kumar · Heather · Carlos"),
]
y = BOT_TOP + Inches(0.4)
for k, v in how_rows:
    add_text(slide, p1_left + Inches(0.15), y, panel_w - Inches(0.3), Inches(0.2),
             k, size=8.5, bold=True, color=TEAL)
    add_text(slide, p1_left + Inches(0.15), y + Inches(0.2), panel_w - Inches(0.3), Inches(0.18),
             v, size=7.5, color=GREY_TEXT)
    y += Inches(0.36)

p2_left = p1_left + panel_w + panel_gap
add_rect(slide, p2_left, BOT_TOP, panel_w, BOT_H, LIGHT_BG, BORDER, 0.75)
add_rect(slide, p2_left, BOT_TOP, panel_w, Inches(0.32), NAVY)
add_text(slide, p2_left, BOT_TOP, panel_w, Inches(0.32),
         "  WHO  ·  ENGAGEMENT MATRIX",
         size=9.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

who_rows = [
    ("Executive Sponsors", "Vitor · Heather  —  Friday status + Tue (Heather)"),
    ("Program Owners", "Carlos · Kumar  —  Friday status"),
    ("Area Champions (×6)", "GIS · Corporate · G&E · Global SAP · Smart Metering · SAP (CS)"),
    ("Heather (fixed)", "1 hour every Tuesday · business-vision validation"),
    ("NTT Team", "Lead + 2 consultants  —  Mon & Thu design · full-time"),
]
y = BOT_TOP + Inches(0.4)
for k, v in who_rows:
    add_text(slide, p2_left + Inches(0.15), y, panel_w - Inches(0.3), Inches(0.18),
             k, size=8.5, bold=True, color=NAVY)
    add_text(slide, p2_left + Inches(0.15), y + Inches(0.17), panel_w - Inches(0.3), Inches(0.18),
             v, size=7.5, color=GREY_TEXT)
    y += Inches(0.29)

p3_left = p2_left + panel_w + panel_gap
add_rect(slide, p3_left, BOT_TOP, panel_w, BOT_H, LIGHT_BG, BORDER, 0.75)
add_rect(slide, p3_left, BOT_TOP, panel_w, Inches(0.32), GREEN)
add_text(slide, p3_left, BOT_TOP, panel_w, Inches(0.32),
         "  WHAT WE DELIVER AT DAY 10",
         size=9.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

deliverables = [
    "Lifecycle Map v1.0  —  5 stages × 4 dimensions blueprint",
    "Hybrid Delivery Model  —  Agile / Waterfall routing rules",
    "End-to-End RACI  —  ownership & hand-off protocols",
    "Gap List  —  ready to operationalize vs. follow-ups",
    "Status-approved v1.0  —  ready to feed Wave 2",
]
y = BOT_TOP + Inches(0.42)
for txt in deliverables:
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                 p3_left + Inches(0.15), y + Inches(0.07),
                                 Inches(0.08), Inches(0.08))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GREEN
    dot.line.fill.background()
    add_text(slide, p3_left + Inches(0.3), y, panel_w - Inches(0.4), Inches(0.24),
             txt, size=8, color=DARK_TEXT)
    y += Inches(0.27)

add_text(slide, Inches(0.4), Inches(7.22), Inches(8), Inches(0.22),
         "© 2026 NTT DATA · Digital Products (Avangrid) · Wave 1 — Lifecycle Map Definition",
         size=7.5, color=GREY_TEXT)
add_text(slide, Inches(10.5), Inches(7.22), Inches(2.5), Inches(0.22),
         "Target close: end of Week 2",
         size=7.5, bold=True, color=AMBER, align=PP_ALIGN.RIGHT)


# ================================================================
# SLIDE 2 — The 8 Guiding Principles (reference)
# ================================================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])

# HEADER
add_rect(s2, 0, 0, prs.slide_width, Inches(0.85), NAVY)
add_text(s2, Inches(0.4), Inches(0.12), Inches(12), Inches(0.4),
         "EIGHT GUIDING PRINCIPLES — TO-BE MODEL REFERENCE",
         size=20, bold=True, color=WHITE)
add_text(s2, Inches(0.4), Inches(0.48), Inches(12), Inches(0.3),
         "The TO-BE Lifecycle Map is built against these 8 principles — each answers a structural issue of the AS-IS",
         size=11, color=RGBColor(0xC7, 0xD3, 0xE4))

# SUBTITLE
add_text(s2, Inches(0.4), Inches(1.05), Inches(12.5), Inches(0.45),
         "From this Baseline: BRANCH-UNIFIED (Apps + CS) and PRODUCT-ORIENTED (Value Streams, not silos).",
         size=11, bold=True, color=NAVY)

PRINCIPLES = [
    ("01", "Single Front Door",
     "One structured intake process for all demand types, replacing the 6+ disconnected channels."),
    ("02", "Value-Based Prioritization",
     "Standardized scoring framework across both branches — objective criteria, not ad-hoc forums."),
    ("03", "Capacity-Aware Assignment",
     "Assignment paired with capacity visibility — not manual handoffs with opaque availability."),
    ("04", "Standardized Execution Platforms",
     "Jira for Agile delivery (small / medium + evolutive); PPM Now for Waterfall delivery (large projects + CAPEX). Consistent tracking across both tracks."),
    ("05", "Integrated Toolchain",
     "Intake + Jira + PPM Now — three core platforms interoperable across Applications and Customer Services."),
    ("06", "Automated Reporting",
     "Real-time dashboards from a single source of truth — replacing manual data gathering."),
    ("07", "Defined RACI per Lifecycle Phase",
     "Unambiguous ownership and handoff protocols at every stage — from intake to deployment."),
    ("08", "Hybrid Delivery Model",
     "Delivery method matched to demand type: Agile (Jira / squads) for small–medium + evolutive work; Waterfall (PDD + PPM Now) for large, CAPEX or regulatory projects. Shared intake and governance."),
]

# Grid 2 cols × 4 rows
grid_top    = Inches(1.65)
col_w       = Inches(6.2)
col_gap     = Inches(0.15)
row_h       = Inches(1.33)
row_gap     = Inches(0.1)
col_left    = [Inches(0.4), Inches(0.4) + col_w + col_gap]

for i, (num, title, desc) in enumerate(PRINCIPLES):
    col = 0 if i < 4 else 1
    row = i if i < 4 else i - 4
    x = col_left[col]
    y = grid_top + row * (row_h + row_gap)

    # card background
    add_rect(s2, x, y, col_w, row_h, LIGHT_BG, BORDER, 0.75)
    # green accent top
    add_rect(s2, x, y, col_w, Inches(0.07), GREEN_DARK)
    # number
    add_text(s2, x + Inches(0.15), y + Inches(0.15), Inches(0.7), Inches(0.5),
             num, size=20, bold=True, color=GREEN_DARK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    # title
    add_text(s2, x + Inches(0.9), y + Inches(0.2), col_w - Inches(1.0), Inches(0.35),
             title, size=12, bold=True, color=GREEN_DARK, anchor=MSO_ANCHOR.MIDDLE)
    # description
    add_text(s2, x + Inches(0.9), y + Inches(0.55), col_w - Inches(1.05), Inches(0.75),
             desc, size=9, color=RGBColor(0x2C, 0x3E, 0x50))

# Footer
add_text(s2, Inches(0.4), Inches(7.22), Inches(8), Inches(0.22),
         "© 2026 NTT DATA · Digital Products (Avangrid) · Reference for TO-BE Lifecycle Map",
         size=7.5, color=GREY_TEXT)
add_text(s2, Inches(10.5), Inches(7.22), Inches(2.5), Inches(0.22),
         "Source: AS-IS Deliverable · Slide 47",
         size=7.5, bold=True, color=AMBER, align=PP_ALIGN.RIGHT)


# ---------- Save ----------
out = "/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento/Next_2_Weeks_Lifecycle_Map_Plan.pptx"
prs.save(out)
print(f"Saved: {out}")
