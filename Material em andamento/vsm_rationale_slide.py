"""
Builds a standalone slide replicating the VSM Demand Flow matrix,
but replaces the Intake→Delivery cell contents with short executive
WHY-messages explaining the maturity rating (Critical / Gap / Partial / Scalable).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- Brand palette ----------
NAVY       = RGBColor(0x2C, 0x5A, 0x88)
NAVY_DEEP  = RGBColor(0x21, 0x46, 0x6C)
TITLE_BLUE = RGBColor(0x2E, 0x6C, 0xB3)
LAVENDER   = RGBColor(0x95, 0x9A, 0xC8)   # row header bg (IT Business Apps group)
TEAL       = RGBColor(0x2B, 0x6B, 0x77)   # row header bg (Customer Services group)
CELL_BG    = RGBColor(0xF4, 0xF5, 0xF7)
BORDER     = RGBColor(0xD8, 0xDD, 0xE3)
DARK_TEXT  = RGBColor(0x24, 0x2A, 0x38)
GREY_TEXT  = RGBColor(0x3F, 0x48, 0x55)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
SECTION_LI = RGBColor(0x5B, 0xB5, 0xB8)   # teal separator line
FOOTER_BAR = RGBColor(0x2C, 0x5A, 0x88)

# Maturity colors
CRITICAL   = RGBColor(0xB6, 0x3B, 0x1F)   # dark red
GAP        = RGBColor(0xD9, 0x6E, 0x25)   # darker orange
PARTIAL    = RGBColor(0xEB, 0xA3, 0x35)   # yellow/orange
SCALABLE   = RGBColor(0x4C, 0x8C, 0x3F)   # green

STATUS = {"critical": CRITICAL, "gap": GAP, "partial": PARTIAL, "scalable": SCALABLE}

# ---------- Canvas ----------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def add_rect(left, top, width, height, fill, line=None, line_w=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w) if line_w else Pt(0.5)
    shp.shadow.inherit = False
    return shp


def add_text(left, top, width, height, text, *,
             size=10, bold=False, color=DARK_TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
             font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = Emu(36000); tf.margin_right = Emu(36000)
    tf.margin_top = Emu(14000);  tf.margin_bottom = Emu(14000)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_dot(left, top, color, diameter=Inches(0.14)):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    d.shadow.inherit = False
    return d


# ================ HEADER BAR ================
add_rect(0, 0, prs.slide_width, Inches(0.82), NAVY)
add_text(Inches(0.3), Inches(0.08), Inches(10), Inches(0.42),
         "CURRENT STATE ANALYSIS",
         size=20, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
add_text(Inches(0.3), Inches(0.48), Inches(10), Inches(0.28),
         "VALUE STREAM MAPPING — RATING RATIONALE",
         size=9, bold=True, color=RGBColor(0xC7, 0xD3, 0xE4), anchor=MSO_ANCHOR.MIDDLE)
# NTT DATA logo placeholder (right side)
add_text(Inches(11.4), Inches(0.2), Inches(1.6), Inches(0.45),
         "⦿ NTT DATA",
         size=13, bold=True, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# ================ TITLE ================
add_text(Inches(0.3), Inches(0.95), Inches(12.75), Inches(0.5),
         "Why each lifecycle stage was rated — the short executive rationale per team",
         size=14, bold=True, color=TITLE_BLUE)

# ================ COLUMN DEFINITIONS ================
TABLE_TOP   = Inches(1.55)
ROW_LABEL_W = Inches(1.25)
SOURCE_W    = Inches(1.60)
TYPES_W     = Inches(1.30)
LIFECYCLE_W = Inches(1.78)   # 5 cols × 1.78 ≈ 8.9 in
GAP_X       = Inches(0.04)

col_lefts   = [Inches(0.30)]
col_lefts.append(col_lefts[-1] + ROW_LABEL_W + GAP_X)     # SOURCE
col_lefts.append(col_lefts[-1] + SOURCE_W + GAP_X)        # TYPES
for _ in range(5):                                        # lifecycle cols
    col_lefts.append(col_lefts[-1] + (TYPES_W if len(col_lefts) == 3 else LIFECYCLE_W) + GAP_X)

col_widths = [ROW_LABEL_W, SOURCE_W, TYPES_W] + [LIFECYCLE_W] * 5

# ================ COLUMN HEADERS ================
HEADER_H = Inches(0.72)

headers = [
    ("", ""),  # row label col has no header
    ("SOURCE",         "Who sends the demand"),
    ("TYPES",          "What kind of demand"),
    ("01 INTAKE",      "Channels"),
    ("02 CLASSIFICATION", "Type, Size, Urgency"),
    ("03 PRIORITIZATION", "Scoring & Criteria"),
    ("04 PORTFOLIO",   "Governance"),
    ("05 DELIVERY",    "Execution"),
]

for i, (h1, h2) in enumerate(headers):
    if i == 0:
        continue
    x = col_lefts[i]; w = col_widths[i]
    add_rect(x, TABLE_TOP, w, HEADER_H, NAVY_DEEP)
    add_text(x, TABLE_TOP, w, Inches(0.42), h1,
             size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(x, TABLE_TOP + Inches(0.36), w, Inches(0.34), h2,
             size=8, color=RGBColor(0xC7, 0xD3, 0xE4), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================ SECTION HEADERS ================
SECTION_H = Inches(0.18)
section_y_app = TABLE_TOP + HEADER_H + Inches(0.06)

def add_section_header(y, label):
    # teal thin line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col_lefts[0], y + Inches(0.03),
                                  prs.slide_width - col_lefts[0] - Inches(0.3), Inches(0.03),
                                  )
    line.fill.solid(); line.fill.fore_color.rgb = SECTION_LI
    line.line.fill.background(); line.shadow.inherit = False
    add_text(col_lefts[0], y + Inches(0.08), Inches(5), SECTION_H,
             label, size=9, bold=True, color=SECTION_LI, anchor=MSO_ANCHOR.TOP)

add_section_header(section_y_app, "IT BUSINESS APPLICATIONS")

# ================ DATA ROWS ================

ROW_H = Inches(0.58)
rows_app = [
    {
        "area": "GIS",
        "group": "app",
        "source": "Electric & Gas Field Ops · OSG",
        "types": "Break/fix · Small · Large",
        "cells": [
            ("partial",  "No SLA, no volume tracking"),
            ("partial",  "No classification rule below PDD threshold"),
            ("partial",  "Heuristic-based; no scoring model"),
            ("partial",  "Informal portfolio management"),
            ("partial",  "Ad-hoc delivery; no release KPIs"),
        ],
    },
    {
        "area": "Corporate",
        "group": "app",
        "source": "Finance · HR · Procurement · Legal · Regulatory",
        "types": "Regulatory · Small/Medium",
        "cells": [
            ("partial",  "Requirements in docs; digital trail breaks"),
            ("gap",      "No model for small/medium demand"),
            ("partial",  "Regulatory = default priority; no scoring"),
            ("scalable", "Governed portfolio with formal approvals"),
            ("partial",  "One-year plan; no release KPIs"),
        ],
    },
    {
        "area": "G&E",
        "group": "app",
        "source": "E&G Ops · Field Services · Outage Mgmt · Regulatory",
        "types": "Version upgrades · Ad-hoc",
        "cells": [
            ("critical", "Single channel; no structured intake"),
            ("critical", "Depends on upgrade type; no rule"),
            ("critical", "Vendor urgency drives priority"),
            ("critical", "Vendor sets timeline & budget"),
            ("critical", "Vendor-owned delivery; 3rd-party risk slow"),
        ],
    },
    {
        "area": "Global SAP\n(Iberdrola)",
        "group": "app",
        "source": "Iberdrola Finance · Controlling · PO · Vendor",
        "types": "Incidents · Small · Minor · Large",
        "cells": [
            ("scalable", "Structured channel taxonomy in place"),
            ("scalable", "Consistent, channel-based classification"),
            ("critical", "Incidents first → strategic starved"),
            ("partial",  "Cadenced review; partial governance"),
            ("partial",  "KT loss on vendor turnover; 1-yr hiring lag"),
        ],
    },
]

rows_cs = [
    {
        "area": "Smart Metering\n(Device & Metering)",
        "group": "cs",
        "source": "Customer Ops · Field Services · Billing · BRS · Vendors",
        "types": "Incidents · Large · Small",
        "cells": [
            ("partial",  "Volumes not tracked"),
            ("scalable", "Consistent template-based classification"),
            ("partial",  "'Intrinsic common sense'; not scored"),
            ("scalable", "Only mature portfolio practice in branch"),
            ("critical", "100% vendor — execution knowledge not owned"),
        ],
    },
    {
        "area": "Customer Services\n(4 teams TBC)",
        "group": "cs",
        "source": "7 OPCOs via BRS · PMO · OT · CX · Digital CoE · Infra · D&A · Regulatory",
        "types": "Regulatory · Large · Small/Medium",
        "cells": [
            ("critical", "No shared channel; each sub-team has its own"),
            ("critical", "No branch model; 2 vacant Sr. Mgr seats"),
            ("gap",      "Director-level routing; no scoring"),
            ("critical", "Only one sub-team has portfolio governance"),
            ("partial",  "Ad-hoc delivery governance"),
        ],
    },
]


def draw_row(y, row):
    # Row label pill
    bg = LAVENDER if row["group"] == "app" else TEAL
    add_rect(col_lefts[0], y, ROW_LABEL_W, ROW_H, bg)
    add_text(col_lefts[0], y, ROW_LABEL_W, ROW_H,
             row["area"], size=9.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # SOURCE cell
    add_rect(col_lefts[1], y, SOURCE_W, ROW_H, CELL_BG, BORDER, 0.5)
    add_text(col_lefts[1] + Inches(0.06), y, SOURCE_W - Inches(0.12), ROW_H,
             row["source"], size=7.5, color=GREY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # TYPES cell
    add_rect(col_lefts[2], y, TYPES_W, ROW_H, CELL_BG, BORDER, 0.5)
    add_text(col_lefts[2] + Inches(0.06), y, TYPES_W - Inches(0.12), ROW_H,
             row["types"], size=7.5, color=GREY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # Lifecycle cells
    for i, (status, text) in enumerate(row["cells"]):
        x = col_lefts[3 + i]
        add_rect(x, y, LIFECYCLE_W, ROW_H, CELL_BG, BORDER, 0.5)
        # dot
        add_dot(x + Inches(0.08), y + Inches(0.12), STATUS[status])
        # text
        add_text(x + Inches(0.30), y + Inches(0.04),
                 LIFECYCLE_W - Inches(0.34), ROW_H - Inches(0.08),
                 text, size=7.5, color=DARK_TEXT, anchor=MSO_ANCHOR.MIDDLE)


y = section_y_app + SECTION_H + Inches(0.04)
for row in rows_app:
    draw_row(y, row)
    y += ROW_H + Inches(0.03)

# Section 2
add_section_header(y + Inches(0.02), "CUSTOMER SERVICES APPS")
y += SECTION_H + Inches(0.08)
for row in rows_cs:
    draw_row(y, row)
    y += ROW_H + Inches(0.03)

# ================ MATURITY LEGEND ================
LEGEND_Y = y + Inches(0.06)
add_rect(Inches(4.6), LEGEND_Y, Inches(4.5), Inches(0.28), NAVY_DEEP)
add_text(Inches(4.7), LEGEND_Y, Inches(1.0), Inches(0.28),
         "MATURITY:", size=8, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
legend_items = [
    ("CRITICAL", CRITICAL),
    ("GAP",      GAP),
    ("PARTIAL",  PARTIAL),
    ("SCALABLE", SCALABLE),
]
lx = Inches(5.75)
for label, color in legend_items:
    add_dot(lx, LEGEND_Y + Inches(0.07), color, diameter=Inches(0.14))
    add_text(lx + Inches(0.18), LEGEND_Y, Inches(0.85), Inches(0.28),
             label, size=8, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    lx += Inches(0.85)

# ================ FOOTER BANNER ================
FOOTER_Y = LEGEND_Y + Inches(0.4)
add_rect(Inches(0.3), FOOTER_Y, prs.slide_width - Inches(0.6), Inches(0.42), FOOTER_BAR)
add_text(Inches(0.3), FOOTER_Y, prs.slide_width - Inches(0.6), Inches(0.42),
         "Rationale: maturity is rated against 4 dimensions (Channels · Roles · Tools · Governance); the badge reflects the weakest dimension — not the average",
         size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Copyright
add_text(Inches(0.3), Inches(7.25), Inches(6), Inches(0.2),
         "© 2026 NTT DATA, Inc.", size=7, color=GREY_TEXT)

# ---------- Save ----------
out = "/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento/VSM_Rating_Rationale.pptx"
prs.save(out)
print(f"Saved: {out}")
