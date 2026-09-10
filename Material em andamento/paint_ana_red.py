"""
Paint all Ana-feedback text in RED on the AS-IS Final 3 deck.
Operates on the existing Final 3 (already integrated) and adds red color to modified paragraphs.
"""
from pptx import Presentation
from pathlib import Path
from lxml import etree

OUT = Path("/Users/gustavohenriquecastellano/Projeto-SAP-Digital-Products/Material em andamento")
SRC = OUT / "2604_AS-IS deliverable vrevisada - Final 3 (Ana integrated).pptx"

a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": a_ns, "p": p_ns}
RED_HEX = "C00000"  # strong red

prs = Presentation(str(SRC))


def shape_full_text(sp_xml):
    return "".join((t.text or "") for t in sp_xml.findall(".//a:t", NS))


def find_shape(slide, contains, occurrence=0):
    matches = [sp for sp in slide.element.findall(".//p:sp", NS) if contains in shape_full_text(sp)]
    return matches[occurrence] if occurrence < len(matches) else None


def color_run_red(run_elem, hex_color=RED_HEX):
    rPr = run_elem.find("a:rPr", NS)
    if rPr is None:
        rPr = etree.SubElement(run_elem, f"{{{a_ns}}}rPr")
        run_elem.insert(0, rPr)
    # Remove existing solidFill
    for sf in rPr.findall("a:solidFill", NS):
        rPr.remove(sf)
    sf = etree.SubElement(rPr, f"{{{a_ns}}}solidFill")
    sc = etree.SubElement(sf, f"{{{a_ns}}}srgbClr")
    sc.set("val", hex_color)
    # Also set bold so the red is unmistakable
    rPr.set("b", "1")


def paint_paragraph_red(para_elem):
    runs = para_elem.findall("a:r", NS)
    for r in runs:
        color_run_red(r)


def paint_para_in_shape(slide, shape_substring, target_substring, occurrence=0, label=""):
    sh = find_shape(slide, shape_substring, occurrence)
    if sh is None:
        print(f"  ✗ [{label}] shape with '{shape_substring[:50]}' NOT found")
        return False
    paras = sh.findall(".//a:p", NS)
    for pa in paras:
        ptext = "".join((t.text or "") for t in pa.findall(".//a:t", NS))
        if target_substring in ptext:
            paint_paragraph_red(pa)
            return True
    print(f"  ✗ [{label}] paragraph containing '{target_substring[:40]}' NOT found")
    return False


# Slide 13 + 40 — PEOPLE bullets
for idx in [12, 39]:
    s = prs.slides[idx]
    print(f"Slide {idx+1}")
    paint_para_in_shape(s, "Limited and slow-to-scale capacity",
                        "Run and Build are blended", label="Run/Build")
    paint_para_in_shape(s, "Limited and slow-to-scale capacity",
                        "Business-side roles missing", label="Biz roles")

# Slide 12 + 39 — OPERATING CONSTRAINTS People bullet (replaced)
for idx in [11, 38]:
    s = prs.slides[idx]
    print(f"Slide {idx+1}")
    paint_para_in_shape(s, "These are the structural and operational limitations",
                        "Run/Build blended", label="OC People")

# Slide 12 + 39 — Conclusions sentence
for idx in [11, 38]:
    s = prs.slides[idx]
    print(f"Slide {idx+1} Conclusions")
    paint_para_in_shape(s, "Our assessment highlights opportunities",
                        "six operating-model blocks", label="Conclusions")

# Slide 33 — Forums What's Missing
print("Slide 33 Forums")
paint_para_in_shape(prs.slides[32], "No portfolio steering forum",
                    "ERP Steering Committee", label="Forums missing")

# Slide 41 — App-Owner Mindset description
print("Slide 41 App-Owner")
paint_para_in_shape(prs.slides[40], "demand is owned at the application level",
                    "supporting roles that anchor a product mindset", label="App-Owner")

# Slide 42 — CS Leadership Gap description
print("Slide 42 CS Leadership")
paint_para_in_shape(prs.slides[41], "Key leadership roles remain unfilled",
                    "Run side is also undefined", label="CS Leadership")

# Slide 15 + 44 — Principle 04 last bullet
for idx in [14, 43]:
    s = prs.slides[idx]
    print(f"Slide {idx+1} Principle 04")
    paint_para_in_shape(s, "Defined roles (Product Owner / Product Manager / Tech Lead)",
                        "domain-based organization", label="P04 last bullet")

prs.save(str(SRC))
print(f"\n✓ Painted red and saved: {SRC.name}")
