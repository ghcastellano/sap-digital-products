"""Assemble TOM Boards & Intake deck from rendered PNGs."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

DECK_DIR = Path(__file__).parent
OUT = DECK_DIR / "TOM_Boards_Intake_Deck.pptx"

# 16:9 widescreen
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

blank = prs.slide_layouts[6]  # blank layout

png_files = sorted([p for p in DECK_DIR.glob("*.png") if p.name[0:2].isdigit()])
print(f"Assembling {len(png_files)} slides")

for png in png_files:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), 0, 0, width=SLIDE_W, height=SLIDE_H)
    print(f"  + {png.name}")

prs.save(str(OUT))
print(f"\nSaved → {OUT}")
print(f"Slides: {len(prs.slides)}")
