"""
Preparation & Planning · v2 (20 slides) — VISUAL redesign

Visual-first language: real diagrams (flows, hierarchies, swim lanes, hubs, decision trees)
instead of card-grids with bullet text.

INPUT  · TOM MAP workbook - final.pptx
OUTPUT · TOM MAP workbook - final v2.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import math, zipfile, shutil

SRC = "/Users/gustavohenriquecastellano/Downloads/Base Apresentações Projeto Digital Product/TOM MAP workbook - final.pptx"
DST = "/Users/gustavohenriquecastellano/Downloads/Base Apresentações Projeto Digital Product/TOM MAP workbook - final v2.pptx"

# ============== Palette ==============
HEADER_BLUE      = RGBColor(0x2C, 0x57, 0x8B)
HEADER_BLUE_DK   = RGBColor(0x1B, 0x3C, 0x66)
AVANGRID_GREEN   = RGBColor(0x00, 0x8C, 0x38)
GREEN_DARK       = RGBColor(0x00, 0x6B, 0x3F)
GREEN_MID        = RGBColor(0x00, 0x89, 0x50)
GREEN_LIGHT      = RGBColor(0x6F, 0xC1, 0x8F)
GREEN_BG         = RGBColor(0xF2, 0xF8, 0xF4)
GREEN_PASTEL     = RGBColor(0xC8, 0xE0, 0xCB)
BLUE_BG          = RGBColor(0xEE, 0xF3, 0xF9)
BLUE_PASTEL      = RGBColor(0xC9, 0xDA, 0xEF)
BLUE_LIGHT       = RGBColor(0x7E, 0x9C, 0xC4)
GREY_BG          = RGBColor(0xF4, 0xF5, 0xF7)
GREY_BORDER      = RGBColor(0xDC, 0xE0, 0xE5)
WHITE            = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT        = RGBColor(0x2C, 0x3E, 0x50)
GREY_LABEL       = RGBColor(0x6B, 0x77, 0x8C)
GREY_MUTED       = RGBColor(0x99, 0x99, 0x99)
ORANGE_DELIVERY  = RGBColor(0xE3, 0x7A, 0x1F)
ORANGE_BG        = RGBColor(0xFD, 0xF3, 0xE6)
PURPLE_HEATHER   = RGBColor(0x6B, 0x3F, 0x8E)
PURPLE_LIGHT     = RGBColor(0xA6, 0x84, 0xC0)
PURPLE_BG        = RGBColor(0xF2, 0xEB, 0xF6)
RED_GATE         = RGBColor(0xC5, 0x4B, 0x4B)
AMBER            = RGBColor(0xE0, 0x9F, 0x3E)

TITLE_FONT = "IberPangea Text Medium"
BODY_FONT  = "Arial"


# ============== Core helpers ==============
def _no_line(shape):
    shape.line.fill.background()

def rect(slide, x, y, w, h, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.shadow.inherit = False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s); return s

def round_rect(slide, x, y, w, h, fill=None, radius=0.04):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.shadow.inherit = False
    s.adjustments[0] = radius
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s); return s

def oval(slide, x, y, w, h, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    s.shadow.inherit = False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s); return s

def chevron(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, y, w, h)
    s.shadow.inherit = False
    s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s); return s

def pent(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, x, y, w, h)
    s.shadow.inherit = False
    s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s); return s

def line(slide, x1, y1, x2, y2, color=GREY_LABEL, weight=1.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln

def text(slide, x, y, w, h, txt, size=10, bold=False, italic=False,
         color=DARK_TEXT, font=BODY_FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = Emu(36000); tf.margin_right = Emu(36000)
    tf.margin_top = Emu(18000); tf.margin_bottom = Emu(18000)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = txt.split("\n") if isinstance(txt, str) else [str(txt)]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return tb

def slide_title(slide, kicker, subtitle):
    text(slide, Inches(0.40), Inches(0.10), Inches(10.5), Inches(0.42),
         kicker, size=22, bold=True, color=AVANGRID_GREEN, font=TITLE_FONT,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(0.40), Inches(0.50), Inches(10.5), Inches(0.32),
         subtitle, size=10, bold=True, color=DARK_TEXT, font=BODY_FONT)
    rect(slide, Inches(0.40), Inches(0.88), Inches(12.55), Emu(15000), AVANGRID_GREEN)

def phase_tracker(slide, current=2):
    base_y = 1.05
    phases = ["1 · Demand Mgmt", "2 · Prep & Planning", "3 · Execution"]
    pill_w, pill_h, gap = 1.55, 0.22, 0.04
    total_w = 3 * pill_w + 2 * gap
    start_x = 13.33 - 0.40 - total_w
    for i, label in enumerate(phases):
        x = start_x + i * (pill_w + gap)
        active = (i + 1) == current
        fill = AVANGRID_GREEN if active else GREY_BG
        col = WHITE if active else GREY_LABEL
        round_rect(slide, Inches(x), Inches(base_y), Inches(pill_w), Inches(pill_h),
                   fill, radius=0.30)
        text(slide, Inches(x), Inches(base_y), Inches(pill_w), Inches(pill_h),
             label, size=7, bold=True, color=col, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def footer_bar(slide, label):
    text(slide, Inches(0.40), Inches(7.18), Inches(8), Inches(0.18),
         "Avangrid · Digital Products · Target Operating Model",
         size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)
    text(slide, Inches(11.5), Inches(7.18), Inches(1.6), Inches(0.18),
         label, size=8, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.RIGHT)

SECTION_OF = {
    1: "FOUNDATIONS", 2: "FOUNDATIONS", 3: "FOUNDATIONS", 4: "FOUNDATIONS",
    5: "DIFFERENTIATION", 6: "DIFFERENTIATION",
    7: "PORTFOLIO + VS", 8: "PORTFOLIO + VS", 9: "PORTFOLIO + VS", 10: "PORTFOLIO + VS",
    11: "FLOWS", 12: "FLOWS", 13: "FLOWS", 14: "FLOWS",
    15: "DEEP DIVES", 16: "DEEP DIVES", 17: "DEEP DIVES",
    18: "CLOSURE", 19: "CLOSURE", 20: "CLOSURE",
}

NEXT_BRIDGE = {
    1:  "↓ NEXT · 02 · purpose, triggers, and outcomes",
    2:  "↓ NEXT · 03 · the end-to-end pipeline visualised",
    3:  "↓ NEXT · 04 · who governs the pipeline · two levels",
    4:  "↓ NEXT · 05 · CS and ITBA structurally compared",
    5:  "↓ NEXT · 06 · meet HEATHER · Product Leadership",
    6:  "↓ NEXT · 07 · Portfolio practices · CS",
    7:  "↓ NEXT · 08 · Portfolio practices · ITBA",
    8:  "↓ NEXT · 09 · zoom into the Value Stream layer · CS",
    9:  "↓ NEXT · 10 · Value Stream layer · ITBA",
    10: "↓ NEXT · 11 · path by size · S (Small)",
    11: "↓ NEXT · 12 · path by size · M (Medium)",
    12: "↓ NEXT · 13 · path by size · L (Large)",
    13: "↓ NEXT · 14 · scope-add shortcut for existing Initiatives",
    14: "↓ NEXT · 15 · refinement deep dive · the work breakdown tree",
    15: "↓ NEXT · 16 · pre-estimation + capacity allocation",
    16: "↓ NEXT · 17 · dual track · Discovery ↔ Delivery",
    17: "↓ NEXT · 18 · rolling-wave roadmap",
    18: "↓ NEXT · 19 · DoR · the gate to Execution",
    19: "↓ NEXT · 20 · the full roles cast",
    20: "↓ NEXT · Phase 3 · EXECUTION",
}

def breadcrumb(slide, n, section):
    text(slide, Inches(0.40), Inches(0.04), Inches(7.0), Inches(0.13),
         f"PREP & PLANNING  ·  {n:02d}/20  ·  {section}",
         size=7, bold=True, color=GREY_LABEL, font=BODY_FONT)

def story_bridge(slide, n):
    by = 7.00
    rect(slide, Inches(0.40), Inches(by), Inches(12.55), Emu(7000), AVANGRID_GREEN)
    text(slide, Inches(0.40), Inches(by + 0.02), Inches(12.55), Inches(0.16),
         NEXT_BRIDGE[n], size=8, italic=True, color=GREEN_DARK, font=BODY_FONT)

def setup(slide, n):
    phase_tracker(slide, current=2)
    breadcrumb(slide, n, SECTION_OF[n])

def avatar(slide, x, y, d, color, initials):
    oval(slide, Inches(x), Inches(y), Inches(d), Inches(d), color)
    text(slide, Inches(x), Inches(y), Inches(d), Inches(d), initials,
         size=int(d * 22), bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def chip(slide, x, y, w, label, color=GREY_BG, text_color=DARK_TEXT, h=0.30, bold=True, size=8):
    round_rect(slide, Inches(x), Inches(y), Inches(w), Inches(h), color, radius=0.30)
    text(slide, Inches(x), Inches(y), Inches(w), Inches(h), label,
         size=size, bold=bold, color=text_color, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================================
# 01 · DIVIDER — 3-phase timeline + 3 pillars
# ============================================================================
def build_pp01_divider(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "PREPARATION & PLANNING",
                "Phase 2 of the TOM  ·  the conversion from prioritized backlog to ready-to-execute work")
    setup(s, 1)

    tl_y = 1.55
    tl_w = 12.55
    tl_x = 0.40
    rect(s, Inches(tl_x), Inches(tl_y + 0.85), Inches(tl_w), Emu(20000), GREY_BORDER)
    phases = [
        ("1", "DEMAND\nMANAGEMENT", "intake → Gate 1 → Gate 2", GREY_MUTED, False),
        ("2", "PREPARATION\n& PLANNING", "refine → size → DoR", AVANGRID_GREEN, True),
        ("3", "EXECUTION", "build → test → release", GREY_MUTED, False),
    ]
    pw = (tl_w - 0.40) / 3
    for i, (num, name, sub, color, active) in enumerate(phases):
        x = tl_x + 0.10 + i * (pw + 0.10)
        d = 1.30 if active else 1.00
        cy = tl_y + 0.40 - d/2 + 0.45
        cx = x + pw/2 - d/2
        oval(s, Inches(cx), Inches(cy), Inches(d), Inches(d), color)
        text(s, Inches(cx), Inches(cy), Inches(d), Inches(d), num,
             size=int(d*30), bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(x), Inches(tl_y + 1.00), Inches(pw), Inches(0.40),
             name, size=12 if active else 11, bold=True,
             color=color if active else DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x), Inches(tl_y + 1.45), Inches(pw), Inches(0.22),
             sub, size=9, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
        if active:
            round_rect(s, Inches(x + pw/2 - 0.65), Inches(tl_y - 0.20),
                       Inches(1.30), Inches(0.22), AVANGRID_GREEN, radius=0.30)
            text(s, Inches(x + pw/2 - 0.65), Inches(tl_y - 0.20), Inches(1.30), Inches(0.22),
                 "▼ YOU ARE HERE", size=8, bold=True, color=WHITE, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    pl_y = 4.00
    pl_h = 2.90
    pl_w = (tl_w - 2 * 0.20) / 3
    pillars = [
        ("HK", PURPLE_HEATHER, "PRODUCT-LED",
         "Heather LaJoie\nDirector · Product Mgmt",
         "Product is NEW for Avangrid · Heather hires + leads new PMs · spans CS and ITBA"),
        ("CS", AVANGRID_GREEN, "CUSTOMER SERVICES",
         "Kumar Jaladanki\n+ middle managers",
         "4 CS DVS · Genesys / CIS / Smart Meter platforms · Digital CoE BRP"),
        ("IT", HEADER_BLUE, "IT BUSINESS APPS",
         "Carlos Garbin\n+ middle managers",
         "5 ITBA VS · SAP-heavy · regulatory + Safety · P&T BRP coupling"),
    ]
    for i, (initials, color, name, who, why) in enumerate(pillars):
        x = tl_x + i * (pl_w + 0.20)
        round_rect(s, Inches(x), Inches(pl_y), Inches(pl_w), Inches(pl_h), WHITE, radius=0.04)
        rect(s, Inches(x), Inches(pl_y), Inches(0.10), Inches(pl_h), color)
        avatar(s, x + pl_w/2 - 0.45, pl_y + 0.20, 0.90, color, initials)
        text(s, Inches(x + 0.20), Inches(pl_y + 1.30), Inches(pl_w - 0.30), Inches(0.30),
             name, size=11, bold=True, color=color, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x + 0.20), Inches(pl_y + 1.65), Inches(pl_w - 0.30), Inches(0.50),
             who, size=10, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x + 0.20), Inches(pl_y + 2.20), Inches(pl_w - 0.30), Inches(0.65),
             why, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    story_bridge(s, 1); footer_bar(s, "01/20")
    return s


# ============================================================================
# 02 · PURPOSE — funnel + WHEN + NON-GOALS
# ============================================================================
def build_pp02_purpose(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "PURPOSE  ·  WHEN  ·  OUTCOMES",
                "What this phase converts  ·  what flows in  ·  what flows out  ·  what we explicitly do not do")
    setup(s, 2)

    fx = 4.00; fy = 1.55; fw_top = 5.30; fw_bot = 2.80; fh = 4.80
    stages = [
        ("PRIORITIZED BACKLOG", "Demands · Pain Score · CoS · Type · OPCO", AVANGRID_GREEN),
        ("REFINEMENT", "Demand → Epic → Stories", GREEN_MID),
        ("PRE-ESTIMATION", "T-shirt + Effort + (Safety)", GREEN_MID),
        ("ARCHITECTURE", "NFRs + vendor + cross-VS deps", HEADER_BLUE),
        ("CAPACITY ALLOCATION", "% per CoS lane · per VS", HEADER_BLUE),
        ("DoR CHECK", "single gate to Execution", GREEN_DARK),
    ]
    n = len(stages)
    band_h = fh / n
    for i, (lbl, sub, color) in enumerate(stages):
        t1 = i / n; t2 = (i + 1) / n
        w1 = fw_top * (1 - t1) + fw_bot * t1
        w2 = fw_top * (1 - t2) + fw_bot * t2
        avg_w = (w1 + w2) / 2
        avg_x = fx + (fw_top - avg_w) / 2
        rect(s, Inches(avg_x), Inches(fy + i * band_h), Inches(avg_w), Inches(band_h - 0.04), color)
        text(s, Inches(avg_x), Inches(fy + i * band_h + 0.04), Inches(avg_w), Inches(0.30),
             lbl, size=11, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(avg_x), Inches(fy + i * band_h + 0.36), Inches(avg_w), Inches(band_h - 0.40),
             sub, size=9, italic=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    out_y = fy + fh + 0.10
    pent(s, Inches(fx + (fw_top - 3.20) / 2), Inches(out_y), Inches(3.20), Inches(0.50), HEADER_BLUE_DK)
    text(s, Inches(fx + (fw_top - 3.20) / 2), Inches(out_y), Inches(3.20), Inches(0.50),
         "READY FOR EXECUTION",
         size=11, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    when_y = 2.10
    round_rect(s, Inches(0.40), Inches(when_y), Inches(3.20), Inches(2.20), GREEN_BG, radius=0.04)
    text(s, Inches(0.55), Inches(when_y + 0.10), Inches(3.0), Inches(0.25),
         "WHEN  ·  TRIGGER", size=11, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(0.55), Inches(when_y + 0.40), Inches(3.0), Inches(0.30),
         "Gate 2 · GO decision",
         size=14, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(0.55), Inches(when_y + 0.78), Inches(3.0), Inches(1.30),
         "Monthly · 90 min\nVP Vitor chairs\nDirectors (Heather · Kumar · Carlos)\nPMs + Architect support",
         size=10, color=DARK_TEXT, font=BODY_FONT)
    line(s, Inches(3.65), Inches(when_y + 1.10), Inches(fx - 0.05), Inches(when_y + 1.10),
         AVANGRID_GREEN, weight=2.0)

    ng_y = 2.10; ng_x = 9.70
    round_rect(s, Inches(ng_x), Inches(ng_y), Inches(3.30), Inches(4.20), GREY_BG, radius=0.04)
    text(s, Inches(ng_x + 0.15), Inches(ng_y + 0.10), Inches(3.0), Inches(0.25),
         "NON-GOALS  ·  NOT US", size=11, bold=True, color=GREY_LABEL, font=BODY_FONT)
    ng_items = [
        ("✕", "Run our own BRP", "couple to P&T BRP + Digital CoE BRP"),
        ("✕", "Define OKRs", "Avangrid does not use them"),
        ("✕", "Run Discovery", "domain / vendor runs Phase 0/1"),
        ("✕", "Plan XL initiatives", "XL → PPM Now (per Heather)"),
    ]
    iy = ng_y + 0.45
    for mark, lbl, body in ng_items:
        oval(s, Inches(ng_x + 0.15), Inches(iy), Inches(0.28), Inches(0.28), GREY_LABEL)
        text(s, Inches(ng_x + 0.15), Inches(iy), Inches(0.28), Inches(0.28),
             mark, size=11, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(ng_x + 0.55), Inches(iy - 0.02), Inches(2.65), Inches(0.22),
             lbl, size=10, bold=True, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(ng_x + 0.55), Inches(iy + 0.20), Inches(2.65), Inches(0.32),
             body, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)
        iy += 0.92

    story_bridge(s, 2); footer_bar(s, "02/20")
    return s


# ============================================================================
# 03 · BIG-PICTURE FLOW · swim lanes
# ============================================================================
def build_pp03_big_picture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "BIG-PICTURE FLOW  ·  swim lanes by owner",
                "Same flow for CS and ITBA  ·  ownership shifts at each stage  ·  read left to right")
    setup(s, 3)

    lane_x = 1.85; lane_w = 11.10; lane_h = 1.40
    lane_y_start = 1.55; gap = 0.10

    lane_meta = [
        ("PORTFOLIO", "Director + middle mgmt", AVANGRID_GREEN, GREEN_BG),
        ("VALUE STREAM", "PM + Architect", HEADER_BLUE, BLUE_BG),
        ("SQUAD", "Squad PM", GREEN_DARK, GREEN_BG),
    ]
    for i, (name, who, color, bg) in enumerate(lane_meta):
        ly = lane_y_start + i * (lane_h + gap)
        round_rect(s, Inches(0.40), Inches(ly), Inches(1.40), Inches(lane_h), color, radius=0.04)
        text(s, Inches(0.40), Inches(ly + 0.12), Inches(1.40), Inches(0.25),
             name, size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(0.40), Inches(ly + 0.42), Inches(1.40), Inches(0.50),
             who, size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        rect(s, Inches(lane_x), Inches(ly), Inches(lane_w), Inches(lane_h), bg)

    stages = [
        (0, "GATE 2 GO", "monthly decision", AVANGRID_GREEN),
        (1, "REFINEMENT", "Epic + Stories", HEADER_BLUE),
        (1, "PRE-ESTIMATION", "T-shirt · Effort", HEADER_BLUE),
        (1, "ARCHITECTURE", "NFRs · vendor", HEADER_BLUE),
        (0, "CAPACITY ALLOC.", "% per CoS", AVANGRID_GREEN),
        (2, "DoR CHECK", "6 criteria gate", GREEN_DARK),
        (2, "READY · PULL", "into Execution", HEADER_BLUE_DK),
    ]
    n = len(stages)
    node_w = (lane_w - (n + 1) * 0.10) / n
    node_h = lane_h - 0.30
    for i, (lane_idx, name, sub, color) in enumerate(stages):
        nx = lane_x + 0.10 + i * (node_w + 0.10)
        ny = lane_y_start + lane_idx * (lane_h + gap) + 0.15
        round_rect(s, Inches(nx), Inches(ny), Inches(node_w), Inches(node_h), color, radius=0.06)
        text(s, Inches(nx), Inches(ny + 0.12), Inches(node_w), Inches(0.30),
             name, size=10, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(nx), Inches(ny + 0.42), Inches(node_w), Inches(0.30),
             sub, size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        if i < n - 1:
            next_lane = stages[i + 1][0]
            cur_x_end = nx + node_w
            next_x = lane_x + 0.10 + (i + 1) * (node_w + 0.10)
            cur_y_mid = ny + node_h / 2
            next_y_mid = lane_y_start + next_lane * (lane_h + gap) + 0.15 + node_h / 2
            line(s, Inches(cur_x_end), Inches(cur_y_mid),
                 Inches(next_x), Inches(next_y_mid), GREY_LABEL, weight=1.5)

    io_y = 6.10; io_h = 0.80
    iow = (12.55 - 0.20) / 2
    round_rect(s, Inches(0.40), Inches(io_y), Inches(iow), Inches(io_h), GREEN_BG, radius=0.04)
    text(s, Inches(0.55), Inches(io_y + 0.06), Inches(iow - 0.30), Inches(0.20),
         "INPUT  ·  from Demand Mgmt", size=9, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(0.55), Inches(io_y + 0.30), Inches(iow - 0.30), Inches(io_h - 0.35),
         "Prioritized Demand · Pain Score · CoS · Type · OPCO · Sponsor · Customer Impact Tier (CS)",
         size=9, color=DARK_TEXT, font=BODY_FONT)
    ox = 0.40 + iow + 0.20
    round_rect(s, Inches(ox), Inches(io_y), Inches(iow), Inches(io_h), BLUE_BG, radius=0.04)
    text(s, Inches(ox + 0.15), Inches(io_y + 0.06), Inches(iow - 0.30), Inches(0.20),
         "OUTPUT  ·  to Execution", size=9, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    text(s, Inches(ox + 0.15), Inches(io_y + 0.30), Inches(iow - 0.30), Inches(io_h - 0.35),
         "DoR-passed Stories · ≤ 5 SP · NFRs noted · capacity reserved · sponsor signed off",
         size=9, color=DARK_TEXT, font=BODY_FONT)

    story_bridge(s, 3); footer_bar(s, "03/20")
    return s


# ============================================================================
# 04 · GOVERNANCE · pyramid
# ============================================================================
def build_pp04_governance(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "GOVERNANCE  ·  TWO LEVELS",
                "Portfolio sits on top of Value Stream  ·  monthly decisions feed continuous work")
    setup(s, 4)

    cx_center = 6.65
    py = 1.55
    top_w = 7.50; top_h = 1.50
    top_x = cx_center - top_w / 2
    round_rect(s, Inches(top_x), Inches(py), Inches(top_w), Inches(top_h), AVANGRID_GREEN, radius=0.06)
    text(s, Inches(top_x), Inches(py + 0.10), Inches(top_w), Inches(0.32),
         "LEVEL 1  ·  PORTFOLIO", size=14, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(top_x), Inches(py + 0.45), Inches(top_w), Inches(0.25),
         "Director + middle managers · monthly Gate 2 · cross-VS view",
         size=9, italic=True, color=GREEN_PASTEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    av_y = py + 0.78
    avatars_top = [("HK", PURPLE_HEATHER), ("KJ", AVANGRID_GREEN), ("CG", HEADER_BLUE), ("MM", GREY_LABEL)]
    av_d = 0.55; av_gap = 0.40
    av_total = len(avatars_top) * av_d + (len(avatars_top) - 1) * av_gap
    av_start = cx_center - av_total / 2
    for i, (init, color) in enumerate(avatars_top):
        ax = av_start + i * (av_d + av_gap)
        oval(s, Inches(ax), Inches(av_y), Inches(av_d), Inches(av_d), color)
        text(s, Inches(ax), Inches(av_y), Inches(av_d), Inches(av_d), init,
             size=11, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    bot_w = 11.00; bot_h = 1.50
    bot_x = cx_center - bot_w / 2
    bot_y = py + top_h + 0.20
    round_rect(s, Inches(bot_x), Inches(bot_y), Inches(bot_w), Inches(bot_h), HEADER_BLUE, radius=0.06)
    text(s, Inches(bot_x), Inches(bot_y + 0.10), Inches(bot_w), Inches(0.32),
         "LEVEL 2  ·  VALUE STREAM", size=14, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(bot_x), Inches(bot_y + 0.45), Inches(bot_w), Inches(0.25),
         "PM (per VS · NEW under Heather) + Architect · continuous within VS",
         size=9, italic=True, color=BLUE_PASTEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    av_y2 = bot_y + 0.80
    n_pm = 9; pm_d = 0.40; pm_gap = 0.18
    pm_total = n_pm * pm_d + (n_pm - 1) * pm_gap
    pm_start = cx_center - pm_total / 2
    for i in range(n_pm):
        ax = pm_start + i * (pm_d + pm_gap)
        col = AVANGRID_GREEN if i < 4 else HEADER_BLUE
        oval(s, Inches(ax), Inches(av_y2), Inches(pm_d), Inches(pm_d), col)
        text(s, Inches(ax), Inches(av_y2), Inches(pm_d), Inches(pm_d), "PM",
             size=8, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    arr_x = cx_center
    line(s, Inches(arr_x - 0.50), Inches(py + top_h), Inches(arr_x - 0.50), Inches(bot_y),
         AVANGRID_GREEN, weight=2)
    text(s, Inches(arr_x - 1.50), Inches(py + top_h - 0.05), Inches(2.0), Inches(0.30),
         "decisions ↓", size=9, bold=True, color=AVANGRID_GREEN, font=BODY_FONT, align=PP_ALIGN.CENTER)
    line(s, Inches(arr_x + 0.50), Inches(bot_y), Inches(arr_x + 0.50), Inches(py + top_h),
         HEADER_BLUE, weight=2)
    text(s, Inches(arr_x + 0.40), Inches(py + top_h - 0.05), Inches(2.0), Inches(0.30),
         "↑ artefacts + asks", size=9, bold=True, color=HEADER_BLUE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    below_y = bot_y + bot_h + 0.30
    below_h = 1.95
    sw2 = (12.55 - 0.20) / 2
    round_rect(s, Inches(0.40), Inches(below_y), Inches(sw2), Inches(below_h), GREEN_BG, radius=0.04)
    text(s, Inches(0.55), Inches(below_y + 0.10), Inches(sw2 - 0.30), Inches(0.25),
         "LEVEL 1 · what gets decided", size=10, bold=True, color=GREEN_DARK, font=BODY_FONT)
    l1 = [("CAPACITY %", "per CoS lane · monthly"),
          ("BACKLOG RANK", "cross-VS · monthly"),
          ("CROSS-VS DEPS", "resolved at Gate 2"),
          ("LICENSES + VENDORS", "renewals · models")]
    iy = below_y + 0.40
    for lbl, sub in l1:
        chip(s, 0.55, iy, 1.95, lbl, AVANGRID_GREEN, WHITE, h=0.24, size=8)
        text(s, Inches(2.60), Inches(iy), Inches(sw2 - 2.25), Inches(0.24),
             sub, size=9, color=DARK_TEXT, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
        iy += 0.35

    cx2 = 0.40 + sw2 + 0.20
    round_rect(s, Inches(cx2), Inches(below_y), Inches(sw2), Inches(below_h), BLUE_BG, radius=0.04)
    text(s, Inches(cx2 + 0.15), Inches(below_y + 0.10), Inches(sw2 - 0.30), Inches(0.25),
         "LEVEL 2 · what gets produced", size=10, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    l2 = [("REFINED BACKLOG", "Epics + Stories per VS"),
          ("ROADMAP", "rolling wave 3·6·12 mo"),
          ("ARCHITECTURE NOTE", "NFRs · vendor model"),
          ("DoR-PASSED STORIES", "ready to pull")]
    iy = below_y + 0.40
    for lbl, sub in l2:
        chip(s, cx2 + 0.15, iy, 2.10, lbl, HEADER_BLUE, WHITE, h=0.24, size=8)
        text(s, Inches(cx2 + 2.40), Inches(iy), Inches(sw2 - 2.30), Inches(0.24),
             sub, size=9, color=DARK_TEXT, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
        iy += 0.35

    story_bridge(s, 4); footer_bar(s, "04/20")
    return s


# ============================================================================
# 05 · CS vs ITBA · side-by-side panels with avatars + chips + formulas
# ============================================================================
def build_pp05_cs_vs_itba(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "CS  vs  ITBA  ·  STRUCTURAL DIFFERENCES",
                "Same flow, different teams  ·  different formulas  ·  different vendor mix  ·  Heather bridges both")
    setup(s, 5)

    hb_y = 1.55
    round_rect(s, Inches(3.40), Inches(hb_y), Inches(6.55), Inches(0.50), PURPLE_HEATHER, radius=0.30)
    avatar(s, 3.55, hb_y + 0.05, 0.40, PURPLE_LIGHT, "HK")
    text(s, Inches(4.10), Inches(hb_y), Inches(5.85), Inches(0.50),
         "HEATHER LaJOIE  ·  Director · Product Mgmt  ·  spans CS and ITBA",
         size=11, bold=True, color=WHITE, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
    line(s, Inches(5.00), Inches(hb_y + 0.50), Inches(3.50), Inches(2.20), PURPLE_HEATHER, weight=1.5)
    line(s, Inches(8.30), Inches(hb_y + 0.50), Inches(9.85), Inches(2.20), PURPLE_HEATHER, weight=1.5)

    col_y = 2.25; col_h = 4.55
    cw = (12.55 - 0.20) / 2

    # CS panel
    round_rect(s, Inches(0.40), Inches(col_y), Inches(cw), Inches(col_h), WHITE, radius=0.04)
    rect(s, Inches(0.40), Inches(col_y), Inches(cw), Inches(0.70), AVANGRID_GREEN)
    avatar(s, 0.60, col_y + 0.10, 0.50, GREEN_DARK, "KJ")
    text(s, Inches(1.20), Inches(col_y + 0.10), Inches(cw - 1.0), Inches(0.30),
         "CS  ·  Customer Services", size=14, bold=True, color=WHITE, font=BODY_FONT)
    text(s, Inches(1.20), Inches(col_y + 0.40), Inches(cw - 1.0), Inches(0.25),
         "Kumar Jaladanki · Director · CS Engineering",
         size=9, italic=True, color=GREEN_PASTEL, font=BODY_FONT)

    cs_dvs = ["Customer Care", "Billing Ops", "Field Ops (CS)", "Smart Metering"]
    cy = col_y + 0.90
    text(s, Inches(0.55), Inches(cy), Inches(cw - 0.30), Inches(0.20),
         "4 VALUE STREAMS", size=8, bold=True, color=GREEN_DARK, font=BODY_FONT)
    cy += 0.22
    chip_w = (cw - 0.30 - 3 * 0.06) / 4
    for i, name in enumerate(cs_dvs):
        chip(s, 0.55 + i * (chip_w + 0.06), cy, chip_w, name, GREEN_BG, GREEN_DARK, h=0.30, size=8)

    cy += 0.50
    round_rect(s, Inches(0.55), Inches(cy), Inches(cw - 0.30), Inches(0.85), GREEN_BG, radius=0.04)
    text(s, Inches(0.65), Inches(cy + 0.06), Inches(cw - 0.50), Inches(0.22),
         "PAIN SCORE  ·  no Safety", size=8, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(0.65), Inches(cy + 0.30), Inches(cw - 0.50), Inches(0.50),
         "(BV·OPCOmult·0.40 + UX·0.60) / Effort",
         size=11, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)

    cy += 1.00
    text(s, Inches(0.55), Inches(cy), Inches(cw - 0.30), Inches(0.20),
         "VENDOR + NFR PROFILE", size=8, bold=True, color=GREEN_DARK, font=BODY_FONT)
    cy += 0.22
    cs_chips = [("Genesys", AVANGRID_GREEN), ("CIS", AVANGRID_GREEN), ("Smart Meter", AVANGRID_GREEN),
                ("CCPA", GREEN_MID), ("Latency", GREEN_MID), ("24/7 ops", GREEN_MID)]
    chip_x = 0.55
    for label, col in cs_chips:
        cw_chip = max(1.10, len(label) * 0.10)
        chip(s, chip_x, cy, cw_chip, label, col, WHITE, h=0.28, size=7)
        chip_x += cw_chip + 0.08
        if chip_x > 0.55 + cw - 1.20:
            chip_x = 0.55; cy += 0.36

    cy = col_y + col_h - 0.42
    round_rect(s, Inches(0.55), Inches(cy), Inches(cw - 0.30), Inches(0.32), GREEN_DARK, radius=0.04)
    text(s, Inches(0.55), Inches(cy), Inches(cw - 0.30), Inches(0.32),
         "BRP  ·  Digital CoE BRP  ·  Heather is liaison",
         size=9, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ITBA panel
    cx2 = 0.40 + cw + 0.20
    round_rect(s, Inches(cx2), Inches(col_y), Inches(cw), Inches(col_h), WHITE, radius=0.04)
    rect(s, Inches(cx2), Inches(col_y), Inches(cw), Inches(0.70), HEADER_BLUE)
    avatar(s, cx2 + 0.20, col_y + 0.10, 0.50, HEADER_BLUE_DK, "CG")
    text(s, Inches(cx2 + 0.80), Inches(col_y + 0.10), Inches(cw - 1.0), Inches(0.30),
         "ITBA  ·  IT Business Apps", size=14, bold=True, color=WHITE, font=BODY_FONT)
    text(s, Inches(cx2 + 0.80), Inches(col_y + 0.40), Inches(cw - 1.0), Inches(0.25),
         "Carlos Garbin · Director · ITBA",
         size=9, italic=True, color=BLUE_PASTEL, font=BODY_FONT)

    itba_vs = ["Corp & Reg", "Iberdrola Corp", "Renewables IT", "Networks IT", "Enabling Team"]
    cy = col_y + 0.90
    text(s, Inches(cx2 + 0.15), Inches(cy), Inches(cw - 0.30), Inches(0.20),
         "5 VALUE STREAMS", size=8, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    cy += 0.22
    chip_w = (cw - 0.30 - 4 * 0.06) / 5
    for i, name in enumerate(itba_vs):
        chip(s, cx2 + 0.15 + i * (chip_w + 0.06), cy, chip_w, name, BLUE_BG, HEADER_BLUE, h=0.30, size=7)

    cy += 0.50
    round_rect(s, Inches(cx2 + 0.15), Inches(cy), Inches(cw - 0.30), Inches(0.85), BLUE_BG, radius=0.04)
    text(s, Inches(cx2 + 0.25), Inches(cy + 0.06), Inches(cw - 0.50), Inches(0.22),
         "PAIN SCORE  ·  Safety isolated", size=8, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    text(s, Inches(cx2 + 0.25), Inches(cy + 0.30), Inches(cw - 0.50), Inches(0.50),
         "(BV·OPCOmult·0.45 + UX·0.30 + Safety·0.25) / Effort",
         size=10, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)

    cy += 1.00
    text(s, Inches(cx2 + 0.15), Inches(cy), Inches(cw - 0.30), Inches(0.20),
         "VENDOR + NFR PROFILE", size=8, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    cy += 0.22
    itba_chips = [("SAP modules", HEADER_BLUE), ("Vendor portals", HEADER_BLUE), ("Custom builds", HEADER_BLUE),
                  ("Regulatory", HEADER_BLUE_DK), ("Audit trail", HEADER_BLUE_DK), ("Safety", HEADER_BLUE_DK)]
    chip_x = cx2 + 0.15
    for label, col in itba_chips:
        cw_chip = max(1.10, len(label) * 0.10)
        chip(s, chip_x, cy, cw_chip, label, col, WHITE, h=0.28, size=7)
        chip_x += cw_chip + 0.08
        if chip_x > cx2 + cw - 1.20:
            chip_x = cx2 + 0.15; cy += 0.36

    cy = col_y + col_h - 0.42
    round_rect(s, Inches(cx2 + 0.15), Inches(cy), Inches(cw - 0.30), Inches(0.32), HEADER_BLUE_DK, radius=0.04)
    text(s, Inches(cx2 + 0.15), Inches(cy), Inches(cw - 0.30), Inches(0.32),
         "BRP  ·  P&T BRP  ·  Carlos + Asier own it",
         size=9, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    story_bridge(s, 5); footer_bar(s, "05/20")
    return s


# ============================================================================
# 06 · HEATHER · hub-and-spoke
# ============================================================================
def build_pp06_heather(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "HEATHER LaJOIE  ·  PRODUCT LEADERSHIP",
                "The hub of Product Management for Digital Products  ·  NEW capability for Avangrid")
    setup(s, 6)

    cx = 6.65; cy = 4.30; hub_d = 1.80
    oval(s, Inches(cx - hub_d/2), Inches(cy - hub_d/2), Inches(hub_d), Inches(hub_d), PURPLE_HEATHER)
    text(s, Inches(cx - hub_d/2), Inches(cy - hub_d/2), Inches(hub_d), Inches(hub_d),
         "HEATHER\nLaJOIE\nDirector\nProduct Mgmt",
         size=11, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    spokes = [
        (90,  "VP VITOR",       "co-chair Gate 2",            HEADER_BLUE_DK),
        (30,  "KUMAR · CS",     "co-Director",                AVANGRID_GREEN),
        (-30, "PMs · per VS",   "9 NEW under Heather",        PURPLE_LIGHT),
        (-90, "DIGITAL CoE BRP","Heather is the liaison",     AMBER),
        (-150,"P&T BRP",        "indirect (Carlos owns)",     HEADER_BLUE),
        (150, "CARLOS · ITBA",  "co-Director",                HEADER_BLUE),
    ]
    spoke_d = 0.90
    spoke_r = 2.85
    for ang_deg, label, sub, color in spokes:
        ang = math.radians(ang_deg)
        sx = cx + spoke_r * math.cos(ang) - spoke_d / 2
        sy = cy - spoke_r * math.sin(ang) - spoke_d / 2
        line(s, Inches(cx + (hub_d/2) * math.cos(ang)),
             Inches(cy - (hub_d/2) * math.sin(ang)),
             Inches(cx + (spoke_r - spoke_d/2) * math.cos(ang)),
             Inches(cy - (spoke_r - spoke_d/2) * math.sin(ang)),
             PURPLE_LIGHT, weight=2)
        oval(s, Inches(sx), Inches(sy), Inches(spoke_d), Inches(spoke_d), color)
        text(s, Inches(sx), Inches(sy + 0.18), Inches(spoke_d), Inches(0.30),
             label, size=8, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(sx - 0.40), Inches(sy + spoke_d + 0.04),
             Inches(spoke_d + 0.80), Inches(0.30),
             sub, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    own_x = 0.40; own_y = 1.55; own_w = 3.30; own_h = 5.30
    round_rect(s, Inches(own_x), Inches(own_y), Inches(own_w), Inches(own_h), WHITE, radius=0.04)
    rect(s, Inches(own_x), Inches(own_y), Inches(own_w), Inches(0.42), PURPLE_HEATHER)
    text(s, Inches(own_x + 0.15), Inches(own_y + 0.06), Inches(own_w - 0.30), Inches(0.30),
         "WHAT SHE OWNS", size=10, bold=True, color=WHITE, font=BODY_FONT)
    owns = [
        "Product Mgmt capability\n(NEW for Avangrid)",
        "Hires + leads new PMs\n(per VS)",
        "Liaison to Digital CoE BRP",
        "Co-Director at Gate 2",
        "Performance metrics\n(with VP Vitor)",
        "Outcomes — not deliverables",
    ]
    iy = own_y + 0.55
    rh = (own_h - 0.65) / len(owns)
    for o in owns:
        oval(s, Inches(own_x + 0.20), Inches(iy + 0.10), Inches(0.20), Inches(0.20), PURPLE_HEATHER)
        text(s, Inches(own_x + 0.50), Inches(iy + 0.05), Inches(own_w - 0.65), Inches(rh - 0.10),
             o, size=9, color=DARK_TEXT, font=BODY_FONT)
        iy += rh

    no_x = 9.65; no_y = 1.55; no_w = 3.30; no_h = 5.30
    round_rect(s, Inches(no_x), Inches(no_y), Inches(no_w), Inches(no_h), WHITE, radius=0.04)
    rect(s, Inches(no_x), Inches(no_y), Inches(no_w), Inches(0.42), GREY_LABEL)
    text(s, Inches(no_x + 0.15), Inches(no_y + 0.06), Inches(no_w - 0.30), Inches(0.30),
         "WHAT SHE DOES NOT OWN", size=10, bold=True, color=WHITE, font=BODY_FONT)
    nots = [
        "Day-to-day capacity\n→ middle managers",
        "CS Engineering delivery\n→ Kumar's team",
        "ITBA Engineering delivery\n→ Carlos's team",
        "Discovery (Phase 0/1)\n→ domain / vendor",
        "OKRs / Avangrid metrics\n→ none defined",
    ]
    iy = no_y + 0.55
    rh = (no_h - 0.65) / len(nots)
    for o in nots:
        oval(s, Inches(no_x + 0.20), Inches(iy + 0.10), Inches(0.20), Inches(0.20), GREY_LABEL)
        text(s, Inches(no_x + 0.50), Inches(iy + 0.05), Inches(no_w - 0.65), Inches(rh - 0.10),
             o, size=9, italic=True, color=GREY_LABEL, font=BODY_FONT)
        iy += rh

    story_bridge(s, 6); footer_bar(s, "06/20")
    return s


# ============================================================================
# 07/08 · PORTFOLIO · CS or ITBA — control-panel visual (clock + 6 satellites)
# ============================================================================
def build_pp_portfolio(prs, n, side):
    is_cs = side == "CS"
    primary = AVANGRID_GREEN if is_cs else HEADER_BLUE
    primary_dk = GREEN_DARK if is_cs else HEADER_BLUE_DK
    bg_light = GREEN_BG if is_cs else BLUE_BG
    director = "Kumar Jaladanki" if is_cs else "Carlos Garbin"
    director_init = "KJ" if is_cs else "CG"
    brp = "Digital CoE BRP" if is_cs else "P&T BRP"
    vs_label = "4 CS DVS" if is_cs else "5 ITBA VS"

    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, f"LEVEL 1  ·  PORTFOLIO  ·  {side}",
                f"Monthly Gate 2 cadence  ·  Director + Heather + middle managers  ·  cross-VS view  ·  {brp} coupling")
    setup(s, n)

    own_y = 1.55; own_h = 1.05
    round_rect(s, Inches(0.40), Inches(own_y), Inches(12.55), Inches(own_h), bg_light, radius=0.04)
    text(s, Inches(0.55), Inches(own_y + 0.08), Inches(12.30), Inches(0.20),
         "PORTFOLIO OWNERS  ·  3 voices · 1 decision",
         size=10, bold=True, color=primary_dk, font=BODY_FONT)
    avs = [
        (director_init, primary, director, "Director · Engineering"),
        ("HK", PURPLE_HEATHER, "Heather LaJoie", "Director · Product"),
        ("MM", GREY_LABEL, "Middle managers", "day-to-day capacity"),
    ]
    av_y = own_y + 0.36; av_w = 4.10
    for i, (init, color, name, role) in enumerate(avs):
        ax = 0.60 + i * (av_w + 0.10)
        oval(s, Inches(ax), Inches(av_y), Inches(0.55), Inches(0.55), color)
        text(s, Inches(ax), Inches(av_y), Inches(0.55), Inches(0.55), init,
             size=11, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(ax + 0.65), Inches(av_y + 0.02), Inches(av_w - 0.70), Inches(0.22),
             name, size=10, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(ax + 0.65), Inches(av_y + 0.24), Inches(av_w - 0.70), Inches(0.28),
             role, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)

    clk_x = 5.45; clk_y = 2.95; clk_d = 2.00
    oval(s, Inches(clk_x), Inches(clk_y), Inches(clk_d), Inches(clk_d), primary)
    oval(s, Inches(clk_x + 0.18), Inches(clk_y + 0.18), Inches(clk_d - 0.36), Inches(clk_d - 0.36), WHITE)
    text(s, Inches(clk_x), Inches(clk_y + 0.40), Inches(clk_d), Inches(0.30),
         "GATE 2", size=14, bold=True, color=primary, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(clk_x), Inches(clk_y + 0.74), Inches(clk_d), Inches(0.25),
         "MONTHLY", size=10, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(clk_x), Inches(clk_y + 1.00), Inches(clk_d), Inches(0.25),
         "90 min", size=9, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(clk_x), Inches(clk_y + 1.25), Inches(clk_d), Inches(0.25),
         "VP Vitor chairs", size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    sat_specs = [
        (0.40, 2.95, "CAPACITY %", "per CoS lane · monthly", primary),
        (0.40, 4.30, "BACKLOG RANK", "Pain Score + CoS + OPCO", primary),
        (0.40, 5.65, "CROSS-VS DEPS", "from BRP outputs", primary_dk),
        (9.85, 2.95, "LICENSES", "renewals + vendor mgmt", primary_dk),
        (9.85, 4.30, "METRICS", "Heather + Vitor (no OKRs)", primary_dk),
        (9.85, 5.65, "BRP COUPLING", brp, primary),
    ]
    for x, y, lbl, sub, color in sat_specs:
        round_rect(s, Inches(x), Inches(y), Inches(3.05), Inches(1.20), WHITE, radius=0.04)
        rect(s, Inches(x), Inches(y), Inches(0.10), Inches(1.20), color)
        text(s, Inches(x + 0.20), Inches(y + 0.10), Inches(2.85), Inches(0.30),
             lbl, size=11, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(x + 0.20), Inches(y + 0.42), Inches(2.85), Inches(0.30),
             sub, size=9, color=DARK_TEXT, font=BODY_FONT)
        line(s, Inches(x + 3.05) if x < 5 else Inches(x),
             Inches(y + 0.60),
             Inches(clk_x) if x < 5 else Inches(clk_x + clk_d),
             Inches(clk_y + clk_d / 2),
             color, weight=1.0)

    foot_y = 6.10
    round_rect(s, Inches(3.60), Inches(foot_y), Inches(6.15), Inches(0.70), bg_light, radius=0.04)
    text(s, Inches(3.60), Inches(foot_y + 0.05), Inches(6.15), Inches(0.20),
         f"FOOTPRINT  ·  {vs_label}",
         size=9, bold=True, color=primary_dk, font=BODY_FONT, align=PP_ALIGN.CENTER)
    if is_cs:
        names = ["Customer Care", "Billing Ops", "Field Ops (CS)", "Smart Metering"]
    else:
        names = ["Corp & Reg", "Iberdrola Corp", "Renewables IT", "Networks IT", "Enabling Team"]
    cw_n = (6.15 - 0.20 - (len(names) - 1) * 0.05) / len(names)
    for i, name in enumerate(names):
        chip(s, 3.70 + i * (cw_n + 0.05), foot_y + 0.30, cw_n, name, primary, WHITE, h=0.28,
             size=7 if not is_cs else 8)

    story_bridge(s, n); footer_bar(s, f"{n:02d}/20")
    return s

def build_pp07_portfolio_cs(prs):  return build_pp_portfolio(prs, 7, "CS")
def build_pp08_portfolio_itba(prs): return build_pp_portfolio(prs, 8, "ITBA")


# ============================================================================
# 09/10 · VS PRACTICES · matrix VS × practice
# ============================================================================
def build_pp_vs(prs, n, side):
    is_cs = side == "CS"
    primary = AVANGRID_GREEN if is_cs else HEADER_BLUE
    primary_dk = GREEN_DARK if is_cs else HEADER_BLUE_DK
    bg_light = GREEN_BG if is_cs else BLUE_BG
    if is_cs:
        vs_names = ["Customer Care", "Billing Ops", "Field Ops (CS)", "Smart Metering"]
    else:
        vs_names = ["Corp & Reg", "Iberdrola Corp", "Renewables IT", "Networks IT", "Enabling Team"]

    practices = [
        ("REFINEMENT", "Demand→Epic→Story"),
        ("PRE-EST.", "T-shirt + Effort"),
        ("ARCHITECTURE", "NFRs + vendor"),
        ("ROADMAP", "rolling wave"),
        ("DoR + HANDOFF", "single gate"),
    ]

    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, f"LEVEL 2  ·  VALUE STREAM  ·  {side}",
                f"PM (per VS · NEW under Heather) + Architect  ·  continuous within VS  ·  {len(vs_names)} streams × 5 practices")
    setup(s, n)

    mx = 0.40; my = 1.55
    cell_h = 0.95
    label_w = 1.85
    pcell_w = (12.55 - label_w) / len(practices)

    rect(s, Inches(mx), Inches(my), Inches(label_w), Inches(0.55), bg_light)
    text(s, Inches(mx), Inches(my), Inches(label_w), Inches(0.55), "VALUE STREAM",
         size=9, bold=True, color=primary_dk, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i, (lbl, sub) in enumerate(practices):
        x = mx + label_w + i * pcell_w
        rect(s, Inches(x), Inches(my), Inches(pcell_w - 0.04), Inches(0.55), primary)
        text(s, Inches(x), Inches(my + 0.04), Inches(pcell_w - 0.04), Inches(0.25),
             lbl, size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x), Inches(my + 0.30), Inches(pcell_w - 0.04), Inches(0.22),
             sub, size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    cell_data_cs = [
        ["NPS · CSAT framing", "no Safety in formula", "Genesys / CIS / SaaS", "rolling 3·6·12 mo", "+CIT P1/P2 noted"],
        ["cycle time framing", "OPCOmult often >1", "CIS · multi-OPCO", "rolling 3·6·12 mo", "+Sponsor signoff"],
        ["self-service framing", "Effort by Architect", "Mobile + GIS + APIs", "rolling 3·6·12 mo", "+UX validated"],
        ["meter-event framing", "vendor-led estim.", "Smart Meter platforms", "rolling 3·6·12 mo", "+vendor SLA noted"],
    ]
    cell_data_itba = [
        ["audit-pass framing", "Safety isolated 25%", "regulatory portals", "deadlines pin commits", "+regulatory flag"],
        ["cost-savings framing", "SAP module deps", "in-house SAP teams", "rolling 3·6·12 mo", "+SAP rollback"],
        ["uptime framing", "Effort + Safety", "vendor + custom", "rolling 3·6·12 mo", "+NFRs perf"],
        ["data-integrity framing", "Architect double-pass", "custom builds", "rolling 3·6·12 mo", "+audit trail"],
        ["enabling framing", "shared services sized", "mixed", "rolling 3·6·12 mo", "+enabling DoR"],
    ]
    cell_data = cell_data_cs if is_cs else cell_data_itba

    ry = my + 0.55
    for ri, vs_name in enumerate(vs_names):
        rect(s, Inches(mx), Inches(ry), Inches(label_w), Inches(cell_h - 0.04), WHITE)
        rect(s, Inches(mx), Inches(ry), Inches(0.10), Inches(cell_h - 0.04), primary)
        text(s, Inches(mx + 0.18), Inches(ry + 0.10), Inches(label_w - 0.25), Inches(cell_h - 0.20),
             vs_name, size=10, bold=True, color=primary_dk, font=BODY_FONT,
             anchor=MSO_ANCHOR.MIDDLE)
        for ci in range(len(practices)):
            x = mx + label_w + ci * pcell_w
            cell_color = WHITE if (ri + ci) % 2 == 0 else bg_light
            rect(s, Inches(x), Inches(ry), Inches(pcell_w - 0.04), Inches(cell_h - 0.04), cell_color)
            txt_val = cell_data[ri][ci] if ri < len(cell_data) and ci < len(cell_data[ri]) else ""
            text(s, Inches(x + 0.06), Inches(ry + 0.10), Inches(pcell_w - 0.16), Inches(cell_h - 0.20),
                 txt_val, size=8, color=DARK_TEXT, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        ry += cell_h

    leg_y = ry + 0.10
    leg_pills = [
        ("PM · per VS",  PURPLE_LIGHT, "outcome framing + roadmap"),
        ("ARCHITECT", primary, "NFRs + cross-VS deps"),
        ("SQUAD PM", primary_dk, "DoR + pull"),
        ("HEATHER", PURPLE_HEATHER, "mentors PMs · sets Product standard"),
    ]
    pw = (12.55 - 3 * 0.10) / 4
    for i, (lbl, color, sub) in enumerate(leg_pills):
        x = 0.40 + i * (pw + 0.10)
        round_rect(s, Inches(x), Inches(leg_y), Inches(pw), Inches(0.50), color, radius=0.04)
        text(s, Inches(x), Inches(leg_y + 0.04), Inches(pw), Inches(0.22),
             lbl, size=9, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x), Inches(leg_y + 0.26), Inches(pw), Inches(0.22),
             sub, size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    story_bridge(s, n); footer_bar(s, f"{n:02d}/20")
    return s

def build_pp09_vs_cs(prs):  return build_pp_vs(prs, 9, "CS")
def build_pp10_vs_itba(prs): return build_pp_vs(prs, 10, "ITBA")


# ============================================================================
# 11/12/13 · PATH BY SIZE · pipeline + CS/ITBA detail
# ============================================================================
def build_pp_path(prs, n, size, sp_range, color, weeks, branch_pattern):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, f"PATH BY SIZE  ·  {size}",
                f"{sp_range}  ·  {weeks}  ·  same logic for CS and ITBA · NFR profile differs")
    setup(s, n)

    bg_x = 0.40; bg_y = 1.55; bg_d = 1.50
    oval(s, Inches(bg_x), Inches(bg_y), Inches(bg_d), Inches(bg_d), color)
    text(s, Inches(bg_x), Inches(bg_y), Inches(bg_d), Inches(bg_d), size,
         size=64, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(bg_x), Inches(bg_y + bg_d + 0.08), Inches(bg_d), Inches(0.22),
         sp_range, size=9, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(bg_x), Inches(bg_y + bg_d + 0.30), Inches(bg_d), Inches(0.22),
         weeks, size=9, bold=True, color=color, font=BODY_FONT, align=PP_ALIGN.CENTER)

    flow_x = 2.10; flow_y = 1.85; flow_w = 10.85; flow_h = 0.95
    n_steps = len(branch_pattern)
    step_w = (flow_w - (n_steps - 1) * 0.15) / n_steps
    for i, (lbl, sub, kind) in enumerate(branch_pattern):
        x = flow_x + i * (step_w + 0.15)
        if kind == "fork":
            round_rect(s, Inches(x), Inches(flow_y), Inches(step_w), Inches(flow_h), AMBER, radius=0.06)
        elif kind == "gate":
            round_rect(s, Inches(x), Inches(flow_y), Inches(step_w), Inches(flow_h), GREEN_DARK, radius=0.06)
        else:
            round_rect(s, Inches(x), Inches(flow_y), Inches(step_w), Inches(flow_h), color, radius=0.06)
        text(s, Inches(x), Inches(flow_y + 0.10), Inches(step_w), Inches(0.32),
             lbl, size=10, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(x), Inches(flow_y + 0.46), Inches(step_w), Inches(0.45),
             sub, size=8, italic=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < n_steps - 1:
            text(s, Inches(x + step_w + 0.01), Inches(flow_y + 0.20), Inches(0.13), Inches(0.50),
                 "→", size=18, bold=True, color=GREY_LABEL, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    box_y = 3.20; box_h = 3.65
    sw2 = (12.55 - 0.20) / 2

    def _side_box(x, y, w, h, side_color, side_label, items):
        round_rect(s, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, radius=0.04)
        rect(s, Inches(x), Inches(y), Inches(w), Inches(0.42), side_color)
        text(s, Inches(x + 0.15), Inches(y + 0.06), Inches(w - 0.30), Inches(0.30),
             side_label, size=11, bold=True, color=WHITE, font=BODY_FONT)
        iy = y + 0.55
        rh = (h - 0.65) / len(items)
        for icon, lbl, body in items:
            round_rect(s, Inches(x + 0.20), Inches(iy + 0.10), Inches(0.40), Inches(0.40), side_color, radius=0.10)
            text(s, Inches(x + 0.20), Inches(iy + 0.10), Inches(0.40), Inches(0.40), icon,
                 size=12, bold=True, color=WHITE, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(s, Inches(x + 0.70), Inches(iy + 0.06), Inches(w - 0.85), Inches(0.20),
                 lbl, size=9, bold=True, color=DARK_TEXT, font=BODY_FONT)
            text(s, Inches(x + 0.70), Inches(iy + 0.26), Inches(w - 0.85), Inches(rh - 0.30),
                 body, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)
            iy += rh

    if size == "S":
        cs_items = [("►", "Examples", "CRM page-load · script update · self-service field"),
                    ("⚙", "Pain profile", "P3/P4 typical · CIT P1/P2 → Auto-Expedite"),
                    ("◯", "Owner cadence", "weekly batch · Squad PM owns DoR"),
                    ("✦", "Watch out", "don't over-batch S items into forced Epics")]
        itba_items = [("►", "Examples", "SAP config tweak · master-data fix · regulatory tweak"),
                      ("⚙", "Pain profile", "Safety usually 1–2 · Effort 1–2"),
                      ("◯", "Owner cadence", "weekly batch · Architect must review SAP changes"),
                      ("✦", "Watch out", "even S can break a SAP module — Architect always")]
    elif size == "M":
        cs_items = [("►", "Examples", "Genesys integ · CIS multi-OPCO · Smart Meter API"),
                    ("⚙", "Pain profile", "BV mod · UX often high · CS = 2-component formula"),
                    ("◯", "Heather's role", "reviews Epic outcome · ensures customer framing"),
                    ("✦", "Watch out", "don't define outcome as 'add feature X'")]
        itba_items = [("►", "Examples", "SAP enhancement · regulatory report · cross-VS portal"),
                      ("⚙", "Pain profile", "Safety 2–4 · ITBA = 3-component formula (Safety 25%)"),
                      ("◯", "Heather's role", "reviews Epic outcome · ensures business framing"),
                      ("✦", "Watch out", "regulatory M MUST be Fixed-Date (A3 enforces)")]
    else:
        cs_items = [("►", "Examples", "Genesys upgrade · CRM migration · Smart Meter rollout"),
                    ("⚙", "Pain profile", "OPCOmult often 1.6 · Effort 4–5"),
                    ("◯", "Heather's role", "presents Initiative roadmap at Gate 2"),
                    ("✦", "Capacity", "reserves 2–3 quarters · Director signs off")]
        itba_items = [("►", "Examples", "SAP module upgrade · regulatory framework · data platform"),
                      ("⚙", "Pain profile", "Safety often 4–5 · ITBA gives Safety 25% weight"),
                      ("◯", "Heather's role", "partners with Carlos on Architecture trade-offs"),
                      ("✦", "Capacity", "reserves 2–3 quarters · Director signs off")]

    _side_box(0.40, box_y, sw2, box_h, AVANGRID_GREEN, "CS  ·  Customer Services", cs_items)
    _side_box(0.40 + sw2 + 0.20, box_y, sw2, box_h, HEADER_BLUE, "ITBA  ·  IT Business Apps", itba_items)

    story_bridge(s, n); footer_bar(s, f"{n:02d}/20")
    return s

def build_pp11_path_S(prs):
    return build_pp_path(prs, 11, "S", "≤ 5 SP", AVANGRID_GREEN, "~1 week",
        [("Story", "no Epic", "step"),
         ("Refine", "Squad-level", "step"),
         ("DoR", "6 criteria", "gate"),
         ("Pull", "Squad pulls", "step")])

def build_pp12_path_M(prs):
    return build_pp_path(prs, 12, "M", "5–20 SP", GREEN_MID, "~1–2 weeks",
        [("A5 Auto-Epic", "Epic + 3 Stories", "fork"),
         ("VS Refine", "PM + Heather framing", "step"),
         ("Architect", "NFR pre-check", "step"),
         ("DoR (×N)", "per Story", "gate"),
         ("Pull", "Squad sequential", "step")])

def build_pp13_path_L(prs):
    return build_pp_path(prs, 13, "L", "20–60 SP", HEADER_BLUE, "~2–4 weeks",
        [("Initiative", "shape · Epics", "fork"),
         ("VS Refine", "PM + Heather", "step"),
         ("Architecture", "FULL NFR pass", "step"),
         ("Capacity", "multi-quarter", "step"),
         ("DoR phased", "early Stories first", "gate"),
         ("Pull", "Squad phased", "step")])


# ============================================================================
# 14 · DEMAND-ON-EXISTING shortcut · decision tree
# ============================================================================
def build_pp14_existing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "DEMAND-ON-EXISTING-INITIATIVE  ·  SHORTCUT",
                "Decision tree by parent status  ·  Squad PM accepts within 48h  ·  bypasses Gate 2 only when Implementing")
    setup(s, 14)

    root_x = 5.40; root_y = 1.55; root_w = 2.55; root_h = 0.80
    round_rect(s, Inches(root_x), Inches(root_y), Inches(root_w), Inches(root_h), HEADER_BLUE_DK, radius=0.10)
    text(s, Inches(root_x), Inches(root_y + 0.10), Inches(root_w), Inches(0.30),
         "WHAT IS THE PARENT", size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(root_x), Inches(root_y + 0.40), Inches(root_w), Inches(0.30),
         "INITIATIVE STATUS?", size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    statuses = [
        ("FUNNEL", "REJECT shortcut", GREY_LABEL, False),
        ("REVIEWING", "REJECT shortcut", GREY_LABEL, False),
        ("ANALYZING", "REJECT shortcut", GREY_LABEL, False),
        ("BACKLOG", "ACCEPT (cautious)", AMBER, True),
        ("IMPLEMENTING", "ACCEPT (preferred)", AVANGRID_GREEN, True),
    ]
    n = len(statuses)
    avail_w = 12.55
    branch_w = (avail_w - (n - 1) * 0.20) / n
    branch_y = 3.20
    for i, (status, decision, color, is_accept) in enumerate(statuses):
        bx = 0.40 + i * (branch_w + 0.20)
        line(s, Inches(root_x + root_w / 2), Inches(root_y + root_h),
             Inches(bx + branch_w / 2), Inches(branch_y), color, weight=1.5)
        round_rect(s, Inches(bx), Inches(branch_y), Inches(branch_w), Inches(0.55), color, radius=0.04)
        text(s, Inches(bx), Inches(branch_y), Inches(branch_w), Inches(0.55),
             status, size=10, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        round_rect(s, Inches(bx), Inches(branch_y + 0.65), Inches(branch_w), Inches(0.40),
                   WHITE, radius=0.30)
        text(s, Inches(bx), Inches(branch_y + 0.65), Inches(branch_w), Inches(0.40),
             decision, size=8, bold=True, color=color, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    detail_y = 4.40; detail_h = 2.40
    owners = ["PM (per VS)", "PM (per VS)", "PM + Architect", "Squad PM", "Squad PM"]
    paths = [
        "submit standalone\n· normal flow",
        "submit standalone\n· attach reference",
        "standalone · let\nPain Score compete",
        "Story under parent\n· re-baseline scope",
        "Story under parent\n· DoR fast-track",
    ]
    cos_v = ["—", "—", "—", "inherits parent", "inherits parent"]
    for i, (status, decision, color, is_accept) in enumerate(statuses):
        bx = 0.40 + i * (branch_w + 0.20)
        round_rect(s, Inches(bx), Inches(detail_y), Inches(branch_w), Inches(detail_h),
                   GREEN_BG if is_accept else GREY_BG, radius=0.04)
        line(s, Inches(bx + branch_w / 2), Inches(branch_y + 1.05),
             Inches(bx + branch_w / 2), Inches(detail_y), color, weight=1.0)
        text(s, Inches(bx + 0.10), Inches(detail_y + 0.08), Inches(branch_w - 0.20), Inches(0.18),
             "OWNER", size=7, bold=True, color=GREY_LABEL, font=BODY_FONT)
        text(s, Inches(bx + 0.10), Inches(detail_y + 0.26), Inches(branch_w - 0.20), Inches(0.22),
             owners[i], size=9, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(bx + 0.10), Inches(detail_y + 0.55), Inches(branch_w - 0.20), Inches(0.18),
             "PATH", size=7, bold=True, color=GREY_LABEL, font=BODY_FONT)
        text(s, Inches(bx + 0.10), Inches(detail_y + 0.72), Inches(branch_w - 0.20), Inches(0.65),
             paths[i], size=8, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(bx + 0.10), Inches(detail_y + 1.45), Inches(branch_w - 0.20), Inches(0.18),
             "CoS IMPACT", size=7, bold=True, color=GREY_LABEL, font=BODY_FONT)
        text(s, Inches(bx + 0.10), Inches(detail_y + 1.62), Inches(branch_w - 0.20), Inches(0.40),
             cos_v[i], size=8, color=DARK_TEXT, font=BODY_FONT)

    story_bridge(s, 14); footer_bar(s, "14/20")
    return s


# ============================================================================
# 15 · REFINEMENT · hierarchy tree + INVEST + 5 moves
# ============================================================================
def build_pp15_refinement(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "REFINEMENT  ·  WORK BREAKDOWN TREE",
                "Demand → Initiative → Epic → Story → Sub-task  ·  outcomes named  ·  INVEST applied at Story level")
    setup(s, 15)

    levels_x = [0.40, 2.65, 5.10, 8.10, 11.10]
    level_w  = [2.05, 2.25, 2.80, 2.80, 2.10]
    levels = [
        ("DEMAND", "intake placeholder", AVANGRID_GREEN),
        ("INITIATIVE", "L only · multi-quarter", HEADER_BLUE_DK),
        ("EPIC", "M/L · outcome holder", HEADER_BLUE),
        ("STORY", "S unit · ≤ 5 SP", GREEN_MID),
        ("SUB-TASK", "tech breakdown", GREY_LABEL),
    ]

    base_y_center = 2.55
    box_h = 0.55

    for i in range(3):
        x = levels_x[i]; w = level_w[i]
        round_rect(s, Inches(x), Inches(base_y_center), Inches(w), Inches(box_h), levels[i][2], radius=0.06)
        text(s, Inches(x), Inches(base_y_center + 0.02), Inches(w), Inches(0.22),
             levels[i][0], size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x), Inches(base_y_center + 0.26), Inches(w), Inches(0.26),
             levels[i][1], size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    story_ys = [1.55, 2.55, 3.55]
    story_w = level_w[3]
    story_x = levels_x[3]
    for j, sy in enumerate(story_ys):
        round_rect(s, Inches(story_x), Inches(sy), Inches(story_w), Inches(box_h), GREEN_MID, radius=0.06)
        text(s, Inches(story_x), Inches(sy + 0.02), Inches(story_w), Inches(0.22),
             f"STORY {j+1}", size=10, bold=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(story_x), Inches(sy + 0.26), Inches(story_w), Inches(0.26),
             "≤ 5 SP · INVEST", size=8, italic=True, color=WHITE, font=BODY_FONT, align=PP_ALIGN.CENTER)
        line(s, Inches(levels_x[2] + level_w[2]),
             Inches(base_y_center + box_h / 2),
             Inches(story_x), Inches(sy + box_h / 2), GREEN_MID, weight=1.5)

    sub_y_offsets = [-0.22, +0.22]
    sub_w = level_w[4]; sub_x = levels_x[4]; sub_h = 0.40
    for j in range(2):
        sy = story_ys[0] + sub_y_offsets[j] - sub_h / 2 + box_h / 2
        round_rect(s, Inches(sub_x), Inches(sy), Inches(sub_w), Inches(sub_h), GREY_LABEL, radius=0.06)
        text(s, Inches(sub_x), Inches(sy), Inches(sub_w), Inches(sub_h),
             f"SUB-TASK {j+1}", size=8, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        line(s, Inches(story_x + story_w),
             Inches(story_ys[0] + box_h / 2),
             Inches(sub_x), Inches(sy + sub_h / 2), GREY_LABEL, weight=1.0)
    text(s, Inches(sub_x), Inches(story_ys[1] + 0.08), Inches(sub_w), Inches(0.40),
         "(sub-tasks per Story\nas needed)",
         size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(sub_x), Inches(story_ys[2] + 0.08), Inches(sub_w), Inches(0.40),
         "(sub-tasks per Story\nas needed)",
         size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    for i in range(2):
        line(s, Inches(levels_x[i] + level_w[i]), Inches(base_y_center + box_h / 2),
             Inches(levels_x[i + 1]), Inches(base_y_center + box_h / 2),
             levels[i][2], weight=2.0)

    inv_y = 4.40
    round_rect(s, Inches(0.40), Inches(inv_y), Inches(12.55), Inches(0.95), GREEN_BG, radius=0.04)
    text(s, Inches(0.55), Inches(inv_y + 0.06), Inches(12.30), Inches(0.20),
         "INVEST  ·  what every Story must satisfy",
         size=10, bold=True, color=GREEN_DARK, font=BODY_FONT)
    invest = [("I", "Independent"), ("N", "Negotiable"), ("V", "Valuable"),
              ("E", "Estimable"), ("S", "Small (≤5 SP)"), ("T", "Testable")]
    cw = (12.55 - 0.30 - 5 * 0.10) / 6
    for i, (letter, lbl) in enumerate(invest):
        x = 0.55 + i * (cw + 0.10)
        oval(s, Inches(x + cw/2 - 0.20), Inches(inv_y + 0.30), Inches(0.40), Inches(0.40), GREEN_DARK)
        text(s, Inches(x + cw/2 - 0.20), Inches(inv_y + 0.30), Inches(0.40), Inches(0.40), letter,
             size=14, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(x), Inches(inv_y + 0.74), Inches(cw), Inches(0.20),
             lbl, size=8, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)

    moves_y = 5.50
    moves = [
        ("01", "OUTCOME FRAMING", "outcome (NPS · cycle · cost) — not deliverable"),
        ("02", "STORY SLICING", "INVEST applied · split if oversized"),
        ("03", "ACCEPTANCE CRIT.", "binary pass/fail · UAT-testable · Sponsor signs"),
        ("04", "DEPENDENCY MAP", "cross-VS · vendor · license — Architect leads"),
        ("05", "OUTCOME-vs-OUTPUT", "final pass before DoR — drift correction"),
    ]
    mw = (12.55 - 4 * 0.10) / 5
    for i, (num, lbl, body) in enumerate(moves):
        x = 0.40 + i * (mw + 0.10)
        round_rect(s, Inches(x), Inches(moves_y), Inches(mw), Inches(1.20), WHITE, radius=0.04)
        rect(s, Inches(x), Inches(moves_y), Inches(0.10), Inches(1.20), HEADER_BLUE)
        text(s, Inches(x + 0.20), Inches(moves_y + 0.06), Inches(0.40), Inches(0.30),
             num, size=14, bold=True, color=HEADER_BLUE, font=BODY_FONT)
        text(s, Inches(x + 0.65), Inches(moves_y + 0.10), Inches(mw - 0.75), Inches(0.25),
             lbl, size=9, bold=True, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(x + 0.20), Inches(moves_y + 0.45), Inches(mw - 0.30), Inches(0.65),
             body, size=8, color=DARK_TEXT, font=BODY_FONT)

    story_bridge(s, 15); footer_bar(s, "15/20")
    return s


# ============================================================================
# 16 · PRE-ESTIMATION + CAPACITY · big formulas + CoS bar
# ============================================================================
def build_pp16_estimation_capacity(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "PRE-ESTIMATION  +  CAPACITY ALLOCATION",
                "Two formulas (CS · ITBA)  ·  capacity reserved per CoS lane  ·  monthly Gate 2 decides %")
    setup(s, 16)

    fl_x = 0.40; fl_y = 1.55; fl_w = 6.30; fl_h = 5.30
    round_rect(s, Inches(fl_x), Inches(fl_y), Inches(fl_w), Inches(fl_h), WHITE, radius=0.04)
    rect(s, Inches(fl_x), Inches(fl_y), Inches(fl_w), Inches(0.42), HEADER_BLUE)
    text(s, Inches(fl_x + 0.15), Inches(fl_y + 0.06), Inches(fl_w - 0.30), Inches(0.30),
         "PAIN SCORE FORMULAS", size=11, bold=True, color=WHITE, font=BODY_FONT)

    cs_y = fl_y + 0.65
    round_rect(s, Inches(fl_x + 0.20), Inches(cs_y), Inches(fl_w - 0.40), Inches(2.10), GREEN_BG, radius=0.04)
    rect(s, Inches(fl_x + 0.20), Inches(cs_y), Inches(0.10), Inches(2.10), AVANGRID_GREEN)
    text(s, Inches(fl_x + 0.40), Inches(cs_y + 0.10), Inches(fl_w - 0.60), Inches(0.25),
         "CS  ·  no Safety  ·  utility-calibrated for customer outcomes",
         size=9, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(fl_x + 0.40), Inches(cs_y + 0.42), Inches(fl_w - 0.60), Inches(0.55),
         "(BV · OPCOmult · 0.40 + UX · 0.60)",
         size=15, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
    rect(s, Inches(fl_x + 0.50), Inches(cs_y + 1.00), Inches(fl_w - 0.80), Emu(20000), DARK_TEXT)
    text(s, Inches(fl_x + 0.40), Inches(cs_y + 1.05), Inches(fl_w - 0.60), Inches(0.55),
         "Effort", size=15, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
    chip(s, fl_x + 0.40, cs_y + 1.65, 1.30, "BV  40%", AVANGRID_GREEN, WHITE, h=0.30, size=8)
    chip(s, fl_x + 1.80, cs_y + 1.65, 1.30, "UX  60%", GREEN_MID, WHITE, h=0.30, size=8)
    chip(s, fl_x + 3.20, cs_y + 1.65, 1.30, "no Safety", GREY_LABEL, WHITE, h=0.30, size=8)

    it_y = cs_y + 2.30
    round_rect(s, Inches(fl_x + 0.20), Inches(it_y), Inches(fl_w - 0.40), Inches(2.30), BLUE_BG, radius=0.04)
    rect(s, Inches(fl_x + 0.20), Inches(it_y), Inches(0.10), Inches(2.30), HEADER_BLUE)
    text(s, Inches(fl_x + 0.40), Inches(it_y + 0.10), Inches(fl_w - 0.60), Inches(0.25),
         "ITBA  ·  Safety isolated · field-crew impact",
         size=9, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    text(s, Inches(fl_x + 0.40), Inches(it_y + 0.42), Inches(fl_w - 0.60), Inches(0.55),
         "(BV·OPCOmult·0.45 + UX·0.30 + Safety·0.25)",
         size=12, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
    rect(s, Inches(fl_x + 0.50), Inches(it_y + 1.05), Inches(fl_w - 0.80), Emu(20000), DARK_TEXT)
    text(s, Inches(fl_x + 0.40), Inches(it_y + 1.10), Inches(fl_w - 0.60), Inches(0.55),
         "Effort", size=15, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)
    chip(s, fl_x + 0.40, it_y + 1.75, 1.20, "BV  45%", HEADER_BLUE, WHITE, h=0.30, size=8)
    chip(s, fl_x + 1.70, it_y + 1.75, 1.20, "UX  30%", BLUE_LIGHT, WHITE, h=0.30, size=8)
    chip(s, fl_x + 3.00, it_y + 1.75, 1.40, "Safety  25%", HEADER_BLUE_DK, WHITE, h=0.30, size=8)

    cr_x = 7.00; cr_y = 1.55; cr_w = 5.95; cr_h = 5.30
    round_rect(s, Inches(cr_x), Inches(cr_y), Inches(cr_w), Inches(cr_h), WHITE, radius=0.04)
    rect(s, Inches(cr_x), Inches(cr_y), Inches(cr_w), Inches(0.42), GREEN_DARK)
    text(s, Inches(cr_x + 0.15), Inches(cr_y + 0.06), Inches(cr_w - 0.30), Inches(0.30),
         "CAPACITY ALLOCATION  ·  per CoS lane",
         size=11, bold=True, color=WHITE, font=BODY_FONT)
    text(s, Inches(cr_x + 0.15), Inches(cr_y + 0.50), Inches(cr_w - 0.30), Inches(0.22),
         "% per lane is decided MONTHLY at Gate 2 (illustrative split below)",
         size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)

    bar_y = cr_y + 0.85
    bar_h = 0.55
    bar_x = cr_x + 0.20
    bar_w = cr_w - 0.40
    lanes = [
        ("STANDARD", 55, AVANGRID_GREEN, "WIP 3–4"),
        ("FIXED-DATE", 20, AMBER, "WIP 2"),
        ("EXPEDITE", 15, RED_GATE, "WIP 1"),
        ("INTANGIBLE", 10, PURPLE_HEATHER, "WIP 1"),
    ]
    cum = 0
    for lbl, pct, color, wip in lanes:
        x = bar_x + (cum / 100) * bar_w
        w = (pct / 100) * bar_w
        rect(s, Inches(x), Inches(bar_y), Inches(w), Inches(bar_h), color)
        text(s, Inches(x), Inches(bar_y), Inches(w), Inches(bar_h), f"{pct}%",
             size=12, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cum += pct

    detail_y = bar_y + bar_h + 0.30
    rh = (cr_h - (detail_y - cr_y) - 0.20) / len(lanes)
    descs = {
        "STANDARD": "default home for most demands",
        "FIXED-DATE": "regulatory or deadline-locked",
        "EXPEDITE": "P1/P2 (CS) · critical (ITBA) · A2 enforces",
        "INTANGIBLE": "tech debt · refactor",
    }
    for i, (lbl, pct, color, wip) in enumerate(lanes):
        ry = detail_y + i * rh
        round_rect(s, Inches(cr_x + 0.20), Inches(ry), Inches(cr_w - 0.40), Inches(rh - 0.10),
                   GREY_BG, radius=0.04)
        rect(s, Inches(cr_x + 0.20), Inches(ry), Inches(0.10), Inches(rh - 0.10), color)
        text(s, Inches(cr_x + 0.40), Inches(ry + 0.06), Inches(2.0), Inches(0.22),
             lbl, size=10, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(cr_x + 2.30), Inches(ry + 0.06), Inches(0.80), Inches(0.22),
             wip, size=8, bold=True, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(cr_x + 0.40), Inches(ry + 0.30), Inches(cr_w - 0.60), Inches(rh - 0.40),
             descs.get(lbl, ""), size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)

    story_bridge(s, 16); footer_bar(s, "16/20")
    return s


# ============================================================================
# 17 · DUAL TRACK · swim lanes
# ============================================================================
def build_pp17_dual_track(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "DUAL TRACK  ·  DISCOVERY ↔ DELIVERY",
                "We are almost always the orange band  ·  Discovery is run by domain or vendor  ·  DoR is the bridge")
    setup(s, 17)

    top_y = 1.55; lane_h = 1.80
    round_rect(s, Inches(0.40), Inches(top_y), Inches(12.55), Inches(lane_h), GREY_BG, radius=0.04)
    rect(s, Inches(0.40), Inches(top_y), Inches(0.18), Inches(lane_h), GREY_LABEL)
    text(s, Inches(0.65), Inches(top_y + 0.10), Inches(5.0), Inches(0.30),
         "DISCOVERY", size=14, bold=True, color=GREY_LABEL, font=BODY_FONT)
    chip(s, 2.60, top_y + 0.13, 1.20, "NOT US", GREY_MUTED, WHITE, h=0.24, size=8)
    text(s, Inches(0.65), Inches(top_y + 0.42), Inches(5.0), Inches(0.20),
         "domain product team / vendor",
         size=9, italic=True, color=GREY_LABEL, font=BODY_FONT)

    p_y = top_y + 0.80; p_h = 0.85
    round_rect(s, Inches(0.65), Inches(p_y), Inches(5.85), Inches(p_h), WHITE, radius=0.04)
    text(s, Inches(0.80), Inches(p_y + 0.06), Inches(5.50), Inches(0.22),
         "PHASE 0  ·  PROBLEM SHAPING", size=10, bold=True, color=GREY_LABEL, font=BODY_FONT)
    text(s, Inches(0.80), Inches(p_y + 0.30), Inches(5.50), Inches(0.50),
         "problem statement · personas · hypotheses (value + risk)",
         size=8, color=DARK_TEXT, font=BODY_FONT)
    text(s, Inches(6.55), Inches(p_y + 0.20), Inches(0.30), Inches(0.50),
         "→", size=18, bold=True, color=GREY_LABEL, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    round_rect(s, Inches(6.85), Inches(p_y), Inches(5.85), Inches(p_h), WHITE, radius=0.04)
    text(s, Inches(7.00), Inches(p_y + 0.06), Inches(5.50), Inches(0.22),
         "PHASE 1  ·  MATERIALIZATION", size=10, bold=True, color=GREY_LABEL, font=BODY_FONT)
    text(s, Inches(7.00), Inches(p_y + 0.30), Inches(5.50), Inches(0.50),
         "solution shape · vendor model · roadmap · NFRs draft · budget",
         size=8, color=DARK_TEXT, font=BODY_FONT)

    db_y = top_y + lane_h + 0.10
    round_rect(s, Inches(4.30), Inches(db_y), Inches(4.70), Inches(0.45), GREEN_DARK, radius=0.20)
    text(s, Inches(4.30), Inches(db_y), Inches(4.70), Inches(0.45),
         "DoR CHECK  ·  the bridge between worlds",
         size=10, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(6.50), Inches(db_y - 0.30), Inches(0.30), Inches(0.30),
         "↓", size=18, bold=True, color=GREEN_DARK, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    bot_y = db_y + 0.55
    round_rect(s, Inches(0.40), Inches(bot_y), Inches(12.55), Inches(lane_h), ORANGE_BG, radius=0.04)
    rect(s, Inches(0.40), Inches(bot_y), Inches(0.18), Inches(lane_h), ORANGE_DELIVERY)
    text(s, Inches(0.65), Inches(bot_y + 0.10), Inches(5.0), Inches(0.30),
         "DELIVERY", size=14, bold=True, color=ORANGE_DELIVERY, font=BODY_FONT)
    chip(s, 2.20, bot_y + 0.13, 1.30, "OUR ZONE", ORANGE_DELIVERY, WHITE, h=0.24, size=8)
    text(s, Inches(0.65), Inches(bot_y + 0.42), Inches(5.0), Inches(0.20),
         "Digital Products squads",
         size=9, italic=True, color=ORANGE_DELIVERY, font=BODY_FONT)

    p_y = bot_y + 0.80; p_h = 0.85
    delivery = [("PHASE 2 · BUILD", "squad pulls Stories · iterations"),
                ("PHASE 3 · TEST + UAT", "internal QA + Sponsor UAT"),
                ("PHASE 4 · RELEASE", "hypercare · Measure & Learn")]
    pw = (12.55 - 0.50 - 2 * 0.20) / 3
    for i, (lbl, body) in enumerate(delivery):
        x = 0.65 + i * (pw + 0.20)
        round_rect(s, Inches(x), Inches(p_y), Inches(pw), Inches(p_h), WHITE, radius=0.04)
        text(s, Inches(x + 0.15), Inches(p_y + 0.06), Inches(pw - 0.30), Inches(0.22),
             lbl, size=10, bold=True, color=ORANGE_DELIVERY, font=BODY_FONT)
        text(s, Inches(x + 0.15), Inches(p_y + 0.30), Inches(pw - 0.30), Inches(0.50),
             body, size=8, color=DARK_TEXT, font=BODY_FONT)
        if i < 2:
            text(s, Inches(x + pw + 0.02), Inches(p_y + 0.20), Inches(0.16), Inches(0.50),
                 "→", size=18, bold=True, color=ORANGE_DELIVERY, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    cn_y = bot_y + lane_h + 0.10
    cn_h = 0.85
    sw2 = (12.55 - 0.20) / 2
    round_rect(s, Inches(0.40), Inches(cn_y), Inches(sw2), Inches(cn_h), GREEN_BG, radius=0.04)
    rect(s, Inches(0.40), Inches(cn_y), Inches(0.10), Inches(cn_h), AVANGRID_GREEN)
    text(s, Inches(0.55), Inches(cn_y + 0.06), Inches(sw2 - 0.30), Inches(0.22),
         "CS  ·  vendor-led Discovery · SaaS-heavy",
         size=10, bold=True, color=GREEN_DARK, font=BODY_FONT)
    text(s, Inches(0.55), Inches(cn_y + 0.30), Inches(sw2 - 0.30), Inches(cn_h - 0.35),
         "Genesys · CIS · Smart Meter platforms — Heather oversees vendor relationships",
         size=8, color=DARK_TEXT, font=BODY_FONT)
    cx2 = 0.40 + sw2 + 0.20
    round_rect(s, Inches(cx2), Inches(cn_y), Inches(sw2), Inches(cn_h), BLUE_BG, radius=0.04)
    rect(s, Inches(cx2), Inches(cn_y), Inches(0.10), Inches(cn_h), HEADER_BLUE)
    text(s, Inches(cx2 + 0.15), Inches(cn_y + 0.06), Inches(sw2 - 0.30), Inches(0.22),
         "ITBA  ·  domain-led Discovery · SAP-heavy",
         size=10, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    text(s, Inches(cx2 + 0.15), Inches(cn_y + 0.30), Inches(sw2 - 0.30), Inches(cn_h - 0.35),
         "SAP CoE · vendor portals · custom builds (Networks IT) — Architect leads NFR alignment",
         size=8, color=DARK_TEXT, font=BODY_FONT)

    story_bridge(s, 17); footer_bar(s, "17/20")
    return s


# ============================================================================
# 18 · ROADMAP · Gantt-style horizontal bands
# ============================================================================
def build_pp18_roadmap(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "ROADMAPPING  ·  ROLLING WAVE",
                "3 mo committed  ·  3–6 mo planned  ·  6–12 mo hypothesized  ·  re-balanced monthly at Gate 2")
    setup(s, 18)

    ax_y = 1.55; ax_w = 12.55; ax_x = 0.40
    rect(s, Inches(ax_x), Inches(ax_y + 0.30), Inches(ax_w), Emu(20000), GREY_BORDER)
    months = ["Now", "+3", "+6", "+9", "+12"]
    for i, m in enumerate(months):
        x = ax_x + i * (ax_w / 4)
        oval(s, Inches(x - 0.05), Inches(ax_y + 0.25), Inches(0.10), Inches(0.10), DARK_TEXT)
        text(s, Inches(x - 0.40), Inches(ax_y + 0.08), Inches(0.80), Inches(0.20),
             m, size=9, bold=True, color=DARK_TEXT, font=BODY_FONT, align=PP_ALIGN.CENTER)

    band_y_start = 2.10
    band_h = 0.80
    bands = [
        ("0–3  ·  COMMITTED", "sized · DoR-ready · capacity reserved", GREEN_DARK, GREEN_PASTEL, 0.0, 0.25),
        ("3–6  ·  PLANNED", "sized at Initiative · NFRs identified", GREEN_MID, GREEN_BG, 0.25, 0.50),
        ("6–12  ·  HYPOTHESIZED", "outcomes named · sizing rough · NOT promises", GREY_LABEL, GREY_BG, 0.50, 1.00),
    ]
    for i, (lbl, sub, color, bg, start_pct, end_pct) in enumerate(bands):
        by = band_y_start + i * (band_h + 0.10)
        bx = ax_x + start_pct * ax_w
        bw = (end_pct - start_pct) * ax_w
        round_rect(s, Inches(bx), Inches(by), Inches(bw), Inches(band_h), bg, radius=0.04)
        text(s, Inches(bx + 0.15), Inches(by + 0.06), Inches(bw - 0.30), Inches(0.25),
             lbl, size=11, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(bx + 0.15), Inches(by + 0.30), Inches(bw - 0.30), Inches(0.30),
             sub, size=9, italic=True, color=color, font=BODY_FONT)

    items_y = band_y_start + 3 * (band_h + 0.10) + 0.20
    items_h = 0.35
    cs_items = [
        ("Genesys script overhaul", 0.05, 0.20, AVANGRID_GREEN),
        ("Self-service password", 0.10, 0.22, AVANGRID_GREEN),
        ("CIS rate-plan migration", 0.30, 0.55, GREEN_MID),
        ("CRM mobile v2", 0.30, 0.50, GREEN_MID),
        ("Voicebot expansion", 0.65, 0.95, GREY_LABEL),
    ]
    text(s, Inches(0.40), Inches(items_y), Inches(2.0), Inches(0.22),
         "CS examples", size=9, bold=True, color=AVANGRID_GREEN, font=BODY_FONT)
    iy = items_y + 0.25
    for lbl, sx, ex, color in cs_items:
        bx = ax_x + sx * ax_w
        bw = (ex - sx) * ax_w
        round_rect(s, Inches(bx), Inches(iy), Inches(bw), Inches(items_h), color, radius=0.04)
        text(s, Inches(bx + 0.10), Inches(iy + 0.04), Inches(bw - 0.20), Inches(items_h - 0.08),
             lbl, size=8, bold=True, color=WHITE, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
        iy += items_h + 0.05

    iy += 0.10
    text(s, Inches(0.40), Inches(iy), Inches(2.0), Inches(0.22),
         "ITBA examples", size=9, bold=True, color=HEADER_BLUE, font=BODY_FONT)
    itba_items = [
        ("PUC tariff banner Q3", 0.06, 0.22, HEADER_BLUE_DK),
        ("SAP FI year-end close", 0.08, 0.20, HEADER_BLUE),
        ("Audit framework refresh", 0.32, 0.55, HEADER_BLUE),
        ("Networks IT data backbone", 0.30, 0.50, HEADER_BLUE),
        ("S/4HANA migration plan", 0.65, 0.95, GREY_LABEL),
    ]
    iy += 0.25
    for lbl, sx, ex, color in itba_items:
        bx = ax_x + sx * ax_w
        bw = (ex - sx) * ax_w
        round_rect(s, Inches(bx), Inches(iy), Inches(bw), Inches(items_h), color, radius=0.04)
        text(s, Inches(bx + 0.10), Inches(iy + 0.04), Inches(bw - 0.20), Inches(items_h - 0.08),
             lbl, size=8, bold=True, color=WHITE, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
        iy += items_h + 0.05

    story_bridge(s, 18); footer_bar(s, "18/20")
    return s


# ============================================================================
# 19 · DOR · gate visual + 6 checkmark circles + pass/fail flow
# ============================================================================
def build_pp19_dor(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "DEFINITION OF READY  ·  DoR",
                "The single gate between Preparation & Planning and Execution  ·  6 base + CS-extra + ITBA-extra")
    setup(s, 19)

    gate_y = 1.55
    round_rect(s, Inches(0.40), Inches(gate_y + 0.40), Inches(1.40), Inches(0.55), HEADER_BLUE, radius=0.30)
    text(s, Inches(0.40), Inches(gate_y + 0.40), Inches(1.40), Inches(0.55),
         "STORY", size=10, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.85), Inches(gate_y + 0.42), Inches(0.30), Inches(0.50),
         "→", size=18, bold=True, color=GREY_LABEL, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    cp_x_start = 2.30; cp_d = 1.20; cp_gap = 0.18
    criteria = [("OUTCOME", "UAT-testable"), ("SIZED", "≤ 5 SP"),
                ("ARCHITECT", "NFRs · vendor"), ("DEPS", "known"),
                ("CAPACITY", "reserved"), ("SPONSOR", "signed off")]
    for i, (lbl, sub) in enumerate(criteria):
        cx_p = cp_x_start + i * (cp_d + cp_gap)
        oval(s, Inches(cx_p), Inches(gate_y + 0.10), Inches(cp_d), Inches(cp_d), AVANGRID_GREEN)
        text(s, Inches(cx_p), Inches(gate_y + 0.10), Inches(cp_d), Inches(cp_d), "✓",
             size=24, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(cx_p - 0.20), Inches(gate_y + cp_d + 0.18),
             Inches(cp_d + 0.40), Inches(0.22), lbl,
             size=9, bold=True, color=GREEN_DARK, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(cx_p - 0.20), Inches(gate_y + cp_d + 0.40),
             Inches(cp_d + 0.40), Inches(0.22), sub,
             size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
        if i < 5:
            text(s, Inches(cx_p + cp_d), Inches(gate_y + 0.45), Inches(0.18), Inches(0.50),
                 "→", size=14, bold=True, color=AVANGRID_GREEN, font=BODY_FONT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    pass_y = gate_y + 0.40
    last_x = cp_x_start + 5 * (cp_d + cp_gap) + cp_d + 0.10
    text(s, Inches(last_x + 0.10), Inches(pass_y), Inches(0.30), Inches(0.50),
         "→", size=18, bold=True, color=GREEN_DARK, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    round_rect(s, Inches(last_x + 0.50), Inches(pass_y), Inches(1.65), Inches(0.55),
               GREEN_DARK, radius=0.30)
    text(s, Inches(last_x + 0.50), Inches(pass_y), Inches(1.65), Inches(0.55),
         "READY", size=11, bold=True, color=WHITE, font=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    fail_y = gate_y + cp_d + 0.95
    rect(s, Inches(2.30), Inches(fail_y), Inches(last_x - 2.20), Emu(20000), RED_GATE)
    line(s, Inches(2.30), Inches(fail_y), Inches(2.30), Inches(gate_y + 0.65),
         RED_GATE, weight=2)
    line(s, Inches(last_x - 0.10), Inches(fail_y), Inches(last_x - 0.10), Inches(gate_y + 0.65),
         RED_GATE, weight=2)
    text(s, Inches(2.30), Inches(fail_y + 0.05), Inches(last_x - 2.20), Inches(0.22),
         "✕  FAIL → loop back to refinement",
         size=9, bold=True, color=RED_GATE, font=BODY_FONT, align=PP_ALIGN.CENTER)

    bot_y = 4.40; bot_h = 2.45
    sw2 = (12.55 - 0.20) / 2
    round_rect(s, Inches(0.40), Inches(bot_y), Inches(sw2), Inches(bot_h), WHITE, radius=0.04)
    rect(s, Inches(0.40), Inches(bot_y), Inches(sw2), Inches(0.42), AVANGRID_GREEN)
    text(s, Inches(0.55), Inches(bot_y + 0.06), Inches(sw2 - 0.30), Inches(0.30),
         "CS-EXTRA  ·  Customer Services additions",
         size=11, bold=True, color=WHITE, font=BODY_FONT)
    cs_extra = [
        ("CIT P1/P2/P3/P4 recorded", "drives Auto-Expedite when high"),
        ("UX validated by Sponsor", "customer-rep voice"),
        ("Vendor SLA noted (if touched)", "Genesys / CIS / Smart Meter"),
        ("Multi-OPCO rollout planned", "when OPCOmult > 1.0"),
    ]
    iy = bot_y + 0.55
    rh = (bot_h - 0.65) / len(cs_extra)
    for lbl, sub in cs_extra:
        oval(s, Inches(0.60), Inches(iy + 0.12), Inches(0.20), Inches(0.20), AVANGRID_GREEN)
        text(s, Inches(0.90), Inches(iy + 0.06), Inches(sw2 - 1.0), Inches(0.20),
             lbl, size=9, bold=True, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(0.90), Inches(iy + 0.26), Inches(sw2 - 1.0), Inches(rh - 0.30),
             sub, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)
        iy += rh

    cx2 = 0.40 + sw2 + 0.20
    round_rect(s, Inches(cx2), Inches(bot_y), Inches(sw2), Inches(bot_h), WHITE, radius=0.04)
    rect(s, Inches(cx2), Inches(bot_y), Inches(sw2), Inches(0.42), HEADER_BLUE)
    text(s, Inches(cx2 + 0.15), Inches(bot_y + 0.06), Inches(sw2 - 0.30), Inches(0.30),
         "ITBA-EXTRA  ·  IT Business Apps additions",
         size=11, bold=True, color=WHITE, font=BODY_FONT)
    itba_extra = [
        ("Regulatory flag clear", "deadline pinned if Fixed-Date"),
        ("Safety scored (1–5)", "Architect responsibility"),
        ("SAP module impact noted", "rollback plan named"),
        ("Audit trail / data integrity", "NFRs documented + traceable"),
    ]
    iy = bot_y + 0.55
    for lbl, sub in itba_extra:
        oval(s, Inches(cx2 + 0.20), Inches(iy + 0.12), Inches(0.20), Inches(0.20), HEADER_BLUE)
        text(s, Inches(cx2 + 0.50), Inches(iy + 0.06), Inches(sw2 - 1.0), Inches(0.20),
             lbl, size=9, bold=True, color=DARK_TEXT, font=BODY_FONT)
        text(s, Inches(cx2 + 0.50), Inches(iy + 0.26), Inches(sw2 - 1.0), Inches(rh - 0.30),
             sub, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT)
        iy += rh

    story_bridge(s, 19); footer_bar(s, "19/20")
    return s


# ============================================================================
# 20 · ROLES CAST · org-chart visual
# ============================================================================
def build_pp20_roles_cast(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title(s, "ROLES CAST  ·  THE FULL ORG",
                "Who does what  ·  CS · ITBA · Heather (Product) · middle managers · NEW PMs · Architect · Squad PMs · Vendor")
    setup(s, 20)

    vp_x = 5.90; vp_y = 1.55; vp_w = 1.55
    avatar(s, vp_x, vp_y, 0.90, HEADER_BLUE_DK, "VV")
    text(s, Inches(vp_x - 0.50), Inches(vp_y + 0.95), Inches(vp_w + 0.40), Inches(0.22),
         "VP  ·  Vitor Cruz", size=10, bold=True, color=HEADER_BLUE_DK, font=BODY_FONT, align=PP_ALIGN.CENTER)
    text(s, Inches(vp_x - 0.50), Inches(vp_y + 1.18), Inches(vp_w + 0.40), Inches(0.22),
         "chairs Gate 2  ·  final escalation",
         size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)

    dir_y = 2.55; dir_d = 0.85
    dirs = [
        (3.00, "KJ", AVANGRID_GREEN, "KUMAR JALADANKI", "Director · CS Engineering"),
        (6.05, "HK", PURPLE_HEATHER, "HEATHER LaJOIE", "Director · Product (NEW)"),
        (9.10, "CG", HEADER_BLUE, "CARLOS GARBIN", "Director · ITBA"),
    ]
    for x, init, color, name, role in dirs:
        avatar(s, x, dir_y, dir_d, color, init)
        text(s, Inches(x - 0.80), Inches(dir_y + dir_d + 0.05), Inches(dir_d + 1.60), Inches(0.22),
             name, size=10, bold=True, color=color, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x - 0.80), Inches(dir_y + dir_d + 0.27), Inches(dir_d + 1.60), Inches(0.22),
             role, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
        line(s, Inches(vp_x + 0.45), Inches(vp_y + 0.90),
             Inches(x + dir_d / 2), Inches(dir_y), GREY_LABEL, weight=1)

    text(s, Inches(4.50), Inches(dir_y + dir_d + 0.55), Inches(4.30), Inches(0.22),
         "★ Heather spans CS and ITBA — Product is the bridge",
         size=8, bold=True, italic=True, color=PURPLE_HEATHER, font=BODY_FONT, align=PP_ALIGN.CENTER)

    mm_y = 4.10; mm_d = 0.55
    mm_data = [
        (2.40, "MM", AVANGRID_GREEN, "CS middle managers", "day-to-day capacity"),
        (10.15, "MM", HEADER_BLUE, "ITBA middle managers", "day-to-day capacity"),
    ]
    for x, init, color, name, role in mm_data:
        avatar(s, x, mm_y, mm_d, color, init)
        text(s, Inches(x - 0.60), Inches(mm_y + mm_d + 0.05), Inches(mm_d + 1.20), Inches(0.22),
             name, size=9, bold=True, color=color, font=BODY_FONT, align=PP_ALIGN.CENTER)
        text(s, Inches(x - 0.60), Inches(mm_y + mm_d + 0.27), Inches(mm_d + 1.20), Inches(0.22),
             role, size=8, italic=True, color=GREY_LABEL, font=BODY_FONT, align=PP_ALIGN.CENTER)
        target_dir = dirs[0] if x < 5 else dirs[2]
        line(s, Inches(target_dir[0] + dir_d / 2), Inches(dir_y + dir_d),
             Inches(x + mm_d / 2), Inches(mm_y), GREY_LABEL, weight=1)

    pm_y = 4.10; pm_d = 0.40
    n_cs_pm = 4; n_itba_pm = 5
    n_pm = n_cs_pm + n_itba_pm
    pm_total_w = n_pm * pm_d + (n_pm - 1) * 0.15
    pm_x_start = 6.65 - pm_total_w / 2
    for i in range(n_pm):
        x = pm_x_start + i * (pm_d + 0.15)
        col = AVANGRID_GREEN if i < n_cs_pm else HEADER_BLUE
        oval(s, Inches(x), Inches(pm_y), Inches(pm_d), Inches(pm_d), col)
        text(s, Inches(x), Inches(pm_y), Inches(pm_d), Inches(pm_d), "PM",
             size=8, bold=True, color=WHITE, font=BODY_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        line(s, Inches(dirs[1][0] + dir_d / 2), Inches(dir_y + dir_d),
             Inches(x + pm_d / 2), Inches(pm_y), PURPLE_LIGHT, weight=0.75)
    text(s, Inches(pm_x_start - 0.40), Inches(pm_y + pm_d + 0.10),
         Inches(pm_total_w + 0.80), Inches(0.20),
         "9 NEW PMs  ·  per VS  ·  hired by Heather",
         size=8, bold=True, color=PURPLE_HEATHER, font=BODY_FONT, align=PP_ALIGN.CENTER)

    bot_y = 5.50; bot_h = 1.40
    others = [
        ("ARCHITECT", "cross-VS · NFRs", AMBER, 0.40),
        ("SQUAD PMs", "per squad · DoR + pull", GREEN_DARK, 3.20),
        ("SPONSOR", "business owner · UAT signoff", HEADER_BLUE, 6.50),
        ("VENDOR / DOMAIN", "runs Discovery · Phase 0/1", GREY_LABEL, 9.60),
    ]
    bw = 3.05
    for lbl, sub, color, x in others:
        round_rect(s, Inches(x), Inches(bot_y), Inches(bw), Inches(bot_h), WHITE, radius=0.04)
        rect(s, Inches(x), Inches(bot_y), Inches(0.10), Inches(bot_h), color)
        text(s, Inches(x + 0.20), Inches(bot_y + 0.10), Inches(bw - 0.30), Inches(0.30),
             lbl, size=11, bold=True, color=color, font=BODY_FONT)
        text(s, Inches(x + 0.20), Inches(bot_y + 0.45), Inches(bw - 0.30), Inches(0.45),
             sub, size=9, color=DARK_TEXT, font=BODY_FONT)
        avatar(s, x + 0.30, bot_y + 0.85, 0.40, color, lbl[:1])
        avatar(s, x + 0.85, bot_y + 0.85, 0.40, color, lbl[:1])
        text(s, Inches(x + 1.40), Inches(bot_y + 0.93), Inches(bw - 1.5), Inches(0.30),
             "...", size=14, bold=True, color=color, font=BODY_FONT)

    story_bridge(s, 20); footer_bar(s, "20/20")
    return s


# ============================================================================
# Driver
# ============================================================================
def main():
    print("Opening:", SRC)
    prs = Presentation(SRC)
    n_before = len(prs.slides)

    print("Building Preparation & Planning v2 · VISUAL redesign (20 slides)...")
    build_pp01_divider(prs);              print("   01")
    build_pp02_purpose(prs);              print("   02")
    build_pp03_big_picture(prs);          print("   03")
    build_pp04_governance(prs);           print("   04")
    build_pp05_cs_vs_itba(prs);           print("   05")
    build_pp06_heather(prs);              print("   06")
    build_pp07_portfolio_cs(prs);         print("   07")
    build_pp08_portfolio_itba(prs);       print("   08")
    build_pp09_vs_cs(prs);                print("   09")
    build_pp10_vs_itba(prs);              print("   10")
    build_pp11_path_S(prs);               print("   11")
    build_pp12_path_M(prs);               print("   12")
    build_pp13_path_L(prs);               print("   13")
    build_pp14_existing(prs);             print("   14")
    build_pp15_refinement(prs);           print("   15")
    build_pp16_estimation_capacity(prs);  print("   16")
    build_pp17_dual_track(prs);           print("   17")
    build_pp18_roadmap(prs);              print("   18")
    build_pp19_dor(prs);                  print("   19")
    build_pp20_roles_cast(prs);           print("   20")

    n_after = len(prs.slides)
    print(f"\nSlides: {n_before} → {n_after}")

    prs.save(DST)

    # de-dup zip
    tmp = DST + ".clean"
    with zipfile.ZipFile(DST, "r") as zin:
        seen = set(); items = []
        for info in zin.infolist():
            if info.filename in seen: continue
            seen.add(info.filename); items.append(info)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in items:
                zout.writestr(info, zin.read(info.filename))
    shutil.move(tmp, DST)
    print(f"\n✓ Saved (deduped): {DST}")


if __name__ == "__main__":
    main()
