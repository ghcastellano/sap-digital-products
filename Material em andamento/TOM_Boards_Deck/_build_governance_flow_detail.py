"""
Editable PPTX: Governance Flow Detail.
4 slides — Team Kanban · Field Progression · Full Flow Overview · Gates Deep Dive.
GREEN TONES ONLY (Avangrid palette · no multi-color · accent in deep green).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "/Users/gustavohenriquecastellano/Downloads/Base Apresentações Projeto Digital Product/Governance Flow Detail.pptx"

# ============ STRICT GREEN PALETTE ============
GREEN_DARKEST = RGBColor(0x00, 0x36, 0x21)
GREEN_DARK = RGBColor(0x00, 0x4D, 0x2E)
GREEN = RGBColor(0x00, 0x6B, 0x3F)
GREEN_MID = RGBColor(0x4F, 0x91, 0x4F)
GREEN_SOFT = RGBColor(0x9E, 0xC7, 0xA0)
GREEN_PASTEL = RGBColor(0xC8, 0xE0, 0xCB)
GREEN_LIGHT = RGBColor(0xE6, 0xF2, 0xEB)
GREEN_VLIGHT = RGBColor(0xF2, 0xF8, 0xF4)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x17, 0x2B, 0x4D)
GREY_LABEL = RGBColor(0x6B, 0x77, 0x8C)
GREY_TEXT = RGBColor(0x42, 0x52, 0x6E)
GREY_BORDER = RGBColor(0xC9, 0xD2, 0xDD)
GREY_BG = RGBColor(0xF4, 0xF5, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.shadow.inherit = False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line
        if line_w is not None: s.line.width = line_w
    return s


def round_rect(slide, x, y, w, h, fill=None, line=None, line_w=None, radius=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.shadow.inherit = False
    s.adjustments[0] = radius
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line
        if line_w is not None: s.line.width = line_w
    return s


def text(slide, x, y, w, h, txt, size=11, bold=False, italic=False, color=TEXT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = Emu(36000); tf.margin_right = Emu(36000)
    tf.margin_top = Emu(18000); tf.margin_bottom = Emu(18000)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tb


def header_band(slide, title, subtitle):
    rect(slide, 0, 0, prs.slide_width, Inches(0.95), GREEN_DARK)
    text(slide, Inches(0.4), Inches(0.13), Inches(11), Inches(0.4),
         title, 22, bold=True, color=WHITE)
    text(slide, Inches(0.4), Inches(0.55), Inches(11), Inches(0.3),
         subtitle, 12, bold=True, color=GREEN_LIGHT)
    text(slide, Inches(11.3), Inches(0.30), Inches(1.7), Inches(0.4),
         "NTT  DATA", 16, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    rect(slide, 0, Inches(0.95), prs.slide_width, Emu(19050), GREEN)


def footer(slide, page_num, total):
    text(slide, Inches(0.4), Inches(7.30), Inches(8), Inches(0.18),
         "Avangrid · Digital Products · Target Operating Model",
         8, italic=True, color=GREY_LABEL)
    text(slide, Inches(12.4), Inches(7.30), Inches(0.8), Inches(0.18),
         f"{page_num:02d} / {total}", 8, color=GREY_LABEL, align=PP_ALIGN.RIGHT)


# ============================================================================
# SLIDE 1 — TEAM KANBAN DETAIL (squad-level · Phases 0-4 + Handoff)
# ============================================================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])
header_band(s1, "TEAM KANBAN  ·  SQUAD-LEVEL FLOW",
            "Cards = Stories (Epic only when Initiative > 3 months) · 5 phases + Discovery↔Delivery handoff")

# Top strip — context
ctx_y = Inches(1.2)
round_rect(s1, Inches(0.4), ctx_y, Inches(12.55), Inches(0.55),
           GREEN_LIGHT, GREEN, Pt(1), 0.04)
text(s1, Inches(0.55), ctx_y + Inches(0.06), Inches(12), Inches(0.18),
     "PURPOSE", 8, bold=True, color=GREEN_DARK)
text(s1, Inches(0.55), ctx_y + Inches(0.22), Inches(12.2), Inches(0.32),
     "How a squad executes work after pulling from the Portfolio Board's READY_FOR_EXEC column. "
     "Two tracks (Discovery + Delivery) feed each other continuously · WIP-limited at every phase.",
     9, color=TEXT)

# 5 Phase cards in horizontal flow + Handoff at bottom
phase_y = Inches(1.95)
phase_h = Inches(4.5)
phase_w = Inches(2.45)
phase_gap = Inches(0.07)

phases = [
    ("Phase 0", "DISCOVERY",
     "Problem framing · users · hypothesis",
     "Always-on · supports Materialization",
     "Product Manager",
     "Validated user problem",
     "continuous · weekly review",
     GREEN),
    ("Phase 1", "MATERIALIZATION",
     "Definition · solutioning · sizing",
     "Output: Roadmap · Budget · Sizing",
     "Product Manager + Tech Lead",
     "Roadmap, Budget, Sizing",
     "2–4 weeks per item",
     GREEN_DARK),
    ("Phase 2", "CREATION",
     "Planning · team / vendor assignment",
     "Output: Outcome definition · KPIs · Strategic alignment",
     "Product Manager",
     "Outcome definition · KPIs",
     "1 sprint · pre-build kick-off",
     GREEN),
    ("Phase 3", "ONGOING",
     "Build · test · deploy iteratively",
     "Kanban flow · daily / replenishment / service review · WIP limits per stage",
     "Squad (PM + Tech Lead + BA + dev/QA)",
     "Released increments",
     "continuous · daily flow review",
     GREEN_DARK),
    ("Phase 4", "MAINTENANCE",
     "Hypercare · stabilization · knowledge",
     "Handoff to BAU within 30 days post-launch",
     "Operations + Tech Lead",
     "Stable production state",
     "30-day hypercare window",
     GREEN_MID),
]

for i, (phase_num, phase_name, desc, sub, owner, output, cadence, fill) in enumerate(phases):
    px = Inches(0.4 + i * (2.45 + 0.07))
    round_rect(s1, px, phase_y, phase_w, phase_h, WHITE, fill, Pt(2), 0.04)
    # Top header
    round_rect(s1, px, phase_y, phase_w, Inches(0.7), fill, radius=0.04)
    text(s1, px, phase_y + Inches(0.05), phase_w, Inches(0.22),
         phase_num, 9, bold=True, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)
    text(s1, px, phase_y + Inches(0.27), phase_w, Inches(0.32),
         phase_name, 14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Body
    by = phase_y + Inches(0.85)
    text(s1, px + Inches(0.15), by, phase_w - Inches(0.3), Inches(0.4),
         desc, 10, bold=True, color=TEXT)
    text(s1, px + Inches(0.15), by + Inches(0.4), phase_w - Inches(0.3), Inches(0.55),
         sub, 8, italic=True, color=GREEN_DARK)

    # OWNER
    text(s1, px + Inches(0.15), by + Inches(1.05), phase_w - Inches(0.3), Inches(0.18),
         "OWNER", 7, bold=True, color=GREY_LABEL)
    text(s1, px + Inches(0.15), by + Inches(1.22), phase_w - Inches(0.3), Inches(0.45),
         owner, 9, color=TEXT)

    # OUTPUT
    text(s1, px + Inches(0.15), by + Inches(1.75), phase_w - Inches(0.3), Inches(0.18),
         "OUTPUT", 7, bold=True, color=GREY_LABEL)
    text(s1, px + Inches(0.15), by + Inches(1.92), phase_w - Inches(0.3), Inches(0.45),
         output, 9, color=TEXT)

    # CADENCE
    text(s1, px + Inches(0.15), by + Inches(2.45), phase_w - Inches(0.3), Inches(0.18),
         "CADENCE", 7, bold=True, color=GREY_LABEL)
    text(s1, px + Inches(0.15), by + Inches(2.62), phase_w - Inches(0.3), Inches(0.45),
         cadence, 9, italic=True, color=GREEN_DARK)

# Connector arrows between phases
for i in range(4):
    arr_x = Inches(0.4 + (i + 1) * (2.45 + 0.07) - 0.05)
    arrow = s1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                arr_x - Inches(0.04),
                                phase_y + Inches(0.27),
                                Inches(0.07),
                                Inches(0.18))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = GREEN_DARK
    arrow.line.fill.background()

# Discovery ↔ Delivery handoff (bottom strip)
hf_y = Inches(6.55)
round_rect(s1, Inches(0.4), hf_y, Inches(12.55), Inches(0.6),
           GREEN_LIGHT, GREEN, Pt(2), 0.04)
text(s1, Inches(0.55), hf_y + Inches(0.05), Inches(12), Inches(0.20),
     "DISCOVERY  ↔  DELIVERY HANDOFF  ·  the gatekeeping rule that protects flow",
     10, bold=True, color=GREEN_DARK)
text(s1, Inches(0.55), hf_y + Inches(0.28), Inches(8), Inches(0.30),
     "Items move from Phase 0/1 (Discovery) to Phase 2/3 (Delivery) ONLY when DoR is met. "
     "Discovery feeds Delivery continuously · never breaks WIP limits.",
     9, color=TEXT)
text(s1, Inches(8.7), hf_y + Inches(0.28), Inches(4.0), Inches(0.30),
     "GATEKEEPER  ·  Product Manager",
     9, bold=True, italic=True, color=GREEN_DARK, align=PP_ALIGN.RIGHT)

footer(s1, 1, 4)


# ============================================================================
# SLIDE 2 — FIELD PROGRESSION (who fills what · when)
# ============================================================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])
header_band(s2, "FIELD PROGRESSION  ·  WHO FILLS WHAT  ·  WHEN",
            "Progressive elaboration — requester does not fill all fields at submit · card grows as it advances")

# 4 columns: SUBMIT · REVIEWING · ANALYZING · APPROVAL GATE
col_y = Inches(1.2)
col_h = Inches(5.95)
col_w = Inches(3.05)
col_gap = Inches(0.10)

stages = [
    ("STAGE 1  ·  SUBMIT",
     "Funnel column",
     "Requester (Sponsor)",
     "—",
     GREEN_PASTEL, GREEN_DARK,
     [
         ("WHAT", "Title · Description · Business Value (narrative) · Acceptance Criteria"),
         ("ROUTING", "Value Stream destination · OPCO"),
         ("PEOPLE", "Sponsor (named user picker)"),
         ("FLAGS", "Regulatory flag (Yes/No) · Customer Impact Tier (if known: P1–P4)"),
     ]),
    ("STAGE 2  ·  REVIEWING",
     "between Gate 1 → Analyzing",
     "Demand Manager + Sponsor",
     "SLA: 2 weeks",
     GREEN_SOFT, GREEN_DARK,
     [
         ("CLASSIFICATION", "Type (Demand / Incident / Change)"),
         ("SIZE (initial)", "T-shirt: S · M · L · XL — best-guess"),
         ("CoS (initial)", "Standard / Fixed-Date / Expedite / Intangible — best-guess"),
         ("LENSES", "Work Type (BAU / Build / Upgrade) · Run / Grow / Transform"),
     ]),
    ("STAGE 3  ·  ANALYZING",
     "between Reviewing → Gate 2",
     "Demand Manager + Architect",
     "Pain Score auto-computed",
     GREEN_MID, WHITE,
     [
         ("SCORES (1–5)", "BV · UX · Safety · Effort — scored at triage"),
         ("PAIN SCORE", "H(x) = (BV·w₁ + UX·w₂ + Safety·w₃) / Effort — auto"),
         ("FINAL SIZE", "T-shirt locked after architect review"),
         ("ARTIFACTS", "Discovery Report · Concept Prototype · Feasibility check"),
     ]),
    ("STAGE 4  ·  GATE 2 (Portfolio)",
     "Approval Gate · before Backlog",
     "Director (Carlos / Kumar)",
     "Cadence: monthly",
     GREEN_DARK, WHITE,
     [
         ("DECISION", "Go / No-Go / Defer"),
         ("CAPACITY", "Slot to Value Stream Team · quarter assignment"),
         ("INPUTS REVIEWED", "Pain Score · Discovery Report · Feasibility · capacity availability"),
         ("OUTPUTS", "Backlog rank · status = Ready_For_Exec"),
     ]),
]

for i, (label, where, who, sla, fill_color, txt_color, blocks) in enumerate(stages):
    cx = Inches(0.4 + i * (3.05 + 0.10))

    # Outer card
    round_rect(s2, cx, col_y, col_w, col_h, WHITE, fill_color, Pt(2), 0.03)

    # Top header
    round_rect(s2, cx, col_y, col_w, Inches(1.05), fill_color, radius=0.03)
    text(s2, cx + Inches(0.15), col_y + Inches(0.08), col_w - Inches(0.3), Inches(0.22),
         label, 11, bold=True, color=txt_color)
    text(s2, cx + Inches(0.15), col_y + Inches(0.32), col_w - Inches(0.3), Inches(0.20),
         "▸ " + where, 8, italic=True, color=txt_color)
    text(s2, cx + Inches(0.15), col_y + Inches(0.55), col_w - Inches(0.3), Inches(0.20),
         "OWNER", 7, bold=True, color=txt_color)
    text(s2, cx + Inches(0.15), col_y + Inches(0.72), col_w - Inches(0.3), Inches(0.20),
         who, 9, bold=True, color=txt_color)
    text(s2, cx + Inches(0.15), col_y + Inches(0.88), col_w - Inches(0.3), Inches(0.18),
         sla, 7, italic=True, color=txt_color)

    # Body — fields added at this stage
    by = col_y + Inches(1.20)
    text(s2, cx + Inches(0.15), by, col_w - Inches(0.3), Inches(0.22),
         "FIELDS ADDED AT THIS STAGE",
         8, bold=True, color=GREY_LABEL)
    by += Inches(0.28)
    for label_b, content in blocks:
        round_rect(s2, cx + Inches(0.15), by, col_w - Inches(0.3), Inches(0.95),
                   GREEN_VLIGHT, fill_color, Pt(0.5), 0.05)
        text(s2, cx + Inches(0.25), by + Inches(0.06), col_w - Inches(0.5), Inches(0.20),
             label_b, 8, bold=True, color=GREEN_DARK)
        text(s2, cx + Inches(0.25), by + Inches(0.26), col_w - Inches(0.5), Inches(0.65),
             content, 9, color=TEXT)
        by += Inches(1.05)

# Connector arrows between stages
for i in range(3):
    arr_x = Inches(0.4 + (i + 1) * (3.05 + 0.10) - 0.05)
    arrow = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                arr_x - Inches(0.05),
                                col_y + Inches(0.4),
                                Inches(0.10),
                                Inches(0.20))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = GREEN_DARK
    arrow.line.fill.background()

footer(s2, 2, 4)


# ============================================================================
# SLIDE 3 — FULL FLOW OVERVIEW (Funnel → Done)
# ============================================================================
s3 = prs.slides.add_slide(prs.slide_layouts[6])
header_band(s3, "GOVERNANCE FLOW  ·  FUNNEL → DONE",
            "End-to-end view · 6 columns · 2 gates · BRP overlay · the full path of an Initiative")

# BRP banner at top
brp_y = Inches(1.15)
round_rect(s3, Inches(0.4), brp_y, Inches(12.55), Inches(0.36),
           GREEN_LIGHT, GREEN, Pt(1), 0.04)
text(s3, Inches(0.4), brp_y, Inches(12.55), Inches(0.36),
     "▲  BIG ROOM PLANNING  ·  quarterly · ½ day · realigns Backlog reorder + Implementing commitments across both Portfolio Boards",
     10, bold=True, italic=True, color=GREEN_DARK,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Main flow — horizontal kanban with detailed boxes
flow_y = Inches(1.85)
flow_h = Inches(4.5)
n_stages = 6
stage_w = (Inches(12.55) - 5 * Inches(0.05)) / n_stages

stages_overview = [
    ("Funnel", "Idea submitted",
     "OWNER", "Sponsor",
     "ENTRY", "Demand placeholder created", GREEN_SOFT, GREEN_DARKEST),
    ("Reviewing", "Business case refined · classification",
     "OWNER", "Sponsor + Demand Manager",
     "SLA", "2 weeks · Type · Size · CoS guess", GREEN_PASTEL, GREEN_DARKEST),
    ("Analyzing", "Sized · scored · feasibility",
     "OWNER", "Demand Manager + Architect",
     "OUTPUT", "Pain Score · final T-shirt · CoS", GREEN_MID, WHITE),
    ("Backlog", "Approved · awaiting capacity",
     "OWNER", "Portfolio Manager",
     "RANKED BY", "CoS + Pain Score", GREEN, WHITE),
    ("Implementing", "Pulled into Squad Kanban",
     "OWNER", "Team PM",
     "TRACKED", "Layer 0 VP Dashboard", GREEN_DARK, WHITE),
    ("Done", "Value realized · review",
     "OWNER", "Sponsor + PM",
     "OUTPUT", "Outcome metrics · lessons", GREEN_DARKEST, WHITE),
]

for i, (title, desc, l1, v1, l2, v2, fill, txc) in enumerate(stages_overview):
    sx = Inches(0.4) + i * (stage_w + Inches(0.05))
    # Outer
    round_rect(s3, sx, flow_y, stage_w, flow_h, WHITE, fill, Pt(2), 0.03)
    # Top header
    rect(s3, sx, flow_y, stage_w, Inches(0.5), fill)
    text(s3, sx, flow_y + Inches(0.04), stage_w, Inches(0.4),
         title, 14, bold=True, color=txc, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Description
    text(s3, sx + Inches(0.1), flow_y + Inches(0.6), stage_w - Inches(0.2), Inches(0.7),
         desc, 9, italic=True, color=GREEN_DARK)
    # OWNER block
    text(s3, sx + Inches(0.1), flow_y + Inches(1.5), stage_w - Inches(0.2), Inches(0.18),
         l1, 7, bold=True, color=GREY_LABEL)
    text(s3, sx + Inches(0.1), flow_y + Inches(1.68), stage_w - Inches(0.2), Inches(0.7),
         v1, 9, color=TEXT)
    # Output block
    text(s3, sx + Inches(0.1), flow_y + Inches(2.55), stage_w - Inches(0.2), Inches(0.18),
         l2, 7, bold=True, color=GREY_LABEL)
    text(s3, sx + Inches(0.1), flow_y + Inches(2.73), stage_w - Inches(0.2), Inches(0.7),
         v2, 9, color=TEXT)

# Gate 1 marker (between Funnel/0 and Reviewing/1)
g1_x = Inches(0.4) + stage_w + Inches(0.025)
rect(s3, g1_x, flow_y, Inches(0.05), flow_h, GREEN_DARKEST)
text(s3, g1_x - Inches(1.0), flow_y - Inches(0.32), Inches(2.1), Inches(0.28),
     "▲  GATE 1 · TRIAGE  ·  weekly", 10, bold=True, color=GREEN_DARKEST, align=PP_ALIGN.CENTER)

# Gate 2 marker (between Analyzing/2 and Backlog/3)
g2_x = Inches(0.4) + 3 * stage_w + 3 * Inches(0.05) - Inches(0.025)
rect(s3, g2_x, flow_y, Inches(0.05), flow_h, GREEN_DARKEST)
text(s3, g2_x - Inches(1.0), flow_y - Inches(0.32), Inches(2.1), Inches(0.28),
     "▲  GATE 2 · PORTFOLIO  ·  monthly", 10, bold=True, color=GREEN_DARKEST, align=PP_ALIGN.CENTER)

# Bottom note
note_y = Inches(6.5)
round_rect(s3, Inches(0.4), note_y, Inches(12.55), Inches(0.55),
           GREEN_VLIGHT, GREEN, Pt(1), 0.04)
text(s3, Inches(0.55), note_y + Inches(0.06), Inches(12), Inches(0.18),
     "BYPASS RULE",
     8, bold=True, color=GREEN_DARK)
text(s3, Inches(0.55), note_y + Inches(0.24), Inches(12.2), Inches(0.30),
     "Size = S / M  ·  bypass this Portfolio Board entirely · enter Squad's Team Kanban directly. "
     "Only L / XL Initiatives flow through Funnel → Reviewing → Analyzing → Gate 2.",
     9, color=TEXT)

footer(s3, 3, 4)


# ============================================================================
# SLIDE 4 — GATES DEEP DIVE (Gate 1 + Gate 2 side by side)
# ============================================================================
s4 = prs.slides.add_slide(prs.slide_layouts[6])
header_band(s4, "GATES DEEP DIVE  ·  WHO · WHEN · HOW",
            "The two governance checkpoints that gate quality (Gate 1) and investment (Gate 2)")

# Two columns side by side
gate_y = Inches(1.2)
gate_h = Inches(5.95)
gate_w = Inches(6.18)
gate_gap = Inches(0.18)

gates = [
    ("GATE 1  ·  TRIAGE",
     "Quality gate · entry validation · DoR check",
     "WEEKLY  ·  30 min",
     GREEN_MID, WHITE,
     [
         ("PURPOSE",
          "Confirms incoming Demand has enough business context to be triaged. "
          "Filters noise · prevents wasted analysis effort downstream."),
         ("WHO PARTICIPATES",
          "▸ CHAIR  ·  Demand Manager (rotating role)\n"
          "▸ ATTENDEES  ·  Champions (one per VS) · Architect (cross-VS standards)\n"
          "▸ OBSERVERS  ·  PM stewards (Heather + P&T) · Sponsor on request"),
         ("WHEN",
          "Every Wednesday 09:00–09:30 ET · standing meeting · cancelled only if zero new Demands"),
         ("HOW · STEP-BY-STEP",
          "1. Demand Manager pulls all NEW + ROUTED tickets from prior 7 days\n"
          "2. For each: validate DoR checklist (Title · Description · BV narrative · AC · VS · OPCO · Sponsor)\n"
          "3. Decision: ACCEPT (→ Reviewing) · BOUNCE (→ Requestor with reason) · MERGE (→ existing item)\n"
          "4. Auto-routing recorded in Jira · @notify Champion of accepted items\n"
          "5. KPI captured: DoR pass rate · bounce rate · channel adoption"),
         ("INPUTS",
          "NEW + ROUTED tickets · DoR checklist · Value Stream policies"),
         ("OUTPUTS",
          "TRIAGED status · BOUNCED status · merge links · weekly KPI report"),
     ]),
    ("GATE 2  ·  PORTFOLIO COMMITTEE",
     "Investment gate · Go/No-Go · capacity slotting",
     "MONTHLY  ·  90 min",
     GREEN_DARK, WHITE,
     [
         ("PURPOSE",
          "Approves investment · slots capacity · re-ranks portfolio backlog by Pain Score + CoS. "
          "Gate where Initiatives become committed work."),
         ("WHO PARTICIPATES",
          "▸ CHAIR  ·  Vitor Cruz (VP IT Business Solutions Delivery)\n"
          "▸ DECIDERS  ·  Carlos Garbin (ITBA Director) · Kumar Jaladanki (CS Director)\n"
          "▸ ATTENDEES  ·  PM stewards (Heather + P&T pool) · Architect · Travis (PMO Lead · for L/XL)\n"
          "▸ OBSERVERS  ·  Sponsors on relevant items only"),
         ("WHEN",
          "First Tuesday of the month · 13:00–14:30 ET · pre-read circulated 3 business days ahead"),
         ("HOW · STEP-BY-STEP",
          "1. Pre-read: all Initiatives in ANALYZING with Pain Score computed + Discovery Report attached\n"
          "2. Walk-through by Pain Score (highest first) · Architect comments on feasibility\n"
          "3. Decision per Initiative: GO (→ Backlog · capacity slot) · NO-GO (→ archive · rationale captured) · DEFER (→ next month)\n"
          "4. Backlog re-ranked by CoS + Pain Score · capacity heat-map updated\n"
          "5. Decisions logged in Jira · @notify Sponsors + Squad Leads"),
         ("INPUTS",
          "TRIAGED Initiatives · Pain Scores · Discovery Reports · capacity available per VS · vendor pipeline status"),
         ("OUTPUTS",
          "Ready_For_Exec status · capacity slot · Backlog rank · Layer 0 VP Dashboard updated"),
     ]),
]

for i, (label, sub, cadence, fill, txc, blocks) in enumerate(gates):
    gx = Inches(0.4 + i * (6.18 + 0.18))

    # Outer
    round_rect(s4, gx, gate_y, gate_w, gate_h, WHITE, fill, Pt(2.5), 0.04)

    # Header (top band)
    round_rect(s4, gx, gate_y, gate_w, Inches(0.85), fill, radius=0.04)
    text(s4, gx + Inches(0.2), gate_y + Inches(0.08), gate_w - Inches(0.4), Inches(0.3),
         label, 16, bold=True, color=txc)
    text(s4, gx + Inches(0.2), gate_y + Inches(0.40), gate_w - Inches(0.4), Inches(0.22),
         sub, 10, italic=True, color=txc)
    text(s4, gx + Inches(0.2), gate_y + Inches(0.61), gate_w - Inches(0.4), Inches(0.22),
         "▸ " + cadence, 10, bold=True, color=GREEN_LIGHT)

    # Body — 6 blocks of detail
    by = gate_y + Inches(1.0)
    block_h = Inches(0.78)
    for label_b, content in blocks:
        round_rect(s4, gx + Inches(0.15), by, gate_w - Inches(0.3), block_h,
                   GREEN_VLIGHT, fill, Pt(0.5), 0.04)
        text(s4, gx + Inches(0.25), by + Inches(0.04), gate_w - Inches(0.5), Inches(0.18),
             label_b, 8, bold=True, color=GREEN_DARKEST)
        text(s4, gx + Inches(0.25), by + Inches(0.20), gate_w - Inches(0.5), block_h - Inches(0.22),
             content, 8, color=TEXT)
        by += block_h + Inches(0.04)

footer(s4, 4, 4)


prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Slides: {len(prs.slides)}")
