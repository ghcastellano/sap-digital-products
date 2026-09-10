"""
PPTX Parser — Extract text, speaker notes, and metadata from PowerPoint files
"""
from pathlib import Path
from datetime import datetime

from pptx import Presentation


def parse_pptx(file_path: Path) -> dict:
    """
    Parse a PPTX file and return structured content.

    Returns:
        {
            'slides': [
                {'number': 1, 'title': '...', 'body': '...', 'notes': '...'},
                ...
            ],
            'metadata': {'slide_count': N, 'created': '...', 'modified': '...'}
        }
    """
    prs = Presentation(str(file_path))

    slides_data = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        body_lines = []
        notes = ""

        # Try to get the title placeholder
        try:
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass

        # Extract all text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        body_lines.append(text)

        # Speaker notes
        if slide.has_notes_slide:
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                pass

        # If no title was found, use the first substantial line
        if not title and body_lines:
            for line in body_lines:
                if len(line) > 5 and not line.isdigit():
                    title = line
                    break

        slides_data.append({
            "number": i,
            "title": title,
            "body": "\n".join(body_lines),
            "notes": notes,
        })

    # File metadata
    core_props = prs.core_properties
    created = ""
    modified = ""
    try:
        if core_props.created:
            created = core_props.created.isoformat()
    except Exception:
        pass
    try:
        if core_props.modified:
            modified = core_props.modified.isoformat()
    except Exception:
        pass

    return {
        "slides": slides_data,
        "metadata": {
            "slide_count": len(slides_data),
            "created": created,
            "modified": modified,
        },
    }
