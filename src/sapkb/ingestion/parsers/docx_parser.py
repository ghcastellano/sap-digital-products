"""
DOCX Parser — Extract text from Word documents
"""
from pathlib import Path

from docx import Document


def parse_docx(file_path: Path) -> dict:
    """
    Parse a DOCX file and return text organized by sections.

    Returns:
        {
            'sections': [{'heading': '...', 'text': '...'}],
            'metadata': {'paragraph_count': N}
        }
    """
    doc = Document(str(file_path))
    sections = []
    current_heading = ""
    current_text = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style and para.style.name.startswith("Heading"):
            # Save previous section
            if current_text:
                sections.append({
                    "heading": current_heading,
                    "text": "\n".join(current_text),
                })
            current_heading = text
            current_text = []
        else:
            current_text.append(text)

    # Save last section
    if current_text:
        sections.append({
            "heading": current_heading,
            "text": "\n".join(current_text),
        })

    return {
        "sections": sections,
        "metadata": {"paragraph_count": len(doc.paragraphs)},
    }
