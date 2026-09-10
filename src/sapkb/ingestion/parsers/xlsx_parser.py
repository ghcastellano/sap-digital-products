"""
XLSX Parser — Extract tabular content from Excel workbooks.

Emits one section per sheet, with the sheet content rendered as a markdown
table so downstream chunking (notes chunker) can index it.
"""
from pathlib import Path

from openpyxl import load_workbook


def _cell_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(file_path: Path) -> dict:
    """Return sheets as markdown-rendered sections.

    {
        'sections': [{'heading': '<sheet name>', 'text': '<markdown table>'}],
        'metadata': {'sheet_count': N, 'total_rows': N}
    }
    """
    wb = load_workbook(str(file_path), data_only=True, read_only=True)
    sections = []
    total_rows = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=False):
            values = [_cell_text(c) for c in row]
            if any(v for v in values):
                rows.append(values)

        total_rows += len(rows)
        if not rows:
            continue

        # Normalize widths
        max_cols = max(len(r) for r in rows)
        rows = [r + [""] * (max_cols - len(r)) for r in rows]

        # Treat first row as header
        header = rows[0]
        body = rows[1:]

        md_lines = []
        md_lines.append("| " + " | ".join(header) + " |")
        md_lines.append("| " + " | ".join("---" for _ in header) + " |")
        for r in body:
            safe = [c.replace("|", "/").replace("\n", " ") for c in r]
            md_lines.append("| " + " | ".join(safe) + " |")

        sections.append({
            "heading": sheet_name,
            "text": "\n".join(md_lines),
        })

    wb.close()

    return {
        "sections": sections,
        "metadata": {
            "sheet_count": len(wb.sheetnames),
            "total_rows": total_rows,
        },
    }
