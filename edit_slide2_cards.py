"""Replace slide 2 cards with 1 OVS + 5 DVS, listing all apps."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FILE = "Material em andamento/Customer Services - Proposed Value Streams - SAFe Reframe v1.pptx"

NAVY  = RGBColor(0x0F, 0x1B, 0x4F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x2C, 0x3E, 0x50)
BLUE_MID = RGBColor(0x2C, 0x57, 0x8B)


def add_dvs_card(slide, x, y, w, h, title, domain, desc, apps_label, apps,
                 *, font_title=10.5, font_apps=8.0):
    """Compact card: navy header bar + white body, mirrors original 3-card style."""
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y + 0.08), Inches(w), Inches(h - 0.08))
    body.fill.solid(); body.fill.fore_color.rgb = WHITE
    body.line.color.rgb = NAVY; body.line.width = Pt(0.5)
    body.adjustments[0] = 0.05

    hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x - 0.02), Inches(y), Inches(w + 0.04), Inches(0.78))
    hdr.fill.solid(); hdr.fill.fore_color.rgb = NAVY
    hdr.line.fill.background()
    hdr.adjustments[0] = 0.10

    # Title
    tb = slide.shapes.add_textbox(Inches(x), Inches(y + 0.10), Inches(w), Inches(0.62))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r = p1.add_run(); r.text = title
    r.font.name = "Arial"; r.font.size = Pt(font_title); r.font.bold = True; r.font.color.rgb = WHITE

    # Domain
    inner_y = y + 0.92
    tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(inner_y), Inches(w - 0.20), Inches(0.20))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r = p1.add_run(); r.text = domain
    r.font.name = "Arial"; r.font.size = Pt(7.5); r.font.bold = True; r.font.italic = True; r.font.color.rgb = DARK

    # Description (italic)
    tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(inner_y + 0.22), Inches(w - 0.20), Inches(0.55))
    tf = tb.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r = p1.add_run(); r.text = desc
    r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.italic = True; r.font.color.rgb = DARK

    # Divider
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x + 0.20), Inches(inner_y + 0.85),
                                Inches(w - 0.40), Inches(0.02))
    ln.fill.solid(); ln.fill.fore_color.rgb = NAVY
    ln.line.fill.background()

    # Apps label
    tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(inner_y + 0.93),
                                  Inches(w - 0.20), Inches(0.20))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r = p1.add_run(); r.text = apps_label
    r.font.name = "Arial"; r.font.size = Pt(7.5); r.font.bold = True; r.font.color.rgb = DARK

    # Apps text
    apps_y = inner_y + 1.15
    tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(apps_y),
                                  Inches(w - 0.20), Inches(y + h - apps_y - 0.10))
    tf = tb.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r = p1.add_run(); r.text = apps
    r.font.name = "Arial"; r.font.size = Pt(font_apps); r.font.color.rgb = DARK


def add_ovs_card(slide, x, y, w, h, title, domain, desc, stages_label, stages):
    """Wide card for the operational VS at top."""
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y + 0.05), Inches(w), Inches(h - 0.05))
    body.fill.solid(); body.fill.fore_color.rgb = WHITE
    body.line.color.rgb = NAVY; body.line.width = Pt(0.5)
    body.adjustments[0] = 0.04

    hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x - 0.02), Inches(y), Inches(w + 0.04), Inches(0.50))
    hdr.fill.solid(); hdr.fill.fore_color.rgb = NAVY
    hdr.line.fill.background()
    hdr.adjustments[0] = 0.06

    # Title left
    tb = slide.shapes.add_textbox(Inches(x + 0.20), Inches(y + 0.08), Inches(w - 3.10), Inches(0.40))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = title
    r.font.name = "Arial"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE

    # Badge right
    tb = slide.shapes.add_textbox(Inches(x + w - 3.10), Inches(y + 0.10), Inches(2.90), Inches(0.36))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.RIGHT
    r = p1.add_run(); r.text = "OPERATIONAL VALUE STREAM (1)"
    r.font.name = "Arial"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = WHITE

    # Domain
    inner_y = y + 0.62
    tb = slide.shapes.add_textbox(Inches(x + 0.20), Inches(inner_y), Inches(w - 0.40), Inches(0.20))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = domain
    r.font.name = "Arial"; r.font.size = Pt(8.5); r.font.bold = True; r.font.italic = True; r.font.color.rgb = DARK

    # Description
    tb = slide.shapes.add_textbox(Inches(x + 0.20), Inches(inner_y + 0.22), Inches(w - 0.40), Inches(0.34))
    tf = tb.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = desc
    r.font.name = "Arial"; r.font.size = Pt(11); r.font.italic = True; r.font.color.rgb = DARK

    # Divider
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x + 0.20), Inches(inner_y + 0.62),
                                Inches(w - 0.40), Inches(0.02))
    ln.fill.solid(); ln.fill.fore_color.rgb = NAVY
    ln.line.fill.background()

    # Stages label
    tb = slide.shapes.add_textbox(Inches(x + 0.20), Inches(inner_y + 0.68),
                                  Inches(w - 0.40), Inches(0.20))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = stages_label
    r.font.name = "Arial"; r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = DARK

    # Stages content
    tb = slide.shapes.add_textbox(Inches(x + 0.20), Inches(inner_y + 0.88),
                                  Inches(w - 0.40), Inches(0.32))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = stages
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = NAVY


def main():
    prs = Presentation(FILE)
    slide = prs.slides[1]

    # remove all card shapes (everything past header bar — idx >= 2)
    shapes_to_remove = list(slide.shapes)[2:]
    for sh in shapes_to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    # OVS card on top
    add_ovs_card(slide,
        x=0.50, y=1.05, w=12.33, h=1.55,
        title="Meter-to-Cash",
        domain="DOMAIN · Customer Service · End-to-End Business Outcome across 7 OPCOs",
        desc=("Capture every meter read · invoice on time · collect cash · "
              "serve every customer interaction — anchored on SAP (Billing · FICA · CRM · Metering)"),
        stages_label="STAGES",
        stages="METER  →  BILL  →  COLLECT  →  SERVE   ·   PLATFORM (transversal enabler)",
    )

    # Label above DVS row
    tb = slide.shapes.add_textbox(Inches(0.50), Inches(2.78), Inches(12.33), Inches(0.25))
    p1 = tb.text_frame.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT
    r = p1.add_run(); r.text = "DEVELOPMENT VALUE STREAMS (5)  ·  ONE PRODUCT MANAGER PER DVS"
    r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = NAVY

    # 5 DVS cards in a row
    dvs_y = 3.10
    dvs_h = 3.95
    n = 5
    margin = 0.50
    gap = 0.10
    total_w = 13.33 - 2 * margin
    card_w = (total_w - (n - 1) * gap) / n  # ≈ 2.39

    dvs_cards = [
        ("Meter Data & Devices",
         "DOMAIN · Customer Service / Metering",
         "Capture every read · zero data loss to billing",
         "Smart Metering\nMDM (Meter Data Mgmt)\nSAP IS-U Device\nGroup 8760 (EDI)\nField meter read interfaces"),
        ("Billing & Invoicing",
         "DOMAIN · Customer Service / Billing",
         "Run the bill on time · every day · every OPCO",
         "SAP ECC (FI/SD)\nStreamserve / OpenText\nAddress Validation (Precisely)\nPrint & e-delivery interfaces"),
        ("Collections · FICA · Tax",
         "DOMAIN · Customer Service / Finance",
         "Convert invoices into cash · comply at line level",
         "SAP FICA\nSOVOS Taxware\nPayment processing interfaces\n(banks · lockbox · 3rd-party)\nDunning & collections workflows"),
        ("Customer Interaction",
         "DOMAIN · Customer Service / CX",
         "Unify CRM · contact center · digital · feedback",
         "SAP CRM\nGenesys (Infosys)\nNuance IVR / CTI\nMedallia (formerly VHT)\nPureConnect / PureEngage (legacy)"),
        ("Platform & Enablers",
         "DOMAIN · Shared Service · Transversal",
         "Keep the factory running for every other DVS",
         "SAP Basis\nSAP Gold Client\nSAP Enable Now (WPB)\nBlue Prism (RPA)\nTricentis TOSCA\nJira"),
    ]

    for i, (title, domain, desc, apps) in enumerate(dvs_cards):
        x = margin + i * (card_w + gap)
        add_dvs_card(slide, x, dvs_y, card_w, dvs_h, title, domain, desc, "ALL APPLICATIONS", apps)

    prs.save(FILE)
    print(f"Saved → {FILE}")
    print(f"Slide 2 shapes now: {len(slide.shapes)}")


if __name__ == "__main__":
    main()
